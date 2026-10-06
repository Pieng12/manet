"""Export the independent range pilot; never invokes the four research metrics."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

SHEETS = ("Overview", "Source Status", "Test Points", "RX Samples", "GPS Track", "Diagnostics")


def cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def finite_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def gps_distance_quality(row: dict, run: dict) -> dict:
    distance = row.get("distance_m")
    accuracy = (row.get("location") or {}).get("accuracy_m")
    source_accuracy = run.get("source_accuracy_m")
    usable = (finite_number(distance) and distance >= 0 and
              finite_number(accuracy) and 0 <= accuracy <= 20)
    known_source = finite_number(source_accuracy) and source_accuracy >= 0
    radii_sum = accuracy + source_accuracy if usable and known_source else None
    quality = ("unavailable_or_inaccurate" if not usable else
               "source_uncertainty_unknown" if not known_source else
               "not_distinguishable_from_location_uncertainty" if distance <= radii_sum else
               "gps_estimate")
    descriptions = {
        "unavailable_or_inaccurate": "Lokasi tidak tersedia atau tidak layak",
        "source_uncertainty_unknown": "Ketidakpastian posisi ESP tidak terukur",
        "not_distinguishable_from_location_uncertainty": "Jarak dekat belum dapat dibedakan dari ketidakpastian lokasi",
        "gps_estimate": "Estimasi GPS; bukan jarak terukur",
    }
    return {"gps_distance_quality": quality, "gps_accuracy_radii_sum_m": radii_sum,
            "gps_distance_explanation": descriptions[quality]}


def export_range(directory: Path) -> Path:
    data = json.loads((directory / "range_trial.json").read_text(encoding="utf-8-sig"))
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8-sig"))
    source_path = directory / "source_status.json"
    source = json.loads(source_path.read_text(encoding="utf-8-sig")) if source_path.exists() else {"stop_confirmed": False}
    events_path = directory / "source_events.json"
    source_events = json.loads(events_path.read_text(encoding="utf-8-sig")) if events_path.exists() else []
    run = data["run"]
    if not run or run["run_id"] != manifest["session_id"]:
        raise ValueError("Range run and manifest do not match")
    common = {k: run.get(k) for k in (
        "run_id", "session_id", "trial_id", "algorithm", "direction", "android_build_id",
        "source_sender_crc", "source_timestamp_ms", "source_latitude", "source_longitude",
        "source_accuracy_m", "source_position_method",
    )}
    common.update(esp_build_id=manifest.get("esp_build_id"), tx_power_actual_dbm=manifest.get("tx_power_actual_dbm"))
    received = data.get("receives", [])
    distances = [r["distance_m"] for r in received if r.get("phase") == "session" and isinstance(r.get("distance_m"), (int, float)) and math.isfinite(r["distance_m"])]
    overview = {**run, "dataset_kind": "range_pilot_not_main_experiment",
                "rx_observation_count": len(received), "coding": "unknown_s2_or_s8",
                "farthest_observed_horizontal_m": max(distances, default=None),
                "source_status_verified": source.get("stop_confirmed", False),
                "source_advertise_failed_count": source.get("advertise_failed_count"),
                "distance_is_estimate": True, "not_a_maximum_range_or_reliability_measurement": True}
    overview.update(timestamp_basis="unix_milliseconds_UTC",
                    max_hp_accuracy_m=20, max_rx_location_time_delta_ms=5000,
                    source_position_uncertainty_unknown=run.get("source_accuracy_m") is None)
    resolved_distances = [r["distance_m"] for r in received if r.get("phase") == "session"
                          and gps_distance_quality(r, run)["gps_distance_quality"] == "gps_estimate"]
    completed_received_points = [point for point in data.get("points", [])
                                 if point.get("status") == "completed" and point.get("receive_count", 0) > 0]
    for column in ("measured_horizontal_m", "measured_3d_m"):
        values = [point[column] for point in completed_received_points
                  if finite_number(point.get(column)) and point[column] >= 0]
        overview["farthest_received_manual_" + ("horizontal_m" if column == "measured_horizontal_m" else "3d_m")] = max(values, default=None)
    overview.update(
        farthest_resolved_gps_estimate_m=max(resolved_distances, default=None),
        gps_distance_note="Jarak GPS merupakan estimasi horizontal. Ringkasan lama tetap disimpan, bukan bukti jarak aktual atau maksimum universal.",
        uncertainty_note="Jumlah radius akurasi HP dan ESP adalah penanda kehati-hatian, bukan batas galat pasti atau interval statistik gabungan. Nilai GPS tidak dikurangi atau dinolkan.",
        manual_measurement_note="Input per titik dipisahkan dari GPS. Metode meteran, denah berskala, dan perkiraan manual harus dibedakan; ringkasan manual hanya dari titik selesai yang menerima paket.",
        height_note="Beda tinggi adalah HP dikurangi ESP, diisi manual. Jarak 3D dihitung hanya jika horizontal dan beda tinggi diisi; bukan dari GPS, RSSI, atau nomor lantai.",
    )
    points = [{**common, **row, "source_run_completed": source.get("completed_duration", False),
               "source_stop_confirmed": source.get("stop_confirmed", False)} for row in data.get("points", [])]
    rx = []
    for row in received:
        location = row.get("location") or {}
        rx.append({**common, **row, "location_timestamp_ms": location.get("timestamp_ms"),
                   "latitude": location.get("latitude"), "longitude": location.get("longitude"),
                   "accuracy_m": location.get("accuracy_m"), "coding": "unknown_s2_or_s8",
                   **gps_distance_quality(row, run)})
    tables = [
        [{"field": key, "value": value} for key, value in overview.items()],
        [{"field": key, "value": value} for key, value in source.items()],
        points, rx,
        [{**common, **row} for row in data.get("positions", [])],
        [{**common, "origin": "android_pilot", **row} for row in data.get("diagnostics", [])]
        + [{**common, "origin": "esp_serial", **row} for row in source_events],
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (name, rows) in enumerate(zip(SHEETS, tables), 1):
        for row in rows:
            for key in ("timestamp_ms", "started_at_ms", "ended_at_ms"):
                if isinstance(row.get(key), (int, float)):
                    row[key.removesuffix("_ms") + "_utc"] = datetime.fromtimestamp(row[key] / 1000, timezone.utc).isoformat(timespec="milliseconds")
        sheet = workbook.create_sheet(name)
        headers = list(dict.fromkeys(key for row in rows for key in row)) or ["status"]
        sheet.append(headers)
        for row in rows:
            sheet.append([cell(row.get(key)) for key in headers])
        for header in sheet[1]:
            header.fill = PatternFill("solid", fgColor="1F4E78")
            header.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if rows:
            table = Table(displayName=f"RangeTable{index}", ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(48, max(16, len(headers[column-1]) + 3))
        for row in sheet.iter_rows(min_row=2):
            for value in row:
                value.alignment = Alignment(vertical="top", wrap_text=True)
        with (directory / (name.lower().replace(" ", "_") + ".csv")).open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows({key: cell(row.get(key)) for key in headers} for row in rows)
    destination = directory / "resqmesh_range_analysis.xlsx"
    workbook.save(destination)
    return destination


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    print(export_range(parser.parse_args().output))
