"""Presentation only: metrics and canonical JSON/CSV remain unchanged."""
from __future__ import annotations

from math import ceil

from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_report import _write_table


METHOD_NAMES = {
    "basic_flooding": "Basic Flooding",
    "trickle": "Trickle (suppression c/k)",
    "trickle_mpl": "Trickle-MPL BLE",
    "trickle_no_suppression": "Trickle tanpa suppression (historis)",
    "trickle_neighbor_status": "Neighbor Status (historis)",
}
SCENARIO_NAMES = {
    "S0_STABLE": "S0 - Jaringan stabil",
    "S1_BRANCH_DELAYED": "S1 - Cabang D/E terlambat aktif",
    "S2_POST_DATA_STOP": "S2 - Late join setelah DATA berhenti",
    "S0_MAIN": "S0 - Semua node siap",
    "S1_DELAYED_RX": "S1 - Scanner D terlambat aktif",
    "S2_LATE_JOIN": "S2 - TX/RX E terlambat aktif",
}
SCENARIO_DETAILS = {
    "S0_STABLE": "Semua node aktif; observasi 180 detik sejak DATA sumber pertama.",
    "S1_BRANCH_DELAYED": "D/E TX/RX OFF sebelum SOS, D lalu E ON pada detik 60; observasi 180 detik.",
    "S2_POST_DATA_STOP": "E TX/RX OFF sebelum SOS, ON detik 300; observasi 420 detik. Stop timer DATA dan repair harus dibuktikan terpisah.",
    "S0_MAIN": "Kondisi utama: seluruh node aktif sejak awal.",
    "S1_DELAYED_RX": "Scanner D OFF lalu ON; delay mengikuti manifest, tanpa menghapus state.",
    "S2_LATE_JOIN": "TX/RX E OFF lalu ON; delay mengikuti manifest, tanpa menghapus state.",
}
FIELD_HELP = {
    "trial_id": ("ID trial", "Identitas attempt; INVALID tetap disimpan, pengganti memiliki ID baru.", "teks"),
    "method": ("Metode", "Kode algoritma; label lengkap ada di Mulai Di Sini.", "kategori"),
    "scenario": ("Skenario", "S0 kondisi utama; S1/S2 kondisi pemulihan pendukung.", "kategori"),
    "block": ("Blok pengacakan", "Blok pengulangan seimbang; kosong jika tidak tercatat, bukan blok nol.", "blok"),
    "result": ("Hasil trial", "SUCCESS: seluruh target menerima; FAILED_DELIVERY: prosedur sah tetapi tidak semua menerima; INVALID: prosedur/bukti tidak sah.", "kategori"),
    "valid": ("Trial sah", "SUCCESS dan FAILED_DELIVERY sah; INVALID tidak masuk ringkasan metrik.", "boolean"),
    "M": ("M - Pesan sumber", "Jumlah MessageKey SOS unik yang dibuat sumber, bukan jumlah burst atau retransmisi.", "pesan"),
    "N": ("N - Target penerima", "Lima ESP target A, B, C, D, E; sumber Android tidak masuk denominator DSR.", "node"),
    "R": ("R - Penerimaan DATA", "Jumlah RX DATA logis valid pada target; burst ulang dihitung, callback ulang burst sama tidak.", "RX logis"),
    "U": ("U - Pasangan unik sukses", "Pasangan MessageKey-target yang menerima setidaknya satu kali. Bukan jumlah seluruh RX.", "pasangan"),
    "dsr_percent": ("DSR (%)", "U/(M*5)*100. Rasio jumlah; FAILED_DELIVERY sah tetap masuk denominator.", "% (0-100)"),
    "e2e_mean_ms": ("Delay pasangan sukses (ms)", "Rata-rata RX fisik pertama - TX DATA sumber pertama setelah koreksi jam; gagal tidak diberi delay nol.", "ms"),
    "ldr_percent": ("LDR (%)", "(R-U)/R*100. CONTROL tidak termasuk; R=0 menghasilkan nilai kosong.", "% (0-100)"),
    "data_tx": ("DATA berhasil mulai", "Burst DATA sukses mulai dalam window pada seluruh enam node, bukan request atau paket RF.", "burst"),
    "control_tx": ("CONTROL berhasil mulai", "Burst STATUS/CONTROL sukses mulai dalam window, termasuk discovery/probe/repair jika ada.", "burst"),
    "network_overhead": ("Total overhead window", "DATA_TX + CONTROL_TX; bukan konsumsi energi atau airtime.", "burst"),
    "setup_control_tx": ("CONTROL persiapan", "STATUS sukses sebelum window trial yang sama; dilaporkan terpisah, tidak disembunyikan.", "burst"),
    "setup_plus_window_tx": ("Persiapan + window", "CONTROL persiapan + total overhead window; bukan seluruh biaya siklus hidup.", "burst"),
    "valid_trials": ("Trial sah", "Jumlah SUCCESS + FAILED_DELIVERY; INVALID dikecualikan.", "trial"),
    "invalid_trials": ("Trial INVALID", "Attempt tidak sah, tetap tersedia pada Invalid Trials dan log mentah.", "attempt"),
    "failed_delivery_trials": ("Trial gagal pengiriman", "FAILED_DELIVERY sah; tidak sama dengan INVALID dan tidak dikeluarkan dari DSR.", "trial"),
    "successful_pairs": ("Pasangan sukses", "Jumlah pasangan pesan-target berhasil; satu trial dapat memiliki lima pasangan.", "pasangan"),
    "failed_pairs": ("Pasangan gagal", "M*5-U, bukan jumlah trial INVALID.", "pasangan"),
    "defined_delay_pairs": ("Pasangan dengan delay", "Jumlah pasangan sukses yang mempunyai delay terdefinisi, bukan jumlah trial.", "pasangan"),
    "recovery_pairs": ("Pasangan dengan recovery", "Pasangan yang mempunyai RX setelah ON terkonfirmasi dan delay recovery terdefinisi.", "pasangan"),
    "recovery_mean_ms": ("Recovery rata-rata (ms)", "Rata-rata RX pertama - event ON terkonfirmasi pada jam monoton target yang sama; bukan waktu proses repair.", "ms"),
    "recovery_delay_ms": ("Delay recovery (ms)", "RX fisik pertama - event ON terkonfirmasi, memakai jam monoton dan clock epoch lokal yang sama. Bukan waktu respons command atau durasi repair.", "ms"),
    "recovery_on_confirmed_at_ms": ("Event ON (epoch ms)", "Timestamp event aktivasi scanner setelah koreksi jam; bukan timestamp respons command. Delay recovery memakai kolom monoton.", "epoch ms"),
    "recovery_on_monotonic_ms": ("Event ON monoton (ms)", "Waktu lokal event aktivasi scanner terkonfirmasi pada target.", "ms monoton"),
    "first_rx_monotonic_ms": ("RX pertama monoton (ms)", "Waktu fisik penerimaan pertama pada jam monoton target yang sama.", "ms monoton"),
    "recovery_clock_domain": ("Domain jam recovery", "Domain jam ON dan RX harus sama; tidak mencampur jam antarperangkat.", "kategori"),
    "recovery_local_boot_id": ("Incarnation lokal recovery", "Identitas aktivasi lokal pada event ON dan RX; bukan boot_id pemancar paket.", "identitas"),
    "recovery_time_basis": ("Dasar waktu recovery", "same_node_monotonic_event: bukti event lokal cocok; unverified: delay kosong, tanpa fallback ke respons command.", "kategori"),
    "recovery_unavailable_reason": ("Alasan recovery kosong", "NOT_PERTURBED: target tidak diperturbasi; ON_EVENT_MISSING/MONOTONIC_MISSING: bukti kurang; CLOCK_DOMAIN_UNVERIFIED/CLOCK_EPOCH_UNVERIFIED: jam tidak cocok; RX_BEFORE_ON: RX mendahului ON.", "kategori"),
    "defined_trials": ("Nilai terdefinisi", "Jumlah nilai numerik untuk statistik ini; recovery memakai pasangan, metrik lain memakai trial.", "nilai"),
    "mean": ("Rata-rata", "Rata-rata nilai terdefinisi; satuan mengikuti kolom unit.", "sesuai metrik"),
    "median": ("Median", "Nilai tengah dari nilai terdefinisi.", "sesuai metrik"),
    "sample_sd": ("SD sampel", "Simpangan baku sampel; kosong jika jumlah nilai kurang dari dua.", "sesuai metrik"),
    "min": ("Minimum", "Nilai terdefinisi terkecil; kosong bukan nol.", "sesuai metrik"),
    "max": ("Maksimum", "Nilai terdefinisi terbesar; kosong bukan nol.", "sesuai metrik"),
    "invalid_reasons": ("Alasan INVALID", "Alasan asli dari manifest/validator; bukan dasar mengubah hasil agar lebih bagus.", "teks/JSON"),
    "status": ("Status validasi", "PASS: bukti mendukung pemeriksaan; FAIL: kontradiksi; INCONCLUSIVE: bukti belum cukup.", "kategori"),
    "observation_started_at_ms": ("Awal window (epoch ms)", "Window dimulai dari advertising sumber pertama yang berhasil, sesuai manifest.", "epoch ms"),
    "observation_ended_at_ms": ("Akhir window (epoch ms)", "Waktu akhir window pada clock controller; bukan durasi.", "epoch ms"),
}
SHEET_HELP = {
    "Ringkasan Metode": "Perbandingan metode-skenario; hanya trial sah. DSR/LDR rasio jumlah, delay rata-rata pasangan sukses.",
    "Hasil Per Trial": "Empat metrik per attempt beserta biaya DATA/CONTROL/setup. Baris INVALID tidak masuk ringkasan.",
    "Grafik Ringkas": "Grafik langsung di Excel: DSR, delay, LDR, total overhead serta komponennya per skenario.",
    "Metric Definitions": "Definisi rumus penelitian, satuan dan catatan penggunaan.",
    "Descriptive Statistics": "Mean, median, SD sampel, min/max; unit dan jumlah nilai terdefinisi harus dibaca bersama.",
    "Receivers": "Bukti per target: penerimaan pertama, delay, ketidakpastian jam dan recovery setelah ON.",
    "Log Validation": "PASS/FAIL/INCONCLUSIVE untuk bukti log; tidak sama dengan hasil delivery trial.",
    "Invalid Trials": "Attempt INVALID dan alasan asli; tidak dihapus dan tidak masuk ringkasan metrik.",
    "Method Parameters": "Parameter yang benar-benar berlaku untuk tiap metode; N/A berarti tidak dipakai.",
    "MPL Parameters": "Parameter timer DATA/CONTROL, discovery, retry, repair dan freshness mode MPL.",
    "Mechanism Diagnostics": "Frekuensi keputusan mekanisme sepanjang trial; bukan overhead RF tambahan.",
    "MPL Diagnostics": "Event timer dan repair MPL untuk audit; alias native bukan TX tambahan.",
    "STATUS Diagnostics": "Alasan dan hasil request/start/failure STATUS; hanya start sukses masuk overhead.",
    "PHY Evidence": "Konfigurasi radio dan batas klaim; Coded tidak otomatis membuktikan S8.",
    "Participation": "Command perturbasi serta konfirmasi scanner/TX/RX ON/OFF tanpa reset state.",
    "Trial Metrics": "Tabel teknis per trial, mempertahankan nama kolom dan semua nilai asal.",
    "Method Scenario Summary": "Agregat teknis lengkap: total, rata-rata trial dan rata-rata pasangan tidak disamakan.",
    "Charts": "Tautan grafik SVG eksternal; tetap tersedia selain grafik langsung di Excel.",
    "Overview": "Metadata manifest lengkap, termasuk konfigurasi bersarang; bukan ringkasan metrik.",
    "All Events": "Log lengkap dengan nama kolom teknis; raw JSON/CSV tetap tersedia tanpa perubahan.",
    "Adaptive Pilot Review": "Perbandingan eksploratif kebijakan historis, bukan uji statistik atau aturan INVALID.",
    "Kamus Kolom": "Arti singkatan, nama kolom teknis, satuan dan perbedaan rata-rata dengan total.",
}

MPL_PARAMETER_HELP = {
    "mpl_data_imin_ms": ("ms", "Interval DATA awal; listen-only sampai I/2, lalu t dipilih acak."),
    "mpl_data_imax_ms": ("ms", "Batas interval DATA efektif, bukan jumlah doubling."),
    "mpl_data_k": ("counter", "DATA normal diizinkan bila c < k; repair override terpisah dan terbatas."),
    "mpl_data_expirations": ("interval", "Jumlah interval DATA selesai sebelum timer berhenti, termasuk interval suppressed; buffer tetap disimpan."),
    "mpl_control_imin_ms": ("ms", "Interval CONTROL awal untuk ringkasan inventory/discovery."),
    "mpl_control_imax_ms": ("ms", "Batas interval CONTROL efektif sebelum doubling berhenti."),
    "mpl_control_k": ("counter", "CONTROL normal diizinkan bila c < k; bootstrap/deficit terlindungi secara terbatas."),
    "mpl_control_expirations": ("interval", "Jumlah interval CONTROL selesai sebelum timer berhenti; bukan batas TX sukses."),
    "mpl_repair_cooldown_ms": ("ms", "Jeda evaluasi ulang repair untuk episode state-peer yang sama."),
    "mpl_repair_budget": ("kesempatan", "Batas repair/reset per episode; bukan batas relay SOS sepanjang umur pesan."),
    "mpl_repair_expiry_ms": ("ms", "Umur episode repair; tidak menghapus SOS dari buffer."),
    "mpl_bootstrap_opportunities": ("kesempatan sukses", "Kesempatan CONTROL discovery yang terlindungi secara terbatas."),
    "mpl_retry_ms": ("ms", "Jeda retry setelah native advertising gagal; gagal tidak dihitung TX sukses."),
    "mpl_retry_limit": ("percobaan/interval", "Batas percobaan native dalam satu interval."),
    "mpl_probe_interval_ms": ("ms", "Jeda probe opsional; tidak aktif jika probe limit nol."),
    "mpl_probe_limit": ("probe", "Batas probe maintenance; nol berarti nonaktif."),
    "mpl_freshness_ms": ("ms", "Usia bukti RX fisik yang masih fresh; kedaluwarsa menjadi UNKNOWN, bukan MISSING/HAVE."),
    "mpl_discovery_jitter_ms": ("ms", "Batas jitter discovery; bukan jitter Basic atau pengganti pemilihan t Trickle."),
}


def mpl_parameter_rows(parameters):
    return [{"parameter": key, "value": value,
             "unit": MPL_PARAMETER_HELP.get(key, ("sesuai field", "Parameter asli manifest."))[0],
             "penjelasan": MPL_PARAMETER_HELP.get(key, ("sesuai field", "Parameter asli manifest."))[1]}
            for key, value in parameters.items()]


def _label(field):
    if field.endswith("_trial_mean"):
        original = field.removesuffix("_trial_mean")
        title, meaning, unit = FIELD_HELP.get(original, (original, "Nilai sesuai tabel teknis.", "sesuai metrik"))
        return f"{title} - rata-rata per trial", f"Rata-rata tanpa bobot antar trial sah. {meaning}", unit
    return FIELD_HELP.get(field, (field, "Field teknis asli; maknanya mengikuti event/manifest asal.", "sesuai field"))


def _readable_rows(rows, mapping):
    return [{title: METHOD_NAMES.get(row.get(field), row.get(field)) if field == "method"
             else SCENARIO_NAMES.get(row.get(field), row.get(field)) if field == "scenario"
             else row.get(field) for field, title in mapping} for row in rows]


def _style(sheet, percent_headers=()):
    sheet.conditional_formatting = ConditionalFormattingList()
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.tabColor = "2675A9"
    sheet.row_dimensions[1].height = 48
    headers = [cell.value for cell in sheet[1]]
    for index, cell in enumerate(sheet[1], 1):
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if cell.value in FIELD_HELP or str(cell.value).endswith("_trial_mean"):
            title, meaning, unit = _label(str(cell.value))
            cell.comment = Comment(f"{title}\n{meaning}\nSatuan: {unit}", "ResQMesh")
        field = str(cell.value)
        width = 24 if field in {"method", "scenario", "Metode", "Skenario"} else 32 if "trial" in field.lower() and "id" in field.lower() else 19
        if field in {"definition", "catatan", "Penjelasan", "penjelasan", "meaning", "evidence", "invalid_reasons", "Alasan INVALID", "Arti", "value", "Catatan", "parameter_notes", "transmission_timing", "termination"}:
            width = 64
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = max(48, max(
        (15 * ceil(len(str(cell.value)) / max(12, sheet.column_dimensions[cell.column_letter].width - 2)) + 12
         for cell in sheet[1]), default=48))
    for row in sheet.iter_rows(min_row=2):
        height = 24
        for cell in row:
            field = str(headers[cell.column - 1])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = '0.00"%"' if field in percent_headers or field.endswith("_percent") or field.endswith("_percent_trial_mean") else '#,##0.00' if isinstance(cell.value, float) else '#,##0'
            if isinstance(cell.value, str):
                width = sheet.column_dimensions[cell.column_letter].width
                height = max(height, min(90, 15 * ceil(len(cell.value) / max(12, width - 2))))
        sheet.row_dimensions[row[0].row].height = height
    for field in ("result", "status", "Hasil trial"):
        if field not in headers or sheet.max_row < 2:
            continue
        column = get_column_letter(headers.index(field) + 1)
        cells = f"{column}2:{column}{sheet.max_row}"
        for value, color in (("SUCCESS", "C6EFCE"), ("PASS", "C6EFCE"),
                             ("FAILED_DELIVERY", "FFEB9C"), ("INCONCLUSIVE", "FFEB9C"),
                             ("INVALID", "FFC7CE"), ("FAIL", "FFC7CE")):
            sheet.conditional_formatting.add(cells, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill("solid", fgColor=color)))


def decorate_neighbor_workbook(workbook, manifest, trials, summaries):
    """Add human-facing views without editing source rows or technical cell values."""
    summary_columns = [
        ("method", "Metode"), ("scenario", "Skenario"), ("valid_trials", "Trial sah"),
        ("invalid_trials", "Attempt INVALID"), ("failed_delivery_trials", "Trial gagal pengiriman"),
        ("dsr_percent", "DSR (%)"), ("e2e_mean_ms", "Delay pasangan sukses (ms)"),
        ("ldr_percent", "LDR (%)"), ("network_overhead_trial_mean", "Total overhead rata-rata (burst/trial)"),
        ("data_tx_trial_mean", "DATA rata-rata (burst/trial)"), ("control_tx_trial_mean", "CONTROL rata-rata (burst/trial)"),
        ("setup_control_tx_trial_mean", "CONTROL persiapan rata-rata (burst/trial)"),
        ("setup_plus_window_tx_trial_mean", "Persiapan + window rata-rata (burst/trial)"),
        ("successful_pairs", "Pasangan sukses"), ("failed_pairs", "Pasangan gagal"),
        ("defined_delay_pairs", "Pasangan dengan delay"), ("recovery_pairs", "Pasangan dengan recovery"),
        ("recovery_mean_ms", "Recovery rata-rata (ms)"),
    ]
    trial_columns = [
        ("trial_id", "ID trial"), ("method", "Metode"), ("scenario", "Skenario"),
        ("block", "Blok"), ("result", "Hasil trial"), ("valid", "Trial sah"),
        ("dsr_percent", "DSR (%)"), ("e2e_mean_ms", "Delay pasangan sukses (ms)"),
        ("ldr_percent", "LDR (%)"), ("network_overhead", "Total overhead window (burst)"),
        ("data_tx", "DATA berhasil mulai (burst)"), ("control_tx", "CONTROL berhasil mulai (burst)"),
        ("setup_control_tx", "CONTROL persiapan (burst)"), ("setup_plus_window_tx", "Persiapan + window (burst)"),
        ("M", "M - Pesan sumber"), ("N", "N - Target penerima"),
        ("R", "R - Penerimaan DATA"), ("U", "U - Pasangan unik sukses"),
        ("failed_pairs", "Pasangan gagal"), ("invalid_reasons", "Alasan INVALID"),
    ]
    for name, data, columns in (("Ringkasan Metode", summaries, summary_columns),
                                ("Hasil Per Trial", trials, trial_columns)):
        _write_table(workbook, name, _readable_rows(data, columns), [label for _, label in columns])
        sheet = workbook[name]
        sheet.freeze_panes = "C2" if name == "Ringkasan Metode" else "D2"
        for index, (field, _) in enumerate(columns if data else [], 1):
            title, meaning, unit = _label(field)
            sheet.cell(1, index).comment = Comment(f"Kolom asal: {field}\n{title}\n{meaning}\nSatuan: {unit}", "ResQMesh")
        _style(sheet, {"DSR (%)", "LDR (%)"})

    glossary = []
    for sheet in list(workbook):
        if sheet.title in {"Ringkasan Metode", "Hasil Per Trial"}:
            continue
        _style(sheet)
        if sheet.title in {"Trial Metrics", "Method Scenario Summary", "Receivers"}:
            sheet.freeze_panes = "D2"
        if sheet.title in {"All Events", "MPL Diagnostics", "Overview"}:
            sheet.sheet_properties.tabColor = "777777"
        for cell in sheet[1]:
            title, meaning, unit = _label(str(cell.value))
            if str(cell.value) in FIELD_HELP or str(cell.value).endswith("_trial_mean"):
                glossary.append({"Tabel": sheet.title, "Kolom teknis": cell.value,
                                 "Nama jelas": title, "Arti": meaning, "Satuan": unit})
    _write_table(workbook, "Kamus Kolom", glossary, ("Tabel", "Kolom teknis", "Nama jelas", "Arti", "Satuan"))
    _style(workbook["Kamus Kolom"])
    workbook["Kamus Kolom"].column_dimensions["D"].width = 90

    if "MPL Parameters" in workbook:
        sheet = workbook["MPL Parameters"]
        for row in sheet.iter_rows(min_row=2):
            field = str(row[0].value)
            row[1].comment = Comment(
                "Satuan: " + ("ms" if field.endswith("_ms") else "jumlah/batas")
                + ". Expiration menghentikan timer, bukan menghapus buffer. Nilai asli manifest.", "ResQMesh")
        sheet.column_dimensions["A"].width = 42

    chart_sheet = workbook.create_sheet("Grafik Ringkas")
    chart_sheet.sheet_view.showGridLines = False
    chart_sheet.sheet_view.zoomScale = 70
    chart_sheet.sheet_format.defaultColWidth = 10
    chart_sheet.sheet_properties.tabColor = "29956B"
    chart_sheet["A1"] = "Perbandingan metode per skenario - hanya trial sah"
    chart_sheet["A1"].font = Font(size=16, bold=True)
    chart_sheet["A2"] = "Nilai kosong bukan nol; delay memakai pasangan sukses. CONTROL persiapan ada di Ringkasan Metode."
    chart_sheet["A3"] = "DATA SINTETIS - bukan hasil perangkat fisik" if manifest.get("synthetic_data") is True else "Baca bukti PHY dan validasi log sebelum menarik kesimpulan."
    summary_sheet = workbook["Ringkasan Metode"]
    headers = [label for _, label in summary_columns]
    for scenario_index, scenario in enumerate(dict.fromkeys(row["scenario"] for row in summaries)):
        indices = [index + 2 for index, row in enumerate(summaries) if row["scenario"] == scenario]
        for metric_index, (title, labels, unit) in enumerate((
            ("DSR", ["DSR (%)"], "%"),
            ("Delay pasangan sukses", ["Delay pasangan sukses (ms)"], "ms"),
            ("LDR", ["LDR (%)"], "%"),
            ("Overhead dalam window", ["DATA rata-rata (burst/trial)", "CONTROL rata-rata (burst/trial)", "Total overhead rata-rata (burst/trial)"], "burst/trial sah"),
        )):
            chart = BarChart()
            chart.title = f"{SCENARIO_NAMES.get(scenario, scenario)}: {title}"
            chart.y_axis.title, chart.x_axis.title = unit, "Metode"
            for label in labels:
                column = headers.index(label) + 1
                chart.add_data(Reference(summary_sheet, min_col=column, max_col=column, min_row=min(indices), max_row=max(indices)), titles_from_data=False)
                from openpyxl.chart.series import SeriesLabel
                chart.series[-1].tx = SeriesLabel(v=label)
            chart.set_categories(Reference(summary_sheet, min_col=1, max_col=1, min_row=min(indices), max_row=max(indices)))
            chart.height, chart.width = 11, 23
            chart.display_blanks = "gap"
            if unit == "%":
                chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
            chart_sheet.add_chart(chart, f"{'A' if metric_index % 2 == 0 else 'N'}{5 + scenario_index * 46 + (metric_index // 2) * 23}")

    guide = workbook.create_sheet("Mulai Di Sini", 0)
    guide.sheet_view.showGridLines = False
    guide.sheet_properties.tabColor = "29956B"
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 32, 100
    guide.merge_cells("A1:B1")
    guide["A1"] = "RESQMESH - LAPORAN ANALISIS EKSPERIMEN"
    guide["A1"].font = Font(size=18, bold=True, color="1F4E78")
    guide.row_dimensions[1].height = 34
    guide.append(["Sesi", manifest.get("session_id")])
    guide.append(["Asal data", "DATA SINTETIS - bukan hasil perangkat fisik" if manifest.get("synthetic_data") is True else "Arsip eksperimen; sah/tidak sah mengikuti hasil trial dan validator."])
    guide.append(["Metode dalam manifest", ", ".join(METHOD_NAMES.get(m, m) for m in dict.fromkeys(row['method'] for row in summaries))])
    guide.append(["Jumlah attempt", len(trials)])
    guide.append(["Attempt sah / INVALID", f"{sum(r.get('valid') is True for r in trials)} / {sum(r.get('result') == 'INVALID' for r in trials)}"])
    guide.append(["Window pengamatan", str(manifest.get("observation_window_seconds", "Lihat manifest tiap trial")) + " detik" if manifest.get("observation_window_seconds") is not None else "Lihat waktu awal/akhir di Trial Metrics"])
    guide.append(["Cara membaca", "Mulai dari Ringkasan Metode, lalu Hasil Per Trial dan Grafik Ringkas. Gunakan Receivers/Log Validation untuk menelusuri bukti."])
    guide.append(["Nilai kosong", "Belum tersedia/tidak terdefinisi; bukan nol. N/A berarti parameter tidak dipakai. Angka tetap numerik, bukan teks."])
    guide.append(["Rata-rata vs total", "Ringkasan Metode menampilkan overhead rata-rata per trial. Method Scenario Summary menyimpan total dan rata-rata dengan kolom berbeda."])
    guide.append(["Trial gagal vs INVALID", "FAILED_DELIVERY yang sah tetap masuk DSR/statistik. INVALID terarsip tetapi dikecualikan dari ringkasan."])
    guide.append(["Batas klaim", "Burst bukan paket RF/energi. Coded PHY bukan otomatis S8. Graph logis bukan bukti isolasi RF atau jangkauan fisik tiga hop."])
    for scenario in dict.fromkeys(row["scenario"] for row in summaries):
        guide.append([SCENARIO_NAMES.get(scenario, scenario), SCENARIO_DETAILS.get(scenario, "Definisi sesuai manifest.")])
    guide.append(["TABEL / SHEET", "ISI DAN KONTEKS"])
    order = ["Mulai Di Sini", "Ringkasan Metode", "Hasil Per Trial", "Grafik Ringkas", "Metric Definitions", "Kamus Kolom",
             "Descriptive Statistics", "Receivers", "Log Validation", "Invalid Trials", "Method Parameters", "MPL Parameters",
             "Mechanism Diagnostics", "MPL Diagnostics", "STATUS Diagnostics", "PHY Evidence", "Participation", "Adaptive Pilot Review",
             "Trial Metrics", "Method Scenario Summary", "Charts", "Overview", "All Events"]
    for name in order[1:]:
        if name not in workbook:
            continue
        guide.append([name, SHEET_HELP.get(name, "Tabel teknis pendukung.")])
        guide.cell(guide.max_row, 1).hyperlink = f"#'{name}'!A1"
        guide.cell(guide.max_row, 1).style = "Hyperlink"
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].font = Font(bold=True, color="1F4E78")
    guide.freeze_panes = "B2"
    for index, name in enumerate(n for n in order if n in workbook):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    workbook.active = 0
