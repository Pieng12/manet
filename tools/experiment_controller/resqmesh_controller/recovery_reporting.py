"""Indonesian recovery views, separate from canonical events and old profiles."""
import csv
import json

from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from openpyxl.worksheet.filters import AutoFilter

from .excel_report import _write_table
from .neighbor_excel_layout import _style, METHOD_NAMES, SCENARIO_NAMES
from .neighbor_reporting import descriptive
from .neighbor_testbed import scenario_design, testbed

STATUS_NAMES = {
    "RECOVERED": "Berhasil menerima setelah ON",
    "NOT_RECOVERED_WITHIN_WINDOW": "Belum menerima hingga akhir window",
    "TIMING_UNVERIFIED": "Bukti waktu belum cukup",
    "NOT_APPLICABLE": "Tidak berlaku",
}


def decorate_recovery(workbook, manifest, trials, summaries, receivers, validation, curves, output):
    design = testbed(manifest)
    def readable(rows):
        return [{**r, **({"Metode": METHOD_NAMES.get(r["Metode"], r["Metode"])} if "Metode" in r else {}),
                 **({"Skenario": SCENARIO_NAMES.get(r["Skenario"], r["Skenario"])} if "Skenario" in r else {})}
                for r in rows]

    guide = workbook["Mulai Di Sini"]
    guide["A1"] = "RESQMESH - PEMULIHAN 180 DETIK " + ("ESP-ONLY (BUKAN DATA UTAMA HP)" if design.esp_only else "HP-ESP")
    guide.row_dimensions[1].height = 54
    explanations = {
        "Window pengamatan": "Semua skenario 180 detik sejak DATA pertama sumber benar-benar mulai. S1 D/E ON detik 60; S2 E ON detik 90.",
        "Cara membaca": "Baca Desain Pengujian, Ringkasan Metode dan Pemulihan Node. Grafik harus dibaca bersama jumlah penerima dan validasi.",
        SCENARIO_NAMES["S2_POST_DATA_STOP"]: "E TX/RX OFF sebelum SOS, ON detik 90. Stop DATA D, buffer dan repair harus dibuktikan; INCONCLUSIVE tidak membatalkan delivery yang sah.",
    }
    for row in guide.iter_rows(min_row=2):
        if row[0].value in explanations:
            row[1].value = explanations[row[0].value]
    guide.append(["Metrik pendukung kelima", "Waktu pemulihan = RX DATA fisik pertama setelah ON - event ON aktual pada jam monoton dan boot epoch target yang sama. Bukan durasi proses repair."])
    guide.append(["Pemulihan tidak teramati", "Delay kosong, bukan nol/180 detik. Baca target layak, menerima, belum menerima dan bukti waktu kurang. Rata-rata hanya atas nilai terdefinisi."])
    guide.append(["Expiration DATA", "3 interval selesai; nominal 8+16+32 = 56 detik tanpa reset, sejak timer masing-masing node mulai. Tidak menjamin stop sebelum E ON."])
    guide.append(["Denominator", f"{len(design.targets)} target; SOURCE {design.source} tidak dihitung sebagai penerima. Biaya mencakup seluruh {len(design.node_ids)} node."])
    guide.append(["Profil / versi pemulihan", f"{manifest['testbed_profile']} / {manifest['recovery_measurement_version']}"])

    design_rows = [{"Kategori": "Perangkat", "Identitas": node,
                    "Peran": "SOURCE" if node == design.source else "RELAY / target",
                    "Tetangga logis": ", ".join(b if a == node else a for a, b in design.edges if node in (a, b))}
                   for node in design.node_ids]
    design_rows += [{"Kategori": "Skenario", "Identitas": SCENARIO_NAMES[s],
                     "Window (detik)": p["observation_window_seconds"], "ON setelah t0 (detik)": p["activate_at_seconds"],
                     "TX/RX OFF sebelum SOS": ", ".join(p["inactive_node_ids"])} for s, p in scenario_design(manifest).items()]
    _write_table(workbook, "Desain Pengujian", design_rows)

    columns = {
        "trial_id": "ID trial", "method": "Metode", "scenario": "Skenario", "receiver": "Target",
        "valid": "Trial sah", "recovery_eligible": "Target pemulihan", "recovery_status": "Hasil pemulihan",
        "recovery_delay_ms": "Waktu pemulihan (ms)", "recovery_received": "DATA teramati setelah ON",
        "recovery_followup_ms": "Sisa window sejak ON (ms)", "recovery_followup_basis": "Dasar sisa window",
        "recovery_on_confirmed_at_ms": "ON aktual (epoch ms)", "recovery_on_monotonic_ms": "ON monoton (ms)",
        "first_rx_monotonic_ms": "RX monoton (ms)", "recovery_first_rx_at_ms": "RX setelah ON (epoch ms)",
        "recovery_clock_domain": "Domain jam", "recovery_local_boot_id": "Boot epoch lokal",
        "recovery_observation_id": "Identitas observasi RX", "recovery_time_basis": "Dasar waktu pemulihan",
        "recovery_unavailable_reason": "Alasan delay kosong",
    }
    human = [{label: STATUS_NAMES.get(r.get(key), r.get(key)) if key == "recovery_status" else r.get(key)
              for key, label in columns.items()} for r in receivers]
    _write_table(workbook, "Pemulihan Node", readable(human))
    stats = []
    for scenario in manifest["neighbor_scenarios"]:
        targets = scenario_design(manifest)[scenario]["inactive_node_ids"]
        for method in manifest["experiment_methods"]:
            for target in ["Semua target pemulihan", *targets]:
                group = [r for r in receivers if r["valid"] and r["scenario"] == scenario
                         and r["method"] == method and r.get("recovery_eligible")
                         and (target == "Semua target pemulihan" or r["receiver"] == target)]
                values = [r["recovery_delay_ms"] for r in group if r.get("recovery_status") == "RECOVERED"]
                stats.append({"Metode": method, "Skenario": scenario, "Target": target,
                              "Target layak": len(group), "Menerima": sum(r.get("recovery_received") is True for r in group),
                              "Belum menerima": sum(r.get("recovery_status") == "NOT_RECOVERED_WITHIN_WINDOW" for r in group),
                              "Bukti waktu kurang": sum(r.get("recovery_status") == "TIMING_UNVERIFIED" for r in group),
                              **{label: descriptive(values)[field] for field, label in
                                 (("defined_trials", "Nilai terdefinisi"), ("mean", "Rata-rata (ms)"),
                                  ("median", "Median (ms)"), ("sample_sd", "SD sampel (ms)"),
                                  ("min", "Minimum (ms)"), ("max", "Maksimum (ms)"))}})
    stats.sort(key=lambda r: (manifest["neighbor_scenarios"].index(r["Skenario"]),
                             r["Target"] != "Semua target pemulihan", r["Target"],
                             manifest["experiment_methods"].index(r["Metode"])))
    _write_table(workbook, "Statistik Pemulihan", readable(stats))
    _write_table(workbook, "Validasi", [{"ID trial": c["trial_id"], "Metode": METHOD_NAMES.get(c["method"], c["method"]),
                 "Skenario": SCENARIO_NAMES.get(c["scenario"], c["scenario"]), "Pemeriksaan": c["check"],
                 "status": c["result"], "Penjelasan": c["evidence"]} for c in validation])
    _write_table(workbook, "Kurva Kumulatif", curves)
    for name in ("Desain Pengujian", "Pemulihan Node", "Statistik Pemulihan", "Validasi", "Kurva Kumulatif"):
        _style(workbook[name])
        guide.append([name, "Analisis pendukung; bukti mentah dan tabel teknis tetap tersedia."])
        guide.cell(guide.max_row, 1).hyperlink = f"#'{name}'!A1"
        guide.cell(guide.max_row, 1).style = "Hyperlink"
    workbook["Pemulihan Node"].freeze_panes = "E2"
    workbook["Statistik Pemulihan"].freeze_panes = "D2"

    # Add columns at the end; never replace existing metric values or units.
    for name, data, fields in (("Ringkasan Metode", summaries,
            (("recovery_eligible_targets", "Target pemulihan layak"), ("recovery_received_targets", "Target pemulihan menerima"),
             ("recovery_unreceived_targets", "Belum menerima hingga akhir"), ("recovery_unverified_targets", "Bukti waktu kurang"),
             ("recovery_mean_ms_trial_mean", "Pemulihan rata-rata per trial (ms)"))),
            ("Hasil Per Trial", trials, (("recovery_eligible_targets", "Target pemulihan layak"),
             ("recovery_received_targets", "Target pemulihan menerima"), ("recovery_mean_ms", "Pemulihan rata-rata target (ms)")))):
        sheet = workbook[name]
        first = sheet.max_column+1
        for i, (field, label) in enumerate(fields, first):
            sheet.cell(1, i, label)
            sheet.cell(1, i).comment = Comment("Pemulihan: nilai kosong bukan nol; baca denominator dan jumlah nilai terdefinisi.", "ResQMesh")
            for r, row in enumerate(data, 2):
                sheet.cell(r, i, row.get(field))
        for table in sheet.tables.values():
            table.ref = sheet.dimensions
            table.tableColumns = []
            table._initialise_columns()
            for col, cell in zip(table.tableColumns, sheet[1]):
                col.name = str(cell.value)
        _style(sheet)
    dictionary = workbook["Kamus Kolom"]
    for key, label in columns.items():
        dictionary.append(["Pemulihan Node", key, label,
                           "Jam monoton lokal untuk delay; sisa window memakai batas wall clock terkoreksi. Nilai kosong bukan nol.",
                           "ms" if key.endswith("_ms") else "sesuai field"])
    for table in dictionary.tables.values():
        table.ref = dictionary.dimensions
        table.autoFilter = AutoFilter(ref=dictionary.dimensions)
    _style(dictionary)
    definitions = workbook["Metric Definitions"]
    definitions.append(["Pemulihan (pendukung)", "RX DATA fisik setelah ON - event ON aktual", "Jam monoton/boot epoch lokal sama; tidak menerima -> kosong; S0 tidak berlaku"])
    for table in definitions.tables.values():
        table.ref = definitions.dimensions
        table.autoFilter = AutoFilter(ref=definitions.dimensions)
    _style(definitions)

    graphs = workbook.create_sheet("Grafik Pemulihan")
    graphs.sheet_view.showGridLines = False
    graphs["A1"] = "Waktu pemulihan dan jumlah penerima - hanya trial sah; nilai kosong bukan nol"
    data = workbook["Statistik Pemulihan"]
    for index, scenario in enumerate(manifest["neighbor_scenarios"]):
        indices = [i+2 for i, r in enumerate(stats) if r["Skenario"] == scenario and r["Target"] == "Semua target pemulihan"]
        headers = [c.value for c in data[1]]
        for j, (title, fields, unit) in enumerate((
            ("Waktu pemulihan", ["Rata-rata (ms)"], "ms; hanya penerimaan dengan waktu terverifikasi"),
            ("Jumlah target", ["Target layak", "Menerima", "Belum menerima", "Bukti waktu kurang"], "target"))):
            chart = BarChart()
            chart.title = SCENARIO_NAMES[scenario] + ": " + title
            chart.y_axis.title = unit
            for field in fields:
                col = headers.index(field)+1
                chart.add_data(Reference(data, min_col=col, min_row=min(indices), max_row=max(indices)), titles_from_data=False)
                from openpyxl.chart.series import SeriesLabel
                chart.series[-1].tx = SeriesLabel(v=field)
            chart.set_categories(Reference(data, min_col=1, min_row=min(indices), max_row=max(indices)))
            chart.display_blanks = "gap"
            chart.width, chart.height = 23, 11
            graphs.add_chart(chart, f"{'A' if j == 0 else 'N'}{4+index*24}")
    cumulative = workbook.create_sheet("Grafik Kumulatif")
    cumulative.sheet_view.showGridLines = False
    cumulative["A1"] = "Kurva per trial sah; waktu sejak DATA pertama sumber; DATA + CONTROL seluruh node"
    curve_sheet = workbook["Kurva Kumulatif"]
    headers = [c.value for c in curve_sheet[1]]
    for i, scenario in enumerate(manifest["neighbor_scenarios"]):
        for j, (field, title, unit) in enumerate((("dsr_percent", "Delivery kumulatif", "%"), ("network_overhead", "Overhead kumulatif", "burst"))):
            chart = ScatterChart()
            chart.title = SCENARIO_NAMES[scenario] + ": " + title
            chart.x_axis.title, chart.y_axis.title = "Detik sejak DATA sumber pertama", unit
            chart.width, chart.height = 23, 11
            for trial in dict.fromkeys(r["trial_id"] for r in curves if r["scenario"] == scenario and r["valid"]):
                indices = [n+2 for n, r in enumerate(curves) if r["trial_id"] == trial]
                x = Reference(curve_sheet, min_col=headers.index("elapsed_seconds")+1, min_row=min(indices), max_row=max(indices))
                y = Reference(curve_sheet, min_col=headers.index(field)+1, min_row=min(indices), max_row=max(indices))
                chart.series.append(Series(y, x, title=trial))
            if unit == "%":
                chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
            cumulative.add_chart(chart, f"{'A' if j == 0 else 'N'}{4+i*24}")
    for sheet in workbook:
        for cell in sheet[1]:
            if cell.value in {"N", "dsr_percent", "data_tx", "failed_pairs", "successful_pairs"}:
                cell.comment = Comment(f"Profil {manifest['testbed_profile']}: {len(design.targets)} target, {len(design.node_ids)} node untuk biaya. SOURCE tidak masuk denominator DSR.", "ResQMesh")
    for row in dictionary.iter_rows(min_row=2):
        if row[1].value in {"N", "dsr_percent", "data_tx", "failed_pairs", "successful_pairs"}:
            row[3].value = f"Profil ini mempunyai {len(design.targets)} target; DSR=U/(M*{len(design.targets)})*100; biaya seluruh {len(design.node_ids)} node."
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    order = ["Mulai Di Sini", "Desain Pengujian", "Ringkasan Metode", "Hasil Per Trial", "Pemulihan Node",
             "Grafik Ringkas", "Grafik Pemulihan", "Grafik Kumulatif", "Validasi", "Kamus Kolom", "Statistik Pemulihan"]
    for i, name in enumerate(order):
        workbook.move_sheet(name, offset=i-workbook.sheetnames.index(name))
    workbook.move_sheet("All Events", offset=len(workbook.sheetnames)-1-workbook.sheetnames.index("All Events"))
    for name, rows in (("recovery_receivers", receivers), ("recovery_statistics", stats)):
        (output/f"{name}.json").write_text(json.dumps(rows, indent=2), encoding="utf-8")
        fields = list(dict.fromkeys(k for r in rows for k in r)) or ["trial_id"]
        with (output/f"{name}.csv").open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
