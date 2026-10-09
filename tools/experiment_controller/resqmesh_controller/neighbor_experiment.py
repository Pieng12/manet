"""Versioned all-target experiment; legacy hop profiles retain their formulas."""
from __future__ import annotations

import hashlib
import base64
import json
import math
import random
import re
import time
import zlib
from pathlib import Path
from typing import Any

from .controller import ExperimentController, TrialSpec
from .config import ConfigError, _valid_research_build_id
from .devices import DeviceError
from .radio import radio_readiness_errors
from . import mpl_config
from .neighbor_testbed import (ESP_ONLY, ESP_SCENARIOS, RECOVERY_ESP, RECOVERY_VERSION,
                              evidence_profile, recovery_profile, scenario_design,
                              scenario_parameters, testbed)

PROFILE = "neighbor_graph_v1"
VERSION = "resqmesh-neighbor-v1"
METHODS = ("basic_flooding", "trickle_no_suppression", "trickle", "trickle_neighbor_status")
SCENARIOS = ("S0_MAIN", "S1_DELAYED_RX", "S2_LATE_JOIN")
SOURCE = "android-source"
TARGETS = ("esp-r1a", "esp-r1b", "esp-r2a", "esp-r2b", "esp-destination")
EDGES = ((SOURCE, TARGETS[0]), (SOURCE, TARGETS[1]), (TARGETS[0], TARGETS[2]),
         (TARGETS[1], TARGETS[2]), (TARGETS[0], TARGETS[3]), (TARGETS[3], TARGETS[4]))
DEFAULTS = {"status_period_ms": 12000, "status_burst_ms": 500, "freshness_ms": 45000,
            "discovery_jitter_ms": 1500, "reset_cooldown_ms": 8000, "neighbor_capacity": 16}
ADAPTIVE_DEFAULTS = {**DEFAULTS, "status_period_ms": 60000, "freshness_ms": 150000,
                     "status_min_period_ms": 15000, "empty_retry_min_ms": 4000,
                     "empty_retry_max_ms": 32000, "data_grace_ms": 10000}


def neighbor_parameters(config: dict[str, Any]) -> dict[str, Any]:
    policy = config.get("neighbor_status_policy", "periodic_v1")
    if policy not in {"periodic_v1", "adaptive_v2"}:
        raise ConfigError("unknown neighbor_status_policy")
    overrides = config.get("neighbor_parameters", {})
    defaults = ADAPTIVE_DEFAULTS if policy == "adaptive_v2" else DEFAULTS
    if not isinstance(overrides, dict) or set(overrides) - set(defaults):
        raise ConfigError("unknown neighbor parameter")
    values = {**defaults, **overrides, "neighbor_status_policy": policy}
    if any(type(v) is not int or v < 0 for k, v in values.items() if k != "neighbor_status_policy"):
        raise ConfigError("neighbor parameters must be nonnegative integers")
    return values


def stable_id(node: str) -> int:
    return zlib.crc32(node.encode("utf-8")) & 0xffffffff


def adjacency(node: str) -> list[int]:
    return [stable_id(b if a == node else a) for a, b in EDGES if node in (a, b)]


def validate_neighbor_config(config: dict[str, Any]) -> None:
    errors = []
    design = testbed(config)
    nodes = config.get("nodes", [])
    ids = [n.get("node_id") for n in nodes]
    if len(ids) != len(design.node_ids) or set(ids) != set(design.node_ids):
        errors.append("required: nodes match the selected testbed")
    numeric = [stable_id(i) for i in ids if isinstance(i, str)]
    if len(set(numeric)) != len(design.node_ids) or 0 in numeric:
        errors.append("stable transmitter ID collision")
    for transport, key in (("adb", "serial"), ("serial", "port")):
        values = [str(n.get(key, "")).strip().upper() for n in nodes if n.get("transport") == transport]
        if any(not v for v in values) or len(set(values)) != len(values):
            errors.append(f"missing/duplicate {key}")
    if any(n.get("transport") != ("adb" if not design.esp_only and n.get("node_id") == design.source else "serial") for n in nodes):
        errors.append("source must be ADB, all targets serial")
    if design.esp_only and not recovery_profile(config):
        if not mpl_config.enabled(config) or tuple(config.get("modes", ())) != mpl_config.METHODS:
            errors.append("ESP-only requires the three MPL comparison methods")
        if config.get("activation_tolerance_ms") != 10000:
            errors.append("ESP-only activation tolerance is frozen at 10000 ms")
        if any(not re.fullmatch(r"COM[1-9][0-9]*", str(n.get("port", "")), re.IGNORECASE) for n in nodes):
            errors.append("ESP-only requires explicit numeric COM ports, not placeholders")
        if any(n.get("role") != ("SOURCE" if n.get("node_id") == design.source else "RELAY") for n in nodes):
            errors.append("ESP-only requires one SOURCE and four RELAY targets")
        if config.get("scenario_parameters") != ESP_SCENARIOS:
            errors.append("ESP-only scenario_parameters must match the frozen 180/180/420-second design")
        if config.get("valid_trials_per_condition", 1) != 1 or config.get("max_attempts_per_condition", 1) != 1:
            errors.append("ESP-only smoke has one attempt per condition, without replacements")
        if config.get("quiet_period_seconds") != 5:
            errors.append("ESP-only quiet period must be five seconds")
    if tuple(config.get("modes", mpl_config.methods(config))) != mpl_config.methods(config):
        errors.append("required methods in canonical order; trial plan randomizes them")
    if 'trickle_mpl' in config.get('modes',[]) and not mpl_config.enabled(config):
        errors.append("MPL scheduler semantics required")
    if mpl_config.enabled(config):
        mpl_config.parameters(config)
        if config.get('buffer_retention', 'persistent_until_supersession_ack_admin') != 'persistent_until_supersession_ack_admin':
            errors.append('MPL timers must not delete persistent protocol buffers')
        p = mpl_config.parameters(config)
        if float(config.get('observation_window_seconds',0)) * 1000 < float(config.get('perturbation_delay_seconds',30))*1000 + p['mpl_discovery_jitter_ms'] + p['mpl_control_imin_ms'] + p['mpl_data_imin_ms']:
            errors.append('MPL discovery and first repair cannot fit recovery window')
        if set(config.get('scenarios',[])) != set(scenario_design(config) if evidence_profile(config) else SCENARIOS):
            errors.append('MPL profile requires all three scenarios')
    scenarios = config.get("scenarios", ["S0_MAIN"])
    if not scenarios or len(scenarios) != len(set(scenarios)) or not set(scenarios) <= set(scenario_design(config) if evidence_profile(config) else SCENARIOS):
        errors.append("invalid scenarios")
    if recovery_profile(config):
        if not mpl_config.enabled(config) or config.get("scenario_parameters") != scenario_design(config):
            errors.append("recovery profile requires the frozen 180/60/90-second design")
        if mpl_config.parameters(config) != {**mpl_config.DEFAULTS, "mpl_data_expirations": 3}:
            errors.append("recovery profile changes DATA expirations only, from 5 to 3")
        if config.get("activation_tolerance_ms") != 10000 or config.get("quiet_period_seconds") != 5:
            errors.append("recovery activation tolerance/quiet period mismatch")
        if config.get("observation_window_seconds") != 180 or config.get("random_seed") != 231402098:
            errors.append("recovery window/seed mismatch")
        if any(n.get("role") != ("SOURCE" if n.get("node_id") == design.source else "RELAY") for n in nodes):
            errors.append("recovery requires one SOURCE and all targets RELAY")
        if any(not re.fullmatch(r"COM[1-9][0-9]*", str(n.get("port", "")), re.IGNORECASE) for n in nodes if n.get("transport") == "serial"):
            errors.append("recovery requires explicit numeric COM ports")
        if design.esp_only and (config.get("valid_trials_per_condition") != 1 or config.get("max_attempts_per_condition") != 1):
            errors.append("ESP recovery smoke has one attempt per condition")
    if config.get("gateway_enabled") or config.get("ack_enabled"):
        errors.append("gateway/completion ACK must be disabled")
    for key in (("firmware_build_id",) if design.esp_only else ("android_build_id", "firmware_build_id")):
        if not _valid_research_build_id(config.get(key)):
            errors.append(f"{key} must identify the frozen source build (12-40 hex)")
    if not design.esp_only and config.get("android_build_id") != config.get("firmware_build_id"):
        errors.append("APK and firmware build IDs differ")
    if config.get("radio_mode", "coded") != "coded":
        errors.append("pilot uses coded; S8 is not verified")
    try:
        window = float(config["observation_window_seconds"])
        delay = float(config.get("perturbation_delay_seconds", 30))
        recovery = float(config.get("minimum_recovery_seconds", 120))
        if not all(math.isfinite(v) for v in (window, delay, recovery)) or window <= 0 or delay <= 0 or recovery <= 0:
            raise ValueError()
        if not recovery_profile(config) and any(s != "S0_MAIN" for s in scenarios) and delay + recovery > window:
            errors.append("perturbation plus recovery must fit the physical t0 window")
        if int(config.get("valid_trials_per_condition", 15)) < 1:
            errors.append("replicates must be positive")
        if int(config.get("max_attempts_per_condition", 45)) < int(config.get("valid_trials_per_condition", 15)):
            errors.append("attempt budget smaller than replicate target")
        if config.get("trial_order", "balanced_randomized") != "balanced_randomized":
            errors.append("new profile requires balanced randomized blocks")
        if int(config.get("clock_tolerance_ms", -1)) < 0:
            errors.append("clock tolerance required")
        p = neighbor_parameters(config)
        if (p["status_period_ms"] < 1000 or not 250 <= p["status_burst_ms"] <= 2000 or
                p["freshness_ms"] < p["status_period_ms"] * 2 or p["discovery_jitter_ms"] < 0 or
                p["reset_cooldown_ms"] < 1000 or not 5 <= p["neighbor_capacity"] <= 64):
            errors.append("unsafe neighbor parameters")
        if p["neighbor_status_policy"] == "adaptive_v2" and (
                not 1000 <= p["status_min_period_ms"] <= p["status_period_ms"] or
                not 1000 <= p["empty_retry_min_ms"] <= p["empty_retry_max_ms"] or
                p["data_grace_ms"] < 1000 or
                p["freshness_ms"] < 2 * (p["status_period_ms"] + p["discovery_jitter_ms"] + p["status_burst_ms"])):
            errors.append("unsafe adaptive neighbor parameters")
        if any(p[k] > 0x7fffffff // 4 for k in ("status_period_ms", "discovery_jitter_ms", "empty_retry_max_ms", "data_grace_ms") if k in p) or p["freshness_ms"] > 0x7fffffff:
            errors.append("neighbor timer exceeds monotonic range")
    except (KeyError, ValueError, TypeError):
        errors.append("invalid duration/parameter")
    if errors:
        raise ConfigError("; ".join(errors))


def neighbor_fingerprint(config: dict[str, Any]) -> str:
    relevant = {k: v for k, v in config.items() if k not in {"session_id", "session_label", "valid_trials_per_condition", "trials_per_condition", "max_attempts_per_condition", "trial_order", "hypotheses"}}
    parameters = neighbor_parameters(config)
    if "neighbor_status_policy" not in config:
        parameters.pop("neighbor_status_policy")
    relevant.update(transport_version=VERSION, design_version=1, measurement_version=1, graph=testbed(config).edges,
                    neighbor_parameters=parameters)
    if mpl_config.enabled(config):
        relevant.update(scheduler_semantics=mpl_config.SEMANTICS,
                        mpl_parameters=mpl_config.parameters(config),buffer_retention='persistent_until_supersession_ack_admin')
    if recovery_profile(config):
        relevant.update(recovery_measurement_version=RECOVERY_VERSION,
                        scenario_parameters=scenario_design(config), **testbed(config).metadata())
    return hashlib.sha256(json.dumps(relevant, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def neighbor_plan(config: dict[str, Any]) -> list[TrialSpec]:
    rng = random.Random(int(config.get("random_seed", 231402095)))
    rows = []
    for block in range(1, int(config.get("valid_trials_per_condition", 15)) + 1):
        combinations = [TrialSpec(m, s, block, block) for m in mpl_config.methods(config) for s in config.get("scenarios", ["S0_MAIN"])]
        rng.shuffle(combinations)
        rows.extend(combinations)
    return rows


def flatten(event: dict[str, Any]) -> dict[str, Any]:
    detail = event.get("detail_json") or event.get("detail") or {}
    if isinstance(detail, str):
        detail = json.loads(detail)
    return {**detail, **event, "node_id": event.get("node_id", event.get("device_id", detail.get("node_id")))}


def _recovery_timing(events: list[dict[str, Any]], record: dict[str, Any], sample: dict[str, Any]) -> dict[str, Any]:
    from .log_merge import _timestamp
    result = {"recovery_on_confirmed_at_ms": None, "recovery_delay_ms": None,
              "recovery_on_monotonic_ms": None, "first_rx_monotonic_ms": None,
              "recovery_clock_domain": None, "recovery_local_boot_id": None,
              "recovery_time_basis": "unverified", "recovery_unavailable_reason": "NOT_PERTURBED"}
    node = sample["node_id"]
    requests = [p for p in record.get("participation", []) if p.get("node") == node
                and p.get("enabled") is True and p.get("response", {}).get("confirmed_enabled") is True]
    if not requests:
        return result
    boot = requests[-1]["response"].get("local_boot_id")
    on_events = [e for e in events if e.get("node_id") == node
                 and e.get("event_type") in {"RX_PARTICIPATION_CHANGED", "NODE_PARTICIPATION_CHANGED"}
                 and e.get("confirmed_enabled") is True and e.get("rx_enabled") is True
                 and (boot is None or e.get("local_boot_id") == boot)]
    if not on_events:
        return {**result, "recovery_unavailable_reason": "ON_EVENT_MISSING"}
    on = max(on_events, key=lambda e: _timestamp(e) or 0)
    result["recovery_on_confirmed_at_ms"] = _timestamp(on)
    for field, event in (("recovery_on_monotonic_ms", on), ("first_rx_monotonic_ms", sample)):
        value = event.get("monotonic_ms")
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value >= 0:
            result[field] = value
    if result["recovery_on_monotonic_ms"] is None or result["first_rx_monotonic_ms"] is None:
        return {**result, "recovery_unavailable_reason": "MONOTONIC_MISSING"}
    domain = on.get("clock_domain")
    if domain not in {"esp_boot_millis", "android_elapsed_realtime"} or sample.get("clock_domain") != domain:
        return {**result, "recovery_unavailable_reason": "CLOCK_DOMAIN_UNVERIFIED"}
    result["recovery_clock_domain"] = domain
    local_boot = on.get("local_boot_id")
    rx_boot = sample.get("local_boot_id")
    if (not isinstance(local_boot, int) or isinstance(local_boot, bool) or local_boot <= 0
            or not isinstance(rx_boot, int) or isinstance(rx_boot, bool) or rx_boot != local_boot):
        return {**result, "recovery_unavailable_reason": "CLOCK_EPOCH_UNVERIFIED"}
    result["recovery_local_boot_id"] = local_boot
    # Command responses can lag scanner activation; only the ON event is the origin.
    delay = result["first_rx_monotonic_ms"] - result["recovery_on_monotonic_ms"]
    if delay < 0:
        return {**result, "recovery_unavailable_reason": "RX_BEFORE_ON"}
    return {**result, "recovery_delay_ms": delay, "recovery_time_basis": "same_node_monotonic_event",
            "recovery_unavailable_reason": None}


def summarize_network(events: list[dict[str, Any]], record: dict[str, Any]) -> dict[str, Any]:
    design = testbed(record)
    source, targets = design.source, design.targets
    from .log_merge import _timestamp
    events = [flatten(e) for e in events]
    events = [e for e in events
              if (record.get("session_id") is None or e.get("session_id") == record["session_id"])
              and (record.get("trial_id") is None or e.get("trial_id") == record["trial_id"])
              and (record.get("scope") is None or e.get("scope") == record["scope"]
                   or (e.get("scope") is None and e.get("event_type") in
                       {"SOS_CREATED", "SOURCE_FIRST_ADVERTISE_STARTED"}))]
    start, end = record.get("observation_started_at_ms"), record.get("observation_ended_at_ms")
    if start is None or end is None:
        raise ValueError("Physical observation window required")
    source_messages = {e.get("message_key") for e in events if e.get("event_type") == "SOS_CREATED" and e.get("node_id") == source and e.get("message_key")}
    starts = {}
    for e in events:
        if e.get("event_type") == "SOURCE_FIRST_ADVERTISE_STARTED" and e.get("node_id") == source and e.get("message_key") in source_messages:
            t = _timestamp(e)
            if t is not None:
                starts[e["message_key"]] = min(t, starts.get(e["message_key"], t))
    accepted, first, first_samples, seen = [], {}, {}, set()
    data_tx, control_tx, setup_control = set(), set(), set()
    for e in events:
        t = _timestamp(e)
        if t is None:
            continue
        node = e.get("node_id", e.get("device_id"))
        if node not in design.node_ids:
            continue
        if record.get("scope") is not None and e.get("scope") != record["scope"]:
            continue
        in_window = start <= t < end
        kind = e.get("event_type")
        burst = (node, e.get("scope"), e.get("transmitter_id"), e.get("boot_id"), e.get("transmission_sequence"))
        complete_id = all(v is not None for v in burst)
        valid_tx = complete_id and e["transmitter_id"] == stable_id(node)
        if kind == "STATUS_BURST_STARTED" and valid_tx:
            if in_window:
                control_tx.add(burst)
            elif t < start:
                setup_control.add(burst)
        if kind == "DATA_BURST_STARTED" and in_window and valid_tx and e.get("message_key") in source_messages:
            data_tx.add(burst)
        if kind != "DATA_RECEIVED" or not in_window or node not in targets or not complete_id or e.get("message_key") not in source_messages:
            continue
        if e["transmitter_id"] not in design.adjacency(node) or burst in seen:
            continue
        seen.add(burst)
        accepted.append(e)
        pair = (e["message_key"], node)
        if pair not in first or t < first[pair]:
            first[pair] = t
            first_samples[pair] = e
    latency = []
    for (message, node), t in sorted(first.items()):
        source_at = starts.get(message)
        clocks = record.get("clock_samples", {})
        bounds = [clocks.get(n, {}).get("uncertainty_ms") for n in (source, node)]
        uncertainty = sum(bounds) if all(v is not None for v in bounds) else None
        latency.append({"message_key": message, "receiver": node, "first_rx_at_ms": t,
                        **_recovery_timing(events, record, first_samples[(message, node)]),
                        "e2e_latency_ms": None if source_at is None else t - source_at,
                        "clock_uncertainty_ms": uncertainty,
                        "clock_tolerance_ms": record.get("clock_tolerance_ms")})
    values = [v["e2e_latency_ms"] for v in latency if v["e2e_latency_ms"] is not None]
    m, r, u = len(source_messages), len(accepted), len(first)
    result = {"measurement_version": "all-node-burst-v1", "M": m, "N": len(targets), "U": u, "R": r,
            "dsr_percent": 100 * u / (m * len(targets)) if m else None,
            "ldr_percent": 100 * (r-u) / r if r else None,
            "e2e_mean_ms": sum(values)/len(values) if values else None,
            "successful_pairs": u, "data_tx": len(data_tx), "control_tx": len(control_tx),
            'failed_pairs':m*len(targets)-u,
            "network_overhead": len(data_tx)+len(control_tx), "setup_control_tx": len(setup_control),
            "setup_plus_window_tx": len(setup_control)+len(data_tx)+len(control_tx),
            "per_receiver": latency}
    if recovery_profile(record):
        from .recovery_metrics import recovery_metrics
        result.update(recovery_metrics(events, record))
    return result


class NeighborExperimentController(ExperimentController):
    def __init__(self, config, nodes, output_dir, sleep=time.sleep):
        self.testbed = testbed(config)
        normalized = {**config, "hypotheses": config.get("scenarios", ["S0_MAIN"]), "modes": list(mpl_config.methods(config))}
        super().__init__(normalized, nodes, output_dir, sleep)

    def _load_manifest(self):
        value = super()._load_manifest()
        value.update(transport_profile=PROFILE, transport_version=VERSION, neighbor_design_version=1,
                     measurement_version="all-node-burst-v1", topology_basis="stable_transmitter_logical_graph_not_RF_isolation",
                     target_valid_trials=len(mpl_config.methods(self.config)) * len(self.config["hypotheses"]) * self.valid_target,
                     neighbor_scenarios=list(self.config["hypotheses"]),
                     graph=self.testbed.edges, target_node_ids=list(self.testbed.targets), neighbor_parameters=neighbor_parameters(self.config),
                     neighbor_status_policy=self.config.get("neighbor_status_policy", "periodic_v1"))
        value["reporting_version"] = "neighbor-descriptive-v1"
        if evidence_profile(self.config):
            value.update(testbed_profile=self.config["testbed_profile"], **self.testbed.metadata(),
                         scenario_parameters=scenario_design(self.config),
                         serial_node_ids=[n["node_id"] for n in self.config["nodes"] if n["transport"] == "serial"],
                         activation_tolerance_ms=self.config["activation_tolerance_ms"],
                         evidence_scope="ESP-only; four targets; not Android main dataset" if self.testbed.esp_only else "HP-ESP; five targets; logical graph")
        if recovery_profile(self.config):
            value.update(recovery_measurement_version=RECOVERY_VERSION,
                         reporting_version="recovery-180-v1")
        value["phy_claim_limit"] = "LE Coded; coding aktual UNKNOWN/UNVERIFIED kecuali ada bukti on-air; bukan otomatis S=8/125 kbps"
        value["cost_basis"] = "Burst logis berhasil dimulai; setup_plus_window_tx bukan seluruh biaya siklus hidup"
        if self.config.get("pilot_baseline") is not None:
            value["pilot_baseline"] = self.config["pilot_baseline"]
        from .config import METHOD_PARAMETERS
        value["method_parameters"] = {**METHOD_PARAMETERS, "trickle_neighbor_status": {
            "scheduler": "trickle", "suppression_basis": "fresh_observed_neighbor_inventory",
            "imin_ms": 8000, "imax_ms": 256000, "imax_doublings":5,
            "k": 1, "burst_ms": 2000,"c_suppression_enabled":False,"neighbor_suppression_enabled":True,
            **value["neighbor_parameters"]}}
        if mpl_config.enabled(self.config):
            value.update(scheduler_semantics=mpl_config.SEMANTICS,
                         mpl_parameters=mpl_config.parameters(self.config),
                         buffer_retention='persistent_until_supersession_ack_admin',
                         experiment_methods=list(mpl_config.METHODS),
                         rfc_claim='BLE adaptation, not full RFC 7731 interoperability')
            value['method_parameters'] = {m:value['method_parameters'][m] for m in mpl_config.METHODS if m != 'trickle_mpl'}
            value['method_parameters']['trickle_mpl'] = {
                'scheduler':'trickle', 'suppression_enabled':True,
                'imin_ms':8000,'imax_ms':256000,'imax_doublings':5,'k':1,'burst_ms':2000,
                'suppression_basis':'c/k with bounded MISSING repair override',
                **value['mpl_parameters']}
        return value

    def _node_topology(self, node, hypothesis):
        return {"role": "SOURCE" if node.node_id == self.testbed.source else "RELAY", "active": True,
                "expected_hop_in": None, "hop_out": 1 if node.node_id == self.testbed.source else 0}

    def observation_window_seconds(self, spec):
        return float(scenario_parameters(self.config, spec.hypothesis)["observation_window_seconds"])

    def trial_metadata(self, spec):
        if not evidence_profile(self.config):
            return {}
        return {"testbed_profile": self.config["testbed_profile"], **self.testbed.metadata(),
                "recovery_measurement_version": self.manifest.get("recovery_measurement_version"),
                "scenario_parameters": scenario_parameters(self.config, spec.hypothesis),
                **({"firmware_build_id": self.config["firmware_build_id"],
                    "android_build_id": self.config.get("android_build_id")} if recovery_profile(self.config) else {})}

    def source_message_key(self, source, spec, response):
        if not self.testbed.esp_only:
            return super().source_message_key(source, spec, response)
        from .controller import canonical_message_key
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            if getattr(source, "transport_error", None):
                raise DeviceError(source.transport_error)
            events = source.diagnostic_events(self.manifest["session_id"], self._device_trial_id(spec))
            created = [e for e in events if e.get("event_type") == "SOS_CREATED"
                       and e.get("node_id") == self.testbed.source
                       and e.get("scope") == stable_id(self._device_trial_id(spec))]
            keys = {canonical_message_key(e.get("message_key")) for e in created}
            if created:
                if len(created) != 1 or len(keys) != 1 or None in keys:
                    raise DeviceError("ESP_SOURCE_SOS_IDENTITY_AMBIGUOUS")
                return next(iter(keys))
            self.sleep(.05)
        raise DeviceError("ESP_SOURCE_SOS_CREATED_MISSING")

    def readiness(self, spec=None, after_reset=False, require_active_trial=False):
        results, failures = {}, []
        for node in self.nodes:
            result = node.command("readiness", {"command_id": self._command_id("neighbor-ready", node.node_id)})
            # Never replace a measured trial offset with a later preflight clock command.
            if spec is None and node.transport == "serial" and not result.get("clock_valid") and not result.get("trial_id"):
                synced = node.command("clock_sync", {"command_id": self._command_id("neighbor-preflight-clock", node.node_id),
                                                      "wall_time_ms": time.time_ns() // 1_000_000})
                if synced.get("ok") is not True:
                    raise DeviceError(f"preflight clock failed: {node.node_id}")
                result = node.command("readiness", {"command_id": self._command_id("neighbor-ready-after-clock", node.node_id)})
            results[node.node_id] = result
            if evidence_profile(self.config):
                self.manifest.setdefault("readiness_results", {})[node.node_id] = result
            self.manifest.setdefault("radio_readiness", {})[node.node_id] = result.get("radio")
            from .neighbor_reporting import phy_evidence
            self.manifest.setdefault("phy_evidence", {})[node.node_id] = phy_evidence(result.get("radio"))
            errors = self._readiness_errors(node, result, spec, after_reset, require_active_trial)
            if errors:
                failures.append(f"{node.node_id}: {'; '.join(errors)}")
        if evidence_profile(self.config):
            self._save_manifest()
        if failures:
            raise DeviceError("readiness failed: " + " | ".join(failures))
        return results

    def recover_interrupted_trials(self, readiness):
        known = {r["device_trial_id"] for r in self.manifest["trials"].values()}
        foreign = {node: r["trial_id"] for node, r in readiness.items()
                   if r.get("trial_id") and r["trial_id"] not in known}
        if foreign:
            raise DeviceError(f"Unarchived/foreign active trial; finish and archive explicitly: {foreign}")
        super().recover_interrupted_trials(readiness)

    def synchronize_clocks(self, spec):
        samples={}
        record=self.manifest['trials'].get(spec.trial_id,{})
        attempts=record.setdefault('clock_sync_attempts',{})
        for node in self.nodes:
            node_attempts=attempts.setdefault(node.node_id,[])
            # Retry only pre-trial measurement; never relax the bound or resync live data.
            for attempt in range(1,4):
                before=time.time_ns()/1_000_000
                if hasattr(node,"host_clock_offset_ms"):
                    offset=float(node.host_clock_offset_ms())
                    after=time.time_ns()/1_000_000
                else:
                    response=node.command("clock_sync",{"command_id":self._command_id("clock",spec.trial_id,node.node_id,str(attempt)),"wall_time_ms":int(before)})
                    if response.get("ok") is not True: raise DeviceError(f"clock sync failed: {node.node_id}")
                    after=time.time_ns()/1_000_000
                    offset=(after-before)/2
                uncertainty=(after-before)/2+1
                accepted=0 <= uncertainty <= self.config['clock_tolerance_ms'] and after >= before
                node_attempts.append({'attempt':attempt,'offset_ms':offset,'uncertainty_ms':uncertainty,
                                      'before_ms':before,'after_ms':after,'accepted':accepted})
                if accepted:
                    break
            else:
                self._save_manifest()
                raise DeviceError(f"clock uncertainty exceeds tolerance: {node.node_id}: {uncertainty:.1f}ms after 3 attempts")
            self.clock_offsets[node.node_id]=offset
            samples[node.node_id]={"offset_ms":offset,"uncertainty_ms":uncertainty,"before_ms":before,"after_ms":after}
        self.manifest['clock_sync']={"valid":True,"captured_at_ms":time.time_ns()//1_000_000,"offset_ms_by_node":self.clock_offsets,"samples":samples}
        self.manifest['trials'].get(spec.trial_id,{}).update(clock_samples=samples)
        self._save_manifest()

    def _readiness_errors(self, node, result, spec, after_reset, require_active_trial):
        errors = radio_readiness_errors(result.get("radio"), "coded", spec is not None)
        if result.get("ok") is not True:
            errors.append("ok=false")
        if result.get("neighbor_design_version") != 1:
            errors.append("neighbor design version mismatch")
        expected_build = self.config["firmware_build_id" if self.testbed.esp_only else "android_build_id"]
        if result.get("build_id") != expected_build:
            errors.append(f"build mismatch (expected={expected_build}, actual={result.get('build_id')})")
        if mpl_config.enabled(self.config) and mpl_config.SEMANTICS not in result.get('supported_scheduler_semantics',[]):
            errors.append('MPL scheduler unavailable')
        if node.transport == "adb" and node.node_id == self.testbed.source and result.get("radio", {}).get("maximum_advertising_data_length", 0) < 90:
            errors.append("adapter cannot fit largest STATUS plus manufacturer header")
        if result.get("protocol_version") != "resqmesh-ble17-v1" or result.get("payload_length") != 17 or result.get("manufacturer_id") != 65535:
            errors.append("inner protocol mismatch")
        if node.transport == "adb" and (not result.get("bluetooth") or
                not all(result.get("permissions", {}).get(p) is True for p in ("scan", "advertise"))):
            errors.append("Bluetooth/permission unavailable")
        if spec is not None and not after_reset and result.get("transport_profile") != PROFILE:
            errors.append("transport profile mismatch")
        epoch = result.get("protocol_epoch", result)
        if epoch.get("valid", epoch.get("epoch_valid")) is not True or result.get("clock_valid", True) is not True:
            errors.append("invalid epoch/clock")
        if require_active_trial and result.get("trial_id") != self._device_trial_id(spec):
            errors.append("trial scope mismatch")
        if result.get("scanner") is not True and not (result.get('advertising') is True and result.get('rx_participation') is True):
            errors.append("scanner registration missing")
        if after_reset and (result.get("packet_pending") or result.get("advertising") or result.get("queue_size") or not result.get("quiet_period_complete", False)):
            errors.append("reset/quiet verification failed")
        if after_reset and result.get("trial_id"):
            errors.append("trial still active after reset")
        if spec is not None and not after_reset:
            expected = {"suppression_enabled": spec.mode in ('trickle','trickle_mpl'), "trickle_imin_ms":8000,
                        "trickle_imax_ms":256000,"trickle_k":1,"burst_duration_ms":2000,
                        "session_id":self.manifest['session_id'], "transmitter_id":stable_id(node.node_id),
                        "allowed_transmitters":self.testbed.adjacency(node.node_id),
                        "neighbor_parameters":neighbor_parameters(self.config),
                        "transport_version":VERSION, "data_frame_length":39}
            if mpl_config.enabled(self.config):
                expected['buffer_retention']='persistent_until_supersession_ack_admin'
                expected.update(scheduler_semantics=mpl_config.SEMANTICS,mpl_parameters=mpl_config.parameters(self.config))
            for key, value in expected.items():
                actual=result.get(key)
                if key=='allowed_transmitters':
                    actual=sorted(actual or []);value=sorted(value)
                if key=='neighbor_parameters' and isinstance(actual, dict):
                    actual={"neighbor_status_policy":"periodic_v1", **actual}
                if actual!=value: errors.append(f"{key} mismatch")
            if require_active_trial and result.get('scope')!=stable_id(self._device_trial_id(spec)):
                errors.append('on-air trial scope mismatch')
            if result.get("mode") != spec.mode or result.get("role") != ("SOURCE" if node.node_id == self.testbed.source else "RELAY"):
                errors.append("mode/role mismatch")
            if result.get("protocol_active") is not True or result.get("gateway_enabled", False) or result.get("ack_enabled", False):
                errors.append("participation/gateway/ACK configuration mismatch")
        return errors

    def configure(self, spec):
        from . import PROTOCOL_EPOCH_ID, PROTOCOL_EPOCH_SECONDS, PROTOCOL_VERSION
        for node in self.nodes:
            args = {"command_id": self._command_id("cfg", spec.trial_id, node.node_id),
                    "session_id": self.manifest["session_id"], "session_code": f"{spec.mode}-{spec.hypothesis}",
                    "node_id": node.node_id, "role": "SOURCE" if node.node_id == self.testbed.source else "RELAY",
                    "target_hop": 0, "hypothesis": spec.hypothesis, "topology": PROFILE,
                    "expected_hop_in": 0, "hop_out": 1 if node.node_id == self.testbed.source else 0,
                    "mode": spec.mode, "build_id": self.config["firmware_build_id" if self.testbed.esp_only else "android_build_id"],
                    "protocol_version": PROTOCOL_VERSION, "transport_profile": PROFILE,
                    "protocol_epoch_id": PROTOCOL_EPOCH_ID, "protocol_epoch_seconds": PROTOCOL_EPOCH_SECONDS,
                    "radio_mode": "coded", "main_experiment": True, "protocol_active": True,
                    "observation_window_ms": int(self.observation_window_seconds(spec)*1000),
                    "rx_burst_gap_ms": int(self.config.get("rx_burst_gap_ms", 1000)),
                    "allowed_transmitters": self.testbed.adjacency(node.node_id),
                    "clock_offset_ms": self.clock_offsets.get(node.node_id, 0),
                    "clock_tolerance_ms": self.config["clock_tolerance_ms"],
                    **neighbor_parameters(self.config)}
            if mpl_config.enabled(self.config):
                args.update(scheduler_semantics=mpl_config.SEMANTICS,**mpl_config.parameters(self.config))
            result = node.command("configure_session", args)
            if result.get("ok") is not True:
                raise DeviceError(f"configuration failed: {node.node_id}: {result}")

    def _perturb(self, spec, enabled):
        if evidence_profile(self.config):
            parameters = scenario_parameters(self.config, spec.hypothesis)
            for node_id in parameters["inactive_node_ids"]:
                self._set_participation(spec, node_id, "set_node_participation", enabled)
            return
        if spec.hypothesis == "S0_MAIN":
            return
        node_id, command = ("esp-r2b", "set_rx_participation") if spec.hypothesis == "S1_DELAYED_RX" else ("esp-destination", "set_node_participation")
        self._set_participation(spec, node_id, command, enabled)

    def _set_participation(self, spec, node_id, command, enabled):
        node = next(n for n in self.nodes if n.node_id == node_id)
        requested = time.time_ns()//1_000_000
        response = node.command(command, {"command_id": self._command_id("participation", spec.trial_id, str(enabled)), "enabled": enabled})
        self.manifest["trials"][spec.trial_id].setdefault("participation", []).append({"node": node_id, "requested_at_ms": requested, "confirmed_at_ms": time.time_ns()//1_000_000, "enabled": enabled, "response": response})
        self._save_manifest()
        if response.get("ok") is not True or response.get("confirmed_enabled") is not enabled:
            raise DeviceError("participation not confirmed")

    def before_trigger(self, spec):
        self._perturb(spec, False)

    def observe_until(self, spec, deadline):
        if evidence_profile(self.config):
            parameters = scenario_parameters(self.config, spec.hypothesis)
            if parameters["activate_at_seconds"] is not None:
                record = self.manifest["trials"][spec.trial_id]
                t0 = record["observation_started_at_ms"]
                due = t0 + parameters["activate_at_seconds"] * 1000
                self._wait_until(time.monotonic() + max(0, (due-time.time_ns()//1_000_000)/1000))
                while time.time_ns()//1_000_000 < due:
                    self.sleep(.001)
                if spec.hypothesis == "S2_POST_DATA_STOP":
                    node = next(n for n in self.nodes if n.node_id == "esp-r2b")
                    response = node.command("get_status", {"command_id": self._command_id("pre-activation", spec.trial_id)})
                    if response.get("ok") is not True or response.get("session_id") != record["session_id"] or response.get("trial_id") != record["device_trial_id"]:
                        raise DeviceError("PRE_ACTIVATION_STATUS_SCOPE_MISMATCH")
                    record["pre_activation_status"] = {"node_id": node.node_id, "response": response,
                                                       "captured_at_ms": time.time_ns()//1_000_000}
                    self._save_manifest()
                self._perturb(spec, True)
            self._wait_until(deadline)
            return
        if spec.hypothesis != "S0_MAIN":
            t0 = self.manifest["trials"][spec.trial_id]["observation_started_at_ms"]
            due = t0 + self.config.get("perturbation_delay_seconds", 30)*1000
            self._wait_until(time.monotonic()+max(0,(due-time.time_ns()//1_000_000)/1000))
            self._perturb(spec, True)
        self._wait_until(deadline)

    def _wait_until(self, deadline):
        while time.monotonic()<deadline:
            for node in self.nodes:
                if getattr(node,'transport_error',None): raise DeviceError(node.transport_error)
            self.sleep(min(.25,max(0,deadline-time.monotonic())))

    def archive_failed_trial(self, spec, record):
        from .devices import write_jsonl
        errors = []
        for node in self.nodes:
            path = self.output_dir / "raw" / spec.trial_id / f"{node.node_id}.jsonl"
            if path.exists():
                continue
            try:
                if hasattr(node, "diagnostic_events"):
                    events = node.diagnostic_events(record["session_id"], record["device_trial_id"])
                else:
                    events = node.collect_events(record["session_id"], record["device_trial_id"])
                write_jsonl(path, [{**e, "device_trial_id": record["device_trial_id"],
                                   "trial_id": spec.trial_id,
                                   "native_clock_offset_ms": e.get("clock_offset_ms"),
                                   "clock_offset_ms": self.clock_offsets.get(node.node_id, 0)} for e in events])
            except Exception as error:
                errors.append(f"{node.node_id}: {error}")
        record["failure_archive_errors"] = errors

    def _evaluate_events(self, spec, events, source_node, message_key, record):
        SOURCE, TARGETS = self.testbed.source, self.testbed.targets
        record.update(scope=stable_id(self._device_trial_id(spec)), scenario=spec.hypothesis,
                      target_node_ids=list(TARGETS), measurement_version="all-node-burst-v1")
        metrics = summarize_network(events, record)
        invalid = [str(e.get("reason", "CONFIG_VIOLATION")) for e in events if e.get("event_type") == "EXPERIMENT_CONFIG_VIOLATION"]
        if metrics["M"] != 1:
            invalid.append("REQUIRES_ONE_UNIQUE_SOS")
        if {e.get("message_key") for e in events if e.get("event_type")=="SOS_CREATED" and e.get("node_id")==SOURCE} != {message_key}:
            invalid.append("SOURCE_MESSAGE_KEY_MISMATCH")
        if any(p["e2e_latency_ms"] is None or p["e2e_latency_ms"] < 0 for p in metrics["per_receiver"]):
            invalid.append("RX_CLOCK_INVALID")
        if not any(e.get("event_type") == "SOURCE_FIRST_ADVERTISE_STARTED" for e in events):
            invalid.append("SOURCE_START_MISSING")
        from .log_merge import _timestamp
        for e in map(flatten, events):
            if e.get("event_type") in {"DATA_RECEIVED", "DATA_BURST_STARTED", "STATUS_BURST_STARTED"}:
                if _timestamp(e) is None or any(e.get(k) is None for k in ("scope", "transmitter_id", "boot_id", "transmission_sequence")):
                    invalid.append("TRANSPORT_OR_NATIVE_CLOCK_EVIDENCE_MISSING")
        for node in (self.testbed.node_ids if self.testbed.esp_only else TARGETS):
            stream = [e.get("event_sequence") for e in events if e.get("node_id") == node]
            if not stream or any(not isinstance(s, int) for s in stream) or sorted(set(stream)) != list(range(min(stream), max(stream)+1)) or len(stream) != len(set(stream)):
                invalid.append(f"EVENT_SEQUENCE_INCOMPLETE:{node}")
        if metrics["data_tx"] < 1:
            invalid.append("SUCCESSFUL_SOURCE_DATA_START_MISSING")
        if not any(e.get("event_type")=="DATA_BURST_STARTED" and e.get("node_id")==SOURCE and e.get("message_key")==message_key for e in events):
            invalid.append("SOURCE_DATA_START_MISSING")
        if evidence_profile(self.config):
            from .esp_only_validation import scenario_checks
            record["scenario_checks"] = scenario_checks(events, record, self.manifest)
            if any(c["check"] in {"ESP_SCENARIO_PARTICIPATION", "RECOVERY_SCENARIO_PARTICIPATION"} and c["result"] != "PASS" for c in record["scenario_checks"]):
                invalid.append("ESP_SCENARIO_PARTICIPATION_UNVERIFIED")
        return ("INVALID" if invalid else "SUCCESS" if metrics["U"] == len(TARGETS) else "FAILED_DELIVERY"), invalid, metrics

    def after_trial_cleanup(self, spec, record):
        if not record.get("terminal") or record.get("result") not in {"SUCCESS", "FAILED_DELIVERY", "INVALID"}:
            return
        source = next(n for n in self.nodes if n.node_id == self.testbed.source)
        if source.transport != "adb":
            return
        final_result = record["result"]
        finalized = record.setdefault("android_finalized_results", {})
        record["ui_trial_result_confirmed"] = finalized.get(source.node_id) == final_result
        # Cleanup can invalidate a result already finalized/exported on the phone.
        if not record["ui_trial_result_confirmed"]:
            try:
                response = source.command("finalize_trial", {
                    "command_id": self._command_id("finalize-final", spec.trial_id, final_result),
                    "trial_id": record["device_trial_id"], "result": final_result,
                    "reason": ",".join(record.get("invalid_reasons", []))})
                if response.get("ok") is not True:
                    raise DeviceError(f"finalize final result failed: {response}")
                finalized[source.node_id] = final_result
                record["ui_trial_result_confirmed"] = True
            except DeviceError as error:
                record["ui_trial_result_error"] = str(error)
            if record["ui_trial_result_confirmed"]:
                try:
                    response = source.command("export_trial", {
                        "command_id": self._command_id("export-final", spec.trial_id, final_result),
                        "session_id": record["session_id"], "trial_id": record["device_trial_id"]})
                    if response.get("ok") is not True:
                        raise DeviceError(f"export final result failed: {response}")
                    record.setdefault("exports", {})[source.node_id] = response
                except DeviceError as error:
                    record["ui_trial_export_error"] = str(error)
        summary = {"measurement_version": "all-node-burst-v1", "N": len(self.testbed.targets),
                   **record.get("evidence", {}), "session_id": record["session_id"],
                   "trial_id": record["device_trial_id"], "method": spec.mode, "scenario": spec.hypothesis,
                   "result": final_result, "invalid_reasons": record.get("invalid_reasons", []),
                   "reset_verified": record.get("reset_verified") is True}
        try:
            response = source.command("store_neighbor_metrics", {
                "command_id": self._command_id("network-summary-final", spec.trial_id, final_result),
                "summary_base64": base64.b64encode(json.dumps(summary, separators=(",", ":")).encode()).decode()})
            record["ui_summary_confirmed"] = response.get("ok") is True
            if not record["ui_summary_confirmed"]:
                record["ui_summary_error"] = str(response)
        except DeviceError as error:
            record["ui_summary_confirmed"] = False
            record["ui_summary_error"] = str(error)
        self._save_manifest()

    def run_trial(self, spec):
        result = super().run_trial(spec)
        if evidence_profile(self.config) and result.get("reset_verified") is not True:
            raise DeviceError("ESP_ONLY_CLEANUP_UNCONFIRMED: stop batch; power off affected ESP before retry")
        return result

    def smoke_report(self):
        rows = [{"mode": m, "hypothesis": s, "result": next((r["result"] for r in self.manifest["trials"].values() if r["mode"] == m and r["hypothesis"] == s), "MISSING")}
                for m in mpl_config.methods(self.config) for s in self.config["hypotheses"]]
        report = {"passed": all(r["result"] in {"SUCCESS", "FAILED_DELIVERY"} for r in rows),
                  "config_fingerprint": self.config_fingerprint, "session_id": self.manifest["session_id"], "conditions": rows}
        if evidence_profile(self.config):
            from .neighbor_validation import validate_logs
            from .log_merge import read_json_events
            checks = validate_logs(read_json_events((self.output_dir / "raw").rglob("*.jsonl")), self.manifest)
            modern = recovery_profile(self.config)
            core_names = {"LOG_COMPLETENESS", "METRICS_RECOMPUTE_MATCH",
                          "RECOVERY_SCENARIO_PARTICIPATION" if modern else "ESP_SCENARIO_PARTICIPATION"}
            if modern:
                core_names.add("RECOVERY_METRICS_RECOMPUTE_MATCH")
            mpl_core = {"MPL_LISTEN_ONLY_AND_BOUNDS", "MPL_C_K_DECISIONS", "MPL_PARAMETERS_MATCH",
                        "MPL_OPPORTUNITY_ACCOUNTING", "MPL_EXPIRATION_AND_DOUBLING"}
            required = {(trial_id, name) for trial_id, r in self.manifest["trials"].items()
                        for name in core_names | (mpl_core if r["mode"] == "trickle_mpl" else set())}
            core_passed = len(required) == len(rows)*len(core_names)+len(self.config["hypotheses"])*len(mpl_core) and all(
                len(matches := [c for c in checks if (c["trial_id"], c["check"]) == pair]) == 1
                and matches[0]["result"] == "PASS" for pair in required)
            mechanism_trials = {trial_id for trial_id, r in self.manifest["trials"].items()
                                if r["mode"] == "trickle_mpl" and r["hypothesis"] == "S2_POST_DATA_STOP"}
            mechanism = [c for c in checks if c["check"] == ("RECOVERY_POST_STOP_REPAIR" if modern else "ESP_POST_STOP_REPAIR") and c["trial_id"] in mechanism_trials]
            mechanism_result = "FAIL" if any(c["result"] == "FAIL" for c in mechanism) else "PASS" if len(mechanism) == 1 and mechanism[0]["result"] == "PASS" else "INCONCLUSIVE"
            report.update(testbed_profile=self.config["testbed_profile"], batch_complete=all(r["result"] in {"SUCCESS", "FAILED_DELIVERY", "INVALID"} for r in rows),
                          delivery_passed=all(r["result"] == "SUCCESS" for r in rows),
                          core_checks_passed=core_passed, mechanism_result=mechanism_result,
                          checks=checks, target_count=len(self.testbed.targets))
            delivery_valid = all(r["result"] in {"SUCCESS", "FAILED_DELIVERY"} for r in rows)
            control_proved = self.testbed.esp_only or any(c["check"] == "ANDROID_CONTROL_RX" and c["result"] == "PASS" for c in checks)
            report["passed"] = (delivery_valid if modern else report["delivery_passed"]) and core_passed and mechanism_result == "PASS" and not any(c["result"] == "FAIL" for c in checks)
            if modern:
                report.update(android_control_rx_result="NOT_APPLICABLE" if self.testbed.esp_only else "PASS" if control_proved else "INCONCLUSIVE",
                              android_build_id=self.config.get("android_build_id"), firmware_build_id=self.config["firmware_build_id"])
                report["passed"] = report["passed"] and control_proved
        (self.output_dir/"smoke_report.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
        import csv
        with (self.output_dir/"smoke_report.csv").open("w",newline="",encoding="utf-8-sig") as f:
            writer=csv.DictWriter(f,fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
        return report


def merge_neighbor(input_dir: Path, output_dir: Path, manifest: dict[str, Any]):
    from .log_merge import read_json_events
    from .excel_report import _write_table
    from openpyxl import Workbook
    from .neighbor_reporting import aggregate, mechanism_counts, phy_evidence, write_charts
    from .neighbor_validation import validate_logs
    raw_events = read_json_events(input_dir.rglob("*.jsonl"))
    design = testbed(manifest)
    TARGETS = design.targets
    if recovery_profile(manifest):
        if (manifest.get("recovery_measurement_version") != RECOVERY_VERSION
                or manifest.get("scenario_parameters") != scenario_design(manifest)
                or manifest.get("mpl_parameters") != {**mpl_config.DEFAULTS, "mpl_data_expirations": 3}):
            raise ConfigError("Recovery manifest version/schedule/parameters mismatch")
    if evidence_profile(manifest):
        for record in manifest["trials"].values():
            if record.get("testbed_profile") != manifest["testbed_profile"] or record.get("config_fingerprint") != manifest.get("config_fingerprint"):
                raise ConfigError("Mixed ESP-only profile/fingerprint in trial archive")
            if record.get("source_node_id") != design.source or record.get("target_node_ids") != list(design.targets):
                raise ConfigError("Mixed source/targets in trial archive")
            testbed(record)
            if recovery_profile(manifest) and any(record.get(k) != manifest.get(k) for k in
                    ("recovery_measurement_version", "firmware_build_id", "android_build_id")):
                raise ConfigError("Mixed recovery version/build in trial archive")
        if any(e.get("session_id") != manifest["session_id"] or e.get("node_id") not in design.node_ids for e in raw_events):
            raise ConfigError("Foreign session/node in ESP-only archive")
        if recovery_profile(manifest) and any(
                e.get("trial_id") not in manifest["trials"]
                or e.get("config_fingerprint", manifest["config_fingerprint"]) != manifest["config_fingerprint"]
                or e.get("build_id", manifest.get("firmware_build_id")) != manifest.get("firmware_build_id")
                for e in raw_events):
            raise ConfigError("Foreign trial/fingerprint/build in recovery archive")
    events = [flatten(e) for e in raw_events if e.get("session_id") == manifest["session_id"] and e.get("trial_id") in manifest["trials"]]
    rows, receivers = [], []
    for trial_id, record in manifest["trials"].items():
        scoped = [e for e in events if e.get("trial_id") == trial_id]
        metrics = summarize_network(scoped, {**record, "testbed_profile": manifest.get("testbed_profile"), "trial_id": trial_id, "session_id": manifest["session_id"], "scope": stable_id(record["device_trial_id"])}) if all(record.get(k) is not None for k in ("observation_started_at_ms", "observation_ended_at_ms")) else {}
        rows.append({"trial_id": trial_id, "method": record["mode"], "scenario": record["hypothesis"], "analysis_group": "utama" if record["hypothesis"] in {"S0_MAIN", "S0_STABLE"} else "pendukung", "result": record["result"], "valid": record["result"] in {"SUCCESS","FAILED_DELIVERY"}, "invalid_reasons": record.get("invalid_reasons",[]), "block": record.get("block"), "observation_started_at_ms": record.get("observation_started_at_ms"), "observation_ended_at_ms": record.get("observation_ended_at_ms"), "config_fingerprint": record.get("config_fingerprint"), **{k:v for k,v in metrics.items() if k not in {"per_receiver", "recovery_receivers"}}})
        for target in TARGETS:
            pair = next((p for p in metrics.get("per_receiver", []) if p["receiver"] == target), {})
            receivers.append({"trial_id": trial_id, "method": record["mode"], "scenario": record["hypothesis"], "valid": rows[-1]["valid"], "receiver": target, "received": bool(pair), "first_rx_at_ms":pair.get("first_rx_at_ms"), "e2e_latency_ms": pair.get("e2e_latency_ms"), "clock_uncertainty_ms": pair.get("clock_uncertainty_ms"), "clock_tolerance_ms": record.get("clock_tolerance_ms"),
                                **{key: pair.get(key) for key in ('recovery_on_confirmed_at_ms', 'recovery_delay_ms',
                                    'recovery_on_monotonic_ms', 'first_rx_monotonic_ms', 'recovery_clock_domain',
                                    'recovery_local_boot_id', 'recovery_time_basis', 'recovery_unavailable_reason')}})
            if recovery_profile(manifest):
                recovered = next((r for r in metrics.get("recovery_receivers", []) if r["receiver"] == target), {})
                receivers[-1].update(recovered)
    scenarios = manifest.get("neighbor_scenarios", sorted({r["scenario"] for r in rows}))
    policy = manifest.get("neighbor_status_policy", "periodic_v1")
    for row in rows:
        row["neighbor_status_policy"] = policy if row["method"] == "trickle_neighbor_status" else "N/A"
    methods = tuple(manifest.get('experiment_methods', METHODS))
    summaries, stats = aggregate(rows, receivers, scenarios, methods, len(TARGETS))
    diagnostics = mechanism_counts(events, manifest["trials"])
    phy = [{"node_id": node, **phy_evidence(radio)} for node, radio in manifest.get("radio_readiness", {}).items()]
    validation = validate_logs(events, manifest)
    output_dir.mkdir(parents=True,exist_ok=True)
    wb=Workbook(); wb.remove(wb.active)
    _write_table(wb,"Overview",[{"field":k,"value":v} for k,v in manifest.items() if k!="trials"])
    _write_table(wb,"Metric Definitions",[
        {"metrik":"DSR","rumus":f"U/(M*{len(TARGETS)})*100","catatan":"FAILED_DELIVERY tetap masuk denominator"},
        {"metrik":"E2E","rumus":"RX native pertama - DATA sumber pertama","catatan":"Rata-rata pasangan sukses; gagal kosong"},
        {"metrik":"LDR","rumus":"(R-U)/R*100","catatan":"STATUS dikecualikan; R=0 kosong"},
        {"metrik":"Overhead","rumus":"DATA_TX + CONTROL_TX","catatan":f"Burst logis berhasil dimulai seluruh {len(design.node_ids)} node, hanya dalam window; bukan energi/paket RF"},
        {"metrik":"Setup + window","rumus":"setup_control_tx + network_overhead","catatan":"STATUS persiapan trial yang sama + window; bukan seluruh biaya siklus hidup"},
        {"metrik":"Statistik per trial","rumus":"Mean, median, SD sampel, min, max atas trial valid","catatan":"FAILED_DELIVERY termasuk; INVALID terpisah; nilai tidak terdefinisi kosong; SD kosong bila n terdefinisi < 2"},
        {"metrik":"Delay per trial vs pasangan","rumus":"e2e_mean_ms_trial_mean vs e2e_mean_ms","catatan":"Rata-rata trial tanpa bobot vs rata-rata semua pasangan sukses; bukan nilai nol untuk kegagalan"}])
    _write_table(wb,"Trial Metrics",rows, ('trial_id','method','scenario','block','result','valid',
                 'dsr_percent','e2e_mean_ms','ldr_percent','network_overhead','data_tx','control_tx',
                 'setup_control_tx','setup_plus_window_tx','M','N','R','U','failed_pairs','invalid_reasons'))
    _write_table(wb,"Method Scenario Summary",summaries, ('method','scenario','analysis_group',
                 'valid_trials','invalid_trials','failed_delivery_trials','dsr_percent','e2e_mean_ms',
                 'ldr_percent','network_overhead_trial_mean','data_tx_trial_mean','control_tx_trial_mean',
                 'setup_control_tx_trial_mean','setup_plus_window_tx_trial_mean','successful_pairs','failed_pairs'))
    pilot_review = []
    baseline = manifest.get("pilot_baseline", {})
    if policy == "adaptive_v2" and baseline:
        for scenario in scenarios:
            neighbor = next(s for s in summaries if s['method']=='trickle_neighbor_status' and s['scenario']==scenario)
            trickle = next(s for s in summaries if s['method']=='trickle' and s['scenario']==scenario)
            old = baseline.get('network_overhead_by_scenario', {}).get(scenario)
            mean = neighbor['network_overhead_trial_mean']
            reduction = 100 * (1-mean/old) if mean is not None and isinstance(old, (int,float)) and old>0 else None
            target = baseline.get('s0_reduction_target_percent',50) if scenario=='S0_MAIN' else None
            pilot_review.append({'scenario':scenario,'policy':policy,'baseline_session_id':baseline.get('session_id'),
                'baseline_trials':baseline.get('trials_per_scenario'),'baseline_overhead':old,
                'neighbor_valid_trials':neighbor['valid_trials'],'neighbor_overhead_mean':mean,
                'trickle_overhead_mean':trickle['network_overhead_trial_mean'],'reduction_percent':reduction,
                'exploration_target_percent':target,'exploration_target_met':reduction>=target if reduction is not None and target is not None else None,
                'dsr_percent':neighbor['dsr_percent'],'e2e_mean_ms':neighbor['e2e_mean_ms'],
                'catatan':'Diagnostik terhadap satu trial lama per skenario; bukan bukti statistik atau aturan INVALID. Periksa DSR, delay dan pemulihan S1/S2.'})
        _write_table(wb,'Adaptive Pilot Review',pilot_review)
    _write_table(wb,"Descriptive Statistics",stats, ('method','scenario','metric','label','unit',
                 'valid_trials','invalid_trials','defined_trials','mean','median','sample_sd','min','max'))
    _write_table(wb,"Mechanism Diagnostics",diagnostics, ('trial_id','method','scenario','basis',
                 'mpl_allowed','mpl_data_suppressed','mpl_control_suppressed','mpl_missed',
                 'mpl_repair_reset','mpl_repair_completed','mpl_timer_stopped','mpl_discovery'))
    reasons = []
    for trial_id in manifest["trials"]:
        scoped = [e for e in events if e.get("trial_id") == trial_id and e.get("event_type") in {"STATUS_BURST_REQUESTED", "STATUS_BURST_STARTED", "STATUS_BURST_FAILED", "STATUS_COALESCED_WITH_DATA"}]
        scoped = list({json.dumps(e,sort_keys=True,separators=(",", ":")):e for e in scoped}.values())
        from collections import Counter
        counts = Counter((e["node_id"], e["event_type"], e.get("reason", e.get("status_reason", "UNSPECIFIED"))) for e in scoped)
        reasons.extend({"trial_id":trial_id,"node_id":node,"event_type":kind,"reason":reason,"count":count,"neighbor_status_policy":policy} for (node,kind,reason),count in sorted(counts.items()))
    _write_table(wb,"STATUS Diagnostics",reasons)
    _write_table(wb,"PHY Evidence",phy)
    _write_table(wb,"Log Validation",validation)
    _write_table(wb,"Receivers",receivers, ('trial_id','method','scenario','receiver','valid','received',
                 'e2e_latency_ms','first_rx_at_ms','recovery_on_confirmed_at_ms','recovery_delay_ms',
                 'recovery_on_monotonic_ms','first_rx_monotonic_ms','recovery_clock_domain',
                 'recovery_local_boot_id','recovery_time_basis','recovery_unavailable_reason',
                 'clock_uncertainty_ms','clock_tolerance_ms'))
    _write_table(wb,"All Events",events, ('session_id','trial_id','node_id','event_sequence',
                 'event_type','timestamp_ms','monotonic_ms','message_key','state_identity','scope',
                 'transmitter_id','boot_id','transmission_sequence'))
    _write_table(wb,"Participation",[{"trial_id":trial_id,**p} for trial_id, record in manifest["trials"].items() for p in record.get("participation",[])])
    from .excel_report import _method_parameter_rows
    parameters = _method_parameter_rows(manifest)
    for row in parameters:
        if row['mode'] == 'trickle_neighbor_status':
            for key in ('basic_wait_ms','jitter_min_ms','jitter_max_ms','k'):
                row[key] = 'N/A'
            row['suppression_enabled'] = True
            row['parameter_notes'] = 'Keputusan DATA berdasarkan HAVE/MISSING/UNKNOWN tetangga teramati; bukan c/k'
            row['transmission_timing'] = 't acak dalam [I/2,I); STATUS terpisah memakai discovery jitter'
            row.update(manifest.get("neighbor_parameters", {}))
            row['neighbor_status_policy'] = policy
            row['status_notes'] = 'Discovery 1 detik + jitter; DATA sukses menggantikan pengumuman state yang sama; retry inventory kosong tidak berhenti' if policy == 'adaptive_v2' else 'STATUS periodik; kebijakan lama'
    _write_table(wb,"Method Parameters",parameters, ('mode','scheduler','suppression_enabled',
                 'imin_ms','imax_ms','imax_doublings','k','burst_ms','basic_wait_ms',
                 'jitter_min_ms','jitter_max_ms','transmission_timing','parameter_notes','termination'))
    _write_table(wb,"Invalid Trials",[r for r in rows if r["result"]=="INVALID"])
    mpl_events=[]
    if mpl_config.enabled(manifest):
        from .neighbor_excel_layout import mpl_parameter_rows
        _write_table(wb,'MPL Parameters',mpl_parameter_rows(manifest['mpl_parameters']),
                     ('parameter','value','unit','penjelasan'))
        mpl_events=[e for e in events if str(e.get('event_type','')).startswith('MPL_')]
        _write_table(wb,'MPL Diagnostics',mpl_events)
        (output_dir/'mpl_diagnostics.json').write_text(json.dumps(mpl_events,indent=2),encoding='utf-8')
    chart_paths = write_charts(output_dir/"charts", summaries, stats, scenarios, methods, manifest.get("synthetic_data") is True)
    _write_table(wb,"Charts",[{"artifact":Path(p).name,"catatan":"Buka SVG; grafik delay menyertakan jumlah pasangan sukses dan DSR agregat","path":str(Path("charts")/Path(p).name)} for p in chart_paths])
    for row in range(2, len(chart_paths)+2):
        wb["Charts"].cell(row, 3).hyperlink = str(Path("charts")/Path(chart_paths[row-2]).name)
        wb["Charts"].cell(row, 3).style = "Hyperlink"
    from .neighbor_excel_layout import decorate_neighbor_workbook
    decorate_neighbor_workbook(wb, manifest, rows, summaries)
    curves = []
    if recovery_profile(manifest):
        from .esp_only_reporting import cumulative_rows
        from .recovery_reporting import decorate_recovery
        curves = cumulative_rows(events, manifest)
        decorate_recovery(wb, manifest, rows, summaries, receivers, validation, curves, output_dir)
    elif design.esp_only:
        from .esp_only_reporting import decorate_esp_only, cumulative_rows
        curves = cumulative_rows(events, manifest)
        decorate_esp_only(wb, manifest, validation, curves)
    path=output_dir/("resqmesh_esp_only_analysis.xlsx" if design.esp_only else "resqmesh_neighbor_analysis.xlsx"); wb.save(path)
    (output_dir/"network_metrics.json").write_text(json.dumps(rows,indent=2),encoding="utf-8")
    (output_dir/"all_events.json").write_text(json.dumps(raw_events,indent=2),encoding="utf-8")
    (output_dir/"method_scenario_summary.json").write_text(json.dumps(summaries,indent=2),encoding="utf-8")
    (output_dir/"manifest_snapshot.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
    for name, data in (("descriptive_statistics", stats), ("mechanism_diagnostics", diagnostics), ("status_diagnostics", reasons), ("adaptive_pilot_review",pilot_review), ("phy_evidence", phy), ("log_validation", validation)) + ((("cumulative_metrics", curves),) if evidence_profile(manifest) else ()):
        (output_dir/f"{name}.json").write_text(json.dumps(data,indent=2),encoding="utf-8")
    import csv
    for name, data in (("trial_metrics",rows),("receivers",receivers),("all_events",events),("method_scenario_summary",summaries), ("descriptive_statistics",stats), ("mechanism_diagnostics",diagnostics), ("status_diagnostics",reasons), ("adaptive_pilot_review",pilot_review), ("phy_evidence",phy), ("log_validation",validation), ('mpl_diagnostics',mpl_events)) + ((("cumulative_metrics", curves),) if evidence_profile(manifest) else ()):
        with (output_dir/f"{name}.csv").open("w",newline="",encoding="utf-8-sig") as f:
            fields=list(dict.fromkeys(k for row in data for k in row)) or ["trial_id"]
            writer=csv.DictWriter(f,fieldnames=fields); writer.writeheader()
            writer.writerows({k: json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in row.items()} for row in data)
    return {"workbook":str(path),"trial_count":len(rows),"event_count":len(events), "charts":chart_paths, "validation":"log_validation.json", "synthetic_data":manifest.get("synthetic_data", False)}
