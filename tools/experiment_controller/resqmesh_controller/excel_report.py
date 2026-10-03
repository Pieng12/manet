from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from .comparisons import descriptive_comparisons


WORKBOOK_NAME = "resqmesh_analysis.xlsx"
EXCEL_MAX_DATA_ROWS = 1_048_575

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SUCCESS_FILL = PatternFill("solid", fgColor="C6EFCE")
_FAILURE_FILL = PatternFill("solid", fgColor="FFC7CE")

METRIC_DEFINITIONS = [
    {
        "metric": "Delivery Success Ratio",
        "abbreviation": "DSR",
        "unit": "rasio (0-1)",
        "level": "algoritma + hop",
        "formula": "SUCCESS / (SUCCESS + FAILED_DELIVERY)",
        "interpretation": "Semakin tinggi semakin baik",
        "validity_rule": "Trial INVALID tidak disertakan",
    },
    {
        "metric": "End-to-End Latency",
        "abbreviation": "E2E",
        "unit": "milidetik",
        "level": "trial dan algoritma + hop",
        "formula": "DESTINATION_FIRST_VALID_RECEIVE - SOURCE_FIRST_ADVERTISE_STARTED",
        "interpretation": "Semakin rendah semakin baik",
        "validity_rule": "Memerlukan jam tersinkronisasi dan identitas pesan/hop yang cocok; waktu sebelum transmisi pertama dilaporkan terpisah",
    },
    {
        "metric": "Logical Duplicate Ratio",
        "abbreviation": "LDR",
        "unit": "rasio (0-1)",
        "level": "trial dan algoritma + hop",
        "formula": "duplicates / (accepted + duplicates)",
        "interpretation": "Semakin rendah semakin baik",
        "validity_rule": "Kosong jika accepted + duplicates = 0",
    },
    {
        "metric": "Network-Wide Transmission Overhead",
        "abbreviation": "Overhead",
        "unit": "burst SOS yang berhasil dimulai / trial valid",
        "level": "algoritma + hop",
        "formula": "jumlah burst SOS unik ADVERTISE_BURST_STARTED / jumlah trial valid",
        "interpretation": "Semakin rendah semakin baik jika keberhasilan pengiriman setara",
        "validity_rule": "Menghitung burst yang berhasil dimulai, bukan permintaan advertising; burst terpotong tetap dihitung sekali",
    },
]

LATENCY_DIAGNOSTIC_FIELDS = (
    "trial_id",
    "algorithm",
    "hypothesis",
    "result",
    "e2e_latency_ms",
    "sos_created_at_ms",
    "source_wait_before_first_advertise_ms",
    "sos_creation_to_destination_ms",
    "missed_opportunities_total",
    "source_bursts_before_destination",
    "delivery_after_last_source_burst_ms",
    "source_attempt_index_to_hop1_accept",
    "source_first_attempt_success",
    "source_retry_wait_ms",
    "hop1_first_accept_elapsed_ms",
    "hop1_accept_after_source_burst_ms",
    "hop1_segment_progress_ms",
    "hop1_upstream_bursts_before_accept",
    "hop1_attribution",
    "hop2_first_accept_elapsed_ms",
    "hop2_segment_progress_ms",
    "hop2_upstream_bursts_before_accept",
    "hop2_attribution",
    "hop3_first_accept_elapsed_ms",
    "hop3_segment_progress_ms",
    "hop3_upstream_bursts_before_accept",
    "hop3_attribution",
)

TRIAL_METRIC_FIELDS = (
    "trial_id",
    "block",
    "algorithm",
    "mode",
    "hypothesis",
    "result",
    "valid",
    "delivery_success",
    "e2e_latency_ms",
    "events_total",
    "events_in_window",
    "events_outside_window",
    "metric_events_outside_window",
    "accepted",
    "duplicates",
    "ldr",
    "transmission_bursts",
    "transmission_opportunities",
    "transmission_allowed",
    "transmission_suppressed",
    "transmission_missed",
    "suppression_rate",
    "other_cancellations",
    "invalid_reasons",
)

DIAGNOSTIC_DEFINITIONS = [
    {"field": "missed_opportunities_total", "unit": "kesempatan",
     "definition": "Jumlah TRICKLE_TX_MISSED dalam seluruh log trial, termasuk sebelum transmisi pertama/t0. Scheduler terlambat melewati akhir interval; bukan suppression atau burst sukses",
     "attribution": "transmission_missed hanya dalam jendela; keduanya bukan penyebut suppression_rate. Tanpa event ini, nol bukan bukti bahwa log lama tidak kehilangan kesempatan"},
    {"field": "method_parameter_applicability", "unit": "kategori",
     "definition": "N/A berarti tidak dipakai scheduler SOS tersebut; kosong berarti tidak tersedia. Method Parameters menunjukkan parameter yang berlaku; Manifest Metadata mempertahankan konfigurasi asli",
     "attribution": "Hanya pelaporan; tidak mengubah manifest atau fingerprint penelitian"},
    {"field": "delivery_success", "unit": "hasil biner",
     "definition": "1 untuk SUCCESS valid, 0 untuk FAILED_DELIVERY valid, kosong untuk INVALID atau percobaan belum selesai",
     "attribution": "Percobaan invalid bukan kegagalan pengiriman dan tidak masuk DSR"},
    {"field": "transmission_allowed", "unit": "keputusan Trickle",
     "definition": "Keputusan allowed dalam jendela pengamatan, bukan awal advertising yang berhasil. Keputusan pertama sumber bisa mendahului t0, sementara burst suksesnya dimulai pada t0",
     "attribution": "Tidak harus sama dengan actual_bursts; Basic tidak memakai keputusan Trickle"},
    {"field": "suppression_rate", "unit": "rasio (0-1)",
     "definition": "Kesempatan Trickle yang mengalami suppression / TRICKLE_TX_OPPORTUNITY dalam jendela yang sama; kosong jika tidak ada kesempatan",
     "attribution": "Agregat memakai rasio jumlah, bukan rata-rata rasio trial; Basic tidak memiliki kesempatan Trickle"},
    {"field": "ldr_mean_per_trial", "unit": "rasio (0-1)",
     "definition": "Rata-rata tanpa pembobotan dari LDR trial valid yang terdefinisi; ldr utama memakai total duplicates / total accepted+duplicates",
     "attribution": "Ringkasan pendukung; tidak menggantikan empat metrik utama"},
    {"field": "transmission_bursts", "unit": "burst advertising logis yang berhasil dimulai",
     "definition": "Awal burst yang benar-benar berhasil; keputusan allowed bukan awal burst; jumlah paket RF per burst tidak diukur",
     "attribution": "Seluruh jaringan"},
    {"field": "other_cancellations", "unit": "burst aktif yang dibatalkan",
     "definition": "ADVERTISE_BURST_CANCELLED; alasan ACK/state/admin/jendela tersimpan di All Events; bukan suppression",
     "attribution": "Hanya diagnostik; bukan penyebut suppression_rate"},
    {"field": "source_wait_before_first_advertise_ms", "unit": "milidetik",
     "definition": "SOURCE_FIRST_ADVERTISE_STARTED - SOS_CREATED pada sumber dan pesan yang sama, setelah koreksi jam; termasuk waktu tunggu scheduler",
     "attribution": "Diagnostik dari event sebelum t0; bukan perubahan E2E utama"},
    {"field": "sos_creation_to_destination_ms", "unit": "milidetik",
     "definition": "DESTINATION_FIRST_VALID_RECEIVE - SOS_CREATED untuk pesan dan hop tujuan yang cocok, setelah koreksi jam",
     "attribution": "Waktu sejak SOS dibuat; kosong jika bukti/jam tidak valid; tidak menggantikan E2E utama"},
    {"field": "burst_stop_reason", "unit": "kategori",
     "definition": "Alasan penghentian yang tercatat pada detail event; jika tidak ada, jangan menganggap dugaan dari durasi sebagai alasan pasti",
     "attribution": "Burst Diagnostics tidak menulis ulang All Events; burst pendek tetap masuk overhead jika berhasil dimulai"},
    {
        "field": "source_bursts_before_destination",
        "unit": "burst",
        "definition": "Event ADVERTISE_BURST_STARTED sumber pada atau sebelum penerimaan tujuan",
        "attribution": "Pengamatan langsung; bukan bukti kausal seluruh rute",
    },
    {
        "field": "delivery_after_last_source_burst_ms",
        "unit": "milidetik",
        "definition": "Waktu penerimaan tujuan dikurangi awal burst sumber terakhir sebelum penerimaan",
        "attribution": "Hanya diagnostik; bukan bukti kausal seluruh rute H2/H3",
    },
    {
        "field": "source_attempt_index_to_hop1_accept",
        "unit": "nomor percobaan",
        "definition": "Jumlah burst sumber yang dimulai sampai penerimaan pertama hop 1",
        "attribution": "Atribusi sumber tunggal; tidak mengidentifikasi paket RF individual",
    },
    {
        "field": "source_first_attempt_success",
        "unit": "boolean",
        "definition": "True jika penerimaan pertama hop 1 terjadi setelah hanya satu burst sumber dimulai",
        "attribution": "Atribusi sumber tunggal; bukan pengukuran keandalan universal",
    },
    {
        "field": "source_retry_wait_ms",
        "unit": "milidetik",
        "definition": "Awal burst sumber terpilih dikurangi awal burst sumber pertama",
        "attribution": "Atribusi sumber tunggal menuju lapisan 1",
    },
    {
        "field": "hopN_first_accept_elapsed_ms",
        "unit": "milidetik",
        "definition": "BLE_PACKET_ACCEPTED pertama pada hop N dikurangi awal advertising sumber pertama",
        "attribution": "Hop 1 bersumber tunggal; hop berikutnya hanya mengidentifikasi lapisan penerima",
    },
    {
        "field": "hopN_segment_progress_ms",
        "unit": "milidetik",
        "definition": "Penerimaan pertama hop N dikurangi penerimaan pertama hop N-1",
        "attribution": "Hop 1 bersumber tunggal; hop berikutnya berupa selisih waktu antarlapisan, bukan link perangkat tertentu",
    },
    {
        "field": "hopN_upstream_bursts_before_accept",
        "unit": "burst",
        "definition": "Burst pembawa hop N yang dimulai pada atau sebelum penerimaan pertama hop N",
        "attribution": "Hanya awal burst teramati; suppression Trickle bukan percobaan transmisi",
    },
    {
        "field": "hopN_attribution",
        "unit": "kategori",
        "definition": "exact_single_source: sumber tunggal; layer_inferred_parallel_relays: atribusi lapisan dengan relay paralel",
        "attribution": "Mencegah klaim link perangkat tertentu saat relay paralel tidak dapat dibedakan",
    },
]


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _ordered_fields(rows: list[dict[str, Any]], preferred: Iterable[str]) -> list[str]:
    available = {key for row in rows for key in row}
    ordered = [field for field in preferred if field in available]
    return ordered + sorted(available.difference(ordered))


def _safe_table_name(sheet_name: str, index: int) -> str:
    compact = "".join(character for character in sheet_name if character.isalnum())
    return f"ResQMesh{compact[:20]}{index}"


def _write_table(
    workbook: Workbook,
    sheet_name: str,
    rows: list[dict[str, Any]],
    preferred_fields: Iterable[str] = (),
) -> None:
    sheet = workbook.create_sheet(sheet_name)
    fields = _ordered_fields(rows, preferred_fields)
    if not fields:
        sheet.append(["status"])
        sheet.append(["Tidak ada data"])
        fields = ["status"]
    else:
        sheet.append(fields)
        for row in rows:
            sheet.append([_excel_value(row.get(field)) for field in fields])

    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 2:
        table = Table(
            displayName=_safe_table_name(sheet_name, len(workbook.worksheets)),
            ref=sheet.dimensions,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    for column_index, field in enumerate(fields, start=1):
        values = [str(field)] + [
            str(_excel_value(row.get(field)) or "") for row in rows[:250]
        ]
        width = min(max(max(map(len, values)) + 2, 11), 48)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        if field in {"dsr", "ldr"}:
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.number_format = "0.0000"
        elif field.endswith("_ms"):
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.number_format = "0"
        if field in {"definition", "attribution", "formula", "interpretation", "validity_rule",
                     "explanation", "overhead_note", "analysis"}:
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    if "result" in fields:
        result_column = get_column_letter(fields.index("result") + 1)
        result_range = f"{result_column}2:{result_column}{sheet.max_row}"
        sheet.conditional_formatting.add(
            result_range,
            CellIsRule(operator="equal", formula=['"SUCCESS"'], fill=_SUCCESS_FILL),
        )
        sheet.conditional_formatting.add(
            result_range,
            CellIsRule(operator="notEqual", formula=['"SUCCESS"'], fill=_FAILURE_FILL),
        )


def _event_sheets(events: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    if not events:
        return [("All Events", [])]
    chunks = []
    for start in range(0, len(events), EXCEL_MAX_DATA_ROWS):
        number = (start // EXCEL_MAX_DATA_ROWS) + 1
        name = "All Events" if number == 1 else f"All Events {number}"
        chunks.append((name, events[start : start + EXCEL_MAX_DATA_ROWS]))
    return chunks


def _method_parameter_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    # Keep the archived configuration/fingerprint intact, including legacy shared fields.
    rows = []
    for mode, parameters in (manifest.get("method_parameters") or {}).items():
        row = {
            "mode": mode,
            **dict.fromkeys(("scheduler", "suppression_enabled", "imin_ms", "imax_ms",
                             "imax_doublings", "k", "burst_ms", "basic_wait_ms",
                             "jitter_min_ms", "jitter_max_ms")),
            **parameters,
        }
        if mode == "basic_flooding":
            for field in ("imin_ms", "imax_ms", "imax_doublings", "k"):
                row[field] = "N/A"
            row["transmission_timing"] = (
                "SOS pertama langsung dapat dijadwalkan; pengulangan setelah burst selesai + jeda Basic + jitter"
            )
            row["parameter_notes"] = "Tidak memakai interval Trickle atau suppression berdasarkan c < k"
        elif mode in {"trickle", "trickle_no_suppression"}:
            for field in ("basic_wait_ms", "jitter_min_ms", "jitter_max_ms"):
                row[field] = "N/A"
            row["transmission_timing"] = (
                "t dipilih acak dalam [I/2, I); interval berlipat dua sampai Imax; tanpa jeda Basic/jitter"
            )
            if mode == "trickle_no_suppression":
                row["k"] = "N/A"
                row["parameter_notes"] = (
                    "Suppression dinonaktifkan; k konfigurasi dipertahankan di Manifest Metadata, "
                    "bukan syarat untuk mengizinkan transmisi"
                )
            else:
                row["parameter_notes"] = "Transmisi hanya jika c < k pada waktu yang dipilih"
        if row.get("termination") == "ACK/newer state/admin deletion; no hard TTL/hop/count":
            row["termination"] = "ACK/state lebih baru/penghapusan administratif; tanpa batas keras TTL/hop/relay count"
        rows.append(row)
    return rows


def _burst_diagnostic_rows(events: list[dict[str, Any]], manifest: dict[str, Any]) -> list[dict[str, Any]]:
    def identity(event):
        return (event.get("trial_id"), event.get("node_id", event.get("device_id")), event.get("burst_id"))

    def detail(event):
        try:
            value = json.loads(event.get("detail_json") or "{}")
            return value if isinstance(value, dict) else {}
        except (TypeError, ValueError):
            return {}

    starts = {identity(e): e for e in events if e.get("event_type") == "ADVERTISE_BURST_STARTED"}
    rows = []
    for end in events:
        if end.get("event_type") not in {"ADVERTISE_BURST_ENDED", "ADVERTISE_BURST_CANCELLED"}:
            continue
        start = starts.get(identity(end))
        if start is None or not end.get("burst_id"):
            continue
        data = detail(end)
        target = data.get("target_duration_ms", detail(start).get("target_duration_ms"))
        if target is None:
            target = (manifest.get("method_parameters", {}).get(start.get("mode"), {})
                      .get("burst_ms"))
        actual = data.get("actual_duration_ms")
        if actual is None:
            a = start.get("monotonic_ms", start.get("elapsed_realtime_ms"))
            b = end.get("monotonic_ms", end.get("elapsed_realtime_ms"))
            if a is not None and b is not None:
                actual = (int(b) - int(a)) & 0xFFFFFFFF
        reason = data.get("stop_reason") or end.get("reason")
        shortened = actual < target if actual is not None and target is not None else None
        rows.append({
            "trial_id": end.get("trial_id"), "algorithm": end.get("mode"),
            "node_id": identity(end)[1], "burst_id": end.get("burst_id"),
            "terminal_event": end.get("event_type"), "target_duration_ms": target,
            "actual_duration_ms": actual, "shortened": shortened, "burst_stop_reason": reason,
            "explanation": ("Alasan penghentian tercatat pada event" if reason else
                            "Burst lebih pendek dari target; alasan tidak tercatat, jangan dianggap pasti selesai normal"
                            if shortened else "Alasan penghentian tidak tercatat"),
            "overhead_note": "Tetap dihitung sekali karena ADVERTISE_BURST_STARTED berhasil; event asli tidak diubah",
        })
    return rows


def write_analysis_workbook(
    output_path: Path,
    *,
    manifest: dict[str, Any],
    events: list[dict[str, Any]],
    trial_summaries: list[dict[str, Any]],
    aggregates: list[dict[str, Any]],
    attempts: list[dict[str, Any]],
    invalid_trials: list[dict[str, Any]],
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    metadata = [
        {"field": "session_id", "value": manifest.get("session_id")},
        {"field": "config_fingerprint", "value": manifest.get("config_fingerprint")},
        {"field": "protocol_version", "value": manifest.get("protocol_version")},
        {"field": "protocol_epoch_id", "value": manifest.get("epoch_id")},
        {"field": "radio_mode", "value": manifest.get("radio_mode")},
        {"field": "radio_readiness", "value": manifest.get("radio_readiness")},
        {"field": "algorithms", "value": ", ".join(sorted({str(row.get("mode")) for row in trial_summaries if row.get("mode")}))},
        {"field": "trial_count", "value": len(trial_summaries)},
        {"field": "event_count", "value": len(events)},
        {"field": "invalid_trial_count", "value": len(invalid_trials)},
        {"field": "trial_order_strategy", "value": manifest.get("trial_order_strategy")},
        {"field": "trial_order_note", "value": (
            "Urutan dikelompokkan per metode/hop untuk smoke; bukan pengacakan blok eksperimen utama"
            if manifest.get("trial_order_strategy") == "blocked" else
            "Setiap blok memuat seluruh kombinasi metode-hop dengan urutan diacak; periksa Trial Order dan kolom block"
            if manifest.get("trial_order_strategy") == "balanced_randomized" else
            "Periksa strategi dan urutan aktual pada manifest; jangan menganggapnya acak tanpa bukti"
        )},
        {"field": "phy_claim_note", "value": "Coded PHY dibuktikan oleh telemetry penerimaan aktual, bukan konfigurasi yang diminta. Coding S2/S8 belum terverifikasi tanpa bukti khusus"},
        {"field": "topology_basis", "value": manifest.get("topology_basis")},
        {"field": "topology_claim_note", "value": "Jika memakai penyaringan hop logis, hasil mengevaluasi mekanisme forwarding, bukan bukti jangkauan fisik tiga hop atau isolasi RF"},
        {"field": "latency_note", "value": "E2E utama dimulai dari transmisi pertama. Latency Diagnostics juga memuat waktu tunggu sumber dan waktu sejak SOS_CREATED; bukan metrik utama pengganti"},
    ]
    _write_table(workbook, "Overview", metadata, ("field", "value"))
    _write_table(workbook, "Trial Order", [
        {"sequence": index, **row}
        for index, row in enumerate(manifest.get("trial_order") or [], 1)
    ], ("sequence", "block", "mode", "hypothesis", "attempt", "trial_id"))
    _write_table(workbook, "Method Parameters", _method_parameter_rows(manifest),
                 ("mode", "scheduler", "suppression_enabled", "imin_ms", "imax_ms",
                  "imax_doublings", "k", "burst_ms", "basic_wait_ms", "jitter_min_ms",
                  "jitter_max_ms", "transmission_timing", "parameter_notes", "termination"))
    for row in workbook["Method Parameters"].iter_rows(min_row=2):
        for cell in row:
            if cell.value == "N/A":
                cell.font = Font(color="666666", italic=True)
            if cell.column >= 12:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    _write_table(workbook, "Method Comparisons", descriptive_comparisons(aggregates),
                 ("comparison", "hypothesis", "metric", "baseline", "treatment", "baseline_value", "treatment_value", "difference_treatment_minus_baseline"))
    _write_table(workbook, "Suppression Summary", [
        {key: row.get(key) for key in ("mode", "hypothesis", "valid_trials", "transmission_opportunities",
          "transmission_allowed", "transmission_suppressed", "transmission_missed", "missed_opportunities_total",
          "suppression_rate", "actual_bursts", "other_cancellations")}
        for row in aggregates
    ])
    _write_table(workbook, "RX PHY", [e for e in events if e.get("event_type") == "BLE_RX_PHY_OBSERVED"],
                 ("session_id", "trial_id", "node_id", "mode", "observation_id", "timestamp_ms", "primary_phy", "secondary_phy", "legacy", "coding"))
    _write_table(
        workbook,
        "Metric Definitions",
        METRIC_DEFINITIONS,
        ("metric", "abbreviation", "unit", "level", "formula", "interpretation", "validity_rule"),
    )
    _write_table(
        workbook,
        "Algorithm Summary",
        aggregates,
        (
            "mode",
            "hypothesis",
            "valid_trials",
            "success_trials",
            "dsr",
            "e2e_count",
            "e2e_min",
            "e2e_max",
            "e2e_mean",
            "e2e_median",
            "e2e_sample_stddev",
            "e2e_q1",
            "e2e_q3",
            "e2e_iqr",
            "ldr",
            "transmission_overhead",
        ),
    )
    trial_rows = []
    for row in trial_summaries:
        values = {
            **row,
            "algorithm": row.get("mode"),
            "delivery_success": (
                {"SUCCESS": 1, "FAILED_DELIVERY": 0}.get(row.get("result"))
                if row.get("valid") is True else None
            ),
        }
        trial_rows.append({field: values.get(field) for field in TRIAL_METRIC_FIELDS})
    _write_table(
        workbook,
        "Trial Metrics",
        trial_rows,
        TRIAL_METRIC_FIELDS,
    )
    latency_rows = [
        {
            field: row.get("mode") if field == "algorithm" else row.get(field)
            for field in LATENCY_DIAGNOSTIC_FIELDS
        }
        for row in trial_summaries
    ]
    _write_table(
        workbook,
        "Latency Diagnostics",
        latency_rows,
        LATENCY_DIAGNOSTIC_FIELDS,
    )
    _write_table(
        workbook,
        "Diagnostic Definitions",
        DIAGNOSTIC_DEFINITIONS,
        ("field", "unit", "definition", "attribution"),
    )
    _write_table(workbook, "Burst Diagnostics", _burst_diagnostic_rows(events, manifest),
                 ("trial_id", "algorithm", "node_id", "burst_id", "terminal_event",
                  "target_duration_ms", "actual_duration_ms", "shortened", "burst_stop_reason",
                  "explanation", "overhead_note"))
    _write_table(
        workbook,
        "Event Type Counts",
        [
            {"event_type": event_type, "count": count}
            for event_type, count in sorted(
                Counter(str(event.get("event_type") or "UNKNOWN") for event in events).items()
            )
        ],
        ("event_type", "count"),
    )
    _write_table(
        workbook,
        "Attempts",
        attempts,
        ("trial_id", "mode", "hypothesis", "attempt", "result", "terminal", "invalid_reasons"),
    )
    _write_table(
        workbook,
        "Invalid Trials",
        invalid_trials,
        ("trial_id", "mode", "hypothesis", "result", "invalid_reasons"),
    )
    manifest_metadata = [
        {"field": key, "value": value}
        for key, value in sorted(manifest.items())
        if key != "trials"
    ]
    _write_table(
        workbook,
        "Manifest Metadata",
        manifest_metadata,
        ("field", "value"),
    )
    manifest_trials = [
        {"trial_id": trial_id, **record}
        for trial_id, record in sorted((manifest.get("trials") or {}).items())
    ]
    _write_table(
        workbook,
        "Manifest Trials",
        manifest_trials,
        ("trial_id", "device_trial_id", "mode", "hypothesis", "attempt", "result", "terminal"),
    )
    for name, rows in _event_sheets(events):
        _write_table(
            workbook,
            name,
            rows,
            (
                "session_id",
                "trial_id",
                "mode",
                "hypothesis",
                "node_id",
                "role",
                "event_type",
                "event_sequence",
                "within_observation_window",
                "timestamp_ms",
                "event_timestamp_ms",
                "message_key",
                "state_identity",
                "observation_id",
                "burst_id",
                "packet_type",
                "status",
                "hop_in",
                "hop_out",
                "rssi",
            ),
        )

    if aggregates:
        modes = sorted({r["mode"] for r in aggregates})
        indexed = {(r["mode"], r["hypothesis"]): r for r in aggregates}
        chart_sheet = workbook.create_sheet("Charts")
        data_sheet = workbook.create_sheet("Chart Data")
        for index, (metric, title, unit) in enumerate([
            ("dsr", "Delivery Success Ratio", "rasio"),
            ("e2e_median", "Median latensi pengiriman", "ms"),
            ("ldr", "Logical Duplicate Ratio agregat", "rasio jumlah"),
            ("transmission_overhead", "Overhead burst logis seluruh jaringan", "burst / trial valid"),
        ]):
            start = data_sheet.max_row + 2
            for offset, values in enumerate([["hypothesis", *modes], *[
                [hop, *[indexed.get((mode, hop), {}).get(metric) for mode in modes]]
                for hop in ("H1", "H2", "H3")
            ]]):
                for column, value in enumerate(values, 1):
                    data_sheet.cell(start + offset, column, value)
            chart = BarChart()
            chart.title, chart.y_axis.title, chart.x_axis.title = title, unit, "Hop logis"
            chart.add_data(Reference(data_sheet, min_col=2, max_col=1 + len(modes), min_row=start, max_row=start + 3), titles_from_data=True)
            chart.set_categories(Reference(data_sheet, min_col=1, min_row=start + 1, max_row=start + 3))
            chart.height, chart.width = 10, 22
            chart_sheet.add_chart(chart, f"A{1 + index * 21}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
