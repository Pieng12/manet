"""ESP-only views; canonical logs and the four metric definitions are retained."""
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment

from .excel_report import _write_table
from .neighbor_excel_layout import _style, SCENARIO_NAMES
from .neighbor_testbed import ESP_SCENARIOS, testbed


def cumulative_rows(events, manifest):
    from .neighbor_experiment import summarize_network
    rows = []
    for trial_id, record in manifest["trials"].items():
        start, end = record.get("observation_started_at_ms"), record.get("observation_ended_at_ms")
        if start is None or end is None:
            continue
        scoped = [e for e in events if e.get("trial_id") == trial_id]
        seconds = (end-start)//1000
        points = sorted({*range(0, seconds+1, 5), seconds})
        for elapsed in points:
            metrics = summarize_network(scoped, {**record, "trial_id": trial_id,
                                        "session_id": manifest["session_id"],
                                        "observation_ended_at_ms": start+elapsed*1000})
            rows.append({"trial_id": trial_id, "method": record["mode"],
                         "scenario": record["hypothesis"], "valid": record.get("result") in {"SUCCESS", "FAILED_DELIVERY"},
                         "elapsed_seconds": elapsed, "received_targets": metrics["U"],
                         "dsr_percent": metrics["dsr_percent"], "data_tx": metrics["data_tx"],
                         "control_tx": metrics["control_tx"], "network_overhead": metrics["network_overhead"]})
    return rows


def decorate_esp_only(workbook, manifest, validation, curves):
    design = testbed(manifest)
    guide = workbook["Mulai Di Sini"]
    guide["A1"] = "RESQMESH - SMOKE ESP-ONLY (BUKAN DATA UTAMA ANDROID)"
    guide.row_dimensions[1].height = 54
    for row in guide.iter_rows(min_row=2):
        if row[0].value == "Window pengamatan":
            row[1].value = "S0/S1: 180 detik; S2: 420 detik. Bandingkan metode dalam skenario yang sama, bukan total lintas durasi."
    guide.append(["Perangkat dan denominator", "Lima ESP: satu SOURCE dan empat target. DSR = U/(M*4)*100; sumber tidak dihitung sebagai target."])
    guide.append(["Batas penerimaan smoke", "Batch lengkap, delivery sukses, dan bukti mekanisme berbeda. INCONCLUSIVE bukan PASS. ESP-only tidak memvalidasi Android."])
    guide.append(["Profil", manifest["testbed_profile"]])
    rows = [{"kategori": "node", "identitas": node, "peran": "SOURCE" if node == design.source else "RELAY / target",
             "tetangga": ", ".join(b if a == node else a for a, b in design.edges if node in (a, b))}
            for node in design.node_ids]
    rows += [{"kategori": "skenario", "identitas": SCENARIO_NAMES[s],
              "durasi_detik": p["observation_window_seconds"], "aktivasi_detik": p["activate_at_seconds"],
              "node_OFF_sebelum_SOS": ", ".join(p["inactive_node_ids"])} for s, p in ESP_SCENARIOS.items()]
    _write_table(workbook, "Topologi dan Skenario", rows)
    _write_table(workbook, "Bukti Skenario", [c for c in validation if c["check"].startswith("ESP_")])
    _write_table(workbook, "Kurva Kumulatif", curves)
    for name in ("Topologi dan Skenario", "Bukti Skenario", "Kurva Kumulatif"):
        _style(workbook[name])
        guide.append([name, "Profil pilot terpisah; kurva dihitung dari log, tanpa mengubah data mentah."])
        guide.cell(guide.max_row, 1).hyperlink = f"#'{name}'!A1"
        guide.cell(guide.max_row, 1).style = "Hyperlink"

    meanings = {
        "N": "Empat target ESP A/C/D/E; SOURCE ESP tidak masuk denominator DSR.",
        "dsr_percent": "U/(M*4)*100. FAILED_DELIVERY sah tetap masuk denominator.",
        "data_tx": "Burst DATA berhasil dimulai pada kelima ESP dalam window; bukan request, paket RF, atau energi.",
        "failed_pairs": "M*4-U. Bukan jumlah trial INVALID.",
        "successful_pairs": "Jumlah pasangan pesan-target sukses; satu trial mempunyai paling banyak empat pasangan.",
    }
    for sheet in workbook:
        for cell in sheet[1]:
            field = str(cell.value)
            base = field.removesuffix("_trial_mean")
            if base in meanings:
                cell.comment = Comment(meanings[base], "ResQMesh")
    glossary = workbook["Kamus Kolom"]
    for row in glossary.iter_rows(min_row=2):
        base = str(row[1].value).removesuffix("_trial_mean")
        if base in meanings:
            row[3].value = meanings[base]

    graphs = workbook.create_sheet("Grafik Kumulatif")
    graphs.sheet_view.showGridLines = False
    graphs["A1"] = "Delivery dan overhead kumulatif - hanya trial sah; waktu sejak DATA sumber pertama (detik)"
    data = workbook["Kurva Kumulatif"]
    headers = [c.value for c in data[1]]
    for index, scenario in enumerate(ESP_SCENARIOS):
        for metric_index, (field, title, unit) in enumerate((
                ("dsr_percent", "Delivery kumulatif", "%"),
                ("network_overhead", "Overhead DATA + CONTROL kumulatif", "burst"))):
            chart = ScatterChart()
            chart.title = SCENARIO_NAMES[scenario] + ": " + title
            chart.x_axis.title, chart.y_axis.title = "Detik sejak DATA sumber pertama", unit
            chart.width, chart.height = 23, 11
            for trial_id in dict.fromkeys(r["trial_id"] for r in curves if r["scenario"] == scenario and r["valid"]):
                indices = [i+2 for i, r in enumerate(curves) if r["trial_id"] == trial_id]
                x = Reference(data, min_col=headers.index("elapsed_seconds")+1, min_row=min(indices), max_row=max(indices))
                y = Reference(data, min_col=headers.index(field)+1, min_row=min(indices), max_row=max(indices))
                chart.series.append(Series(y, x, title=trial_id))
            if unit == "%":
                chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
            graphs.add_chart(chart, f"{'A' if metric_index == 0 else 'N'}{4+index*24}")
    guide.append(["Grafik Kumulatif", "Kurva per trial sah; tidak mencampur skenario dan durasi yang berbeda."])
    from openpyxl.styles import Alignment
    for row in guide.iter_rows(min_row=2):
        guide.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.move_sheet("Topologi dan Skenario", offset=1-workbook.sheetnames.index("Topologi dan Skenario"))
    workbook.move_sheet("Bukti Skenario", offset=2-workbook.sheetnames.index("Bukti Skenario"))
