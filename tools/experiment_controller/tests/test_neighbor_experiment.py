import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from resqmesh_controller.config import ConfigError, validate_config, research_fingerprint
from resqmesh_controller.controller import build_plan, TrialSpec, BatchIncompleteError
from resqmesh_controller.devices import DeviceError, SerialNode
from resqmesh_controller.neighbor_experiment import (
    SOURCE, TARGETS, METHODS, SCENARIOS, EDGES, adjacency, stable_id,
    NeighborExperimentController, summarize_network, merge_neighbor,
    neighbor_parameters, neighbor_plan,
)
from test_controller import FakeNode

ROOT = Path(__file__).resolve().parents[3]

def config(full=False):
    name = 'full' if full else 'main'
    value=json.loads((ROOT/f'tools/experiment_controller/config.neighbor.{name}.example.json').read_text())
    value['session_id']='software-fixture-neighbor'
    value['android_build_id']=value['firmware_build_id']='0123456789ab'
    return value

def record():
    return {'session_id':'fixture','device_trial_id':'fixture-t1','mode':'trickle_neighbor_status',
            'hypothesis':'S0_MAIN','result':'FAILED_DELIVERY','clock_tolerance_ms':100,
            'observation_started_at_ms':1000,'observation_ended_at_ms':2000,
            'scope':stable_id('fixture-t1')}

def ev(kind,node,at,sequence=1,tx=None,scope=None):
    return {'event_type':kind,'node_id':node,'session_id':'fixture','trial_id':'t1',
            'message_key':'123:456','timestamp_ms':at,'clock_sync_valid':True,'clock_offset_ms':0,
            'scope':stable_id('fixture-t1') if scope is None else scope,
            'transmitter_id':stable_id(node) if tx is None else tx,'boot_id':1,'transmission_sequence':sequence}

def received(node,sequence=1):
    return ev('DATA_RECEIVED',node,1100,sequence,adjacency(node)[0])

def source_events():
    return [ev('SOS_CREATED',SOURCE,900),ev('SOURCE_FIRST_ADVERTISE_STARTED',SOURCE,1000),ev('DATA_BURST_STARTED',SOURCE,1000)]

class NeighborFake(FakeNode):
    def command(self,name,args):
        result=super().command(name,args)
        if name in {'readiness','get_status'}:
            result.update(build_id=self.config['android_build_id'],neighbor_design_version=1,
                          transport_profile='neighbor_graph_v1')
            result['radio']['maximum_advertising_data_length'] = 1650
            from resqmesh_controller.neighbor_experiment import neighbor_parameters
            result.update(transmitter_id=stable_id(self.node_id),allowed_transmitters=adjacency(self.node_id),
                          scope=stable_id(self.trial_id) if self.trial_id else None,
                          transport_version='resqmesh-neighbor-v1', data_frame_length=39,
                          neighbor_parameters=neighbor_parameters(self.config))
        if name in {'set_rx_participation','set_node_participation'}:
            result['confirmed_enabled']=args['enabled']
        return result

    def collect_events(self,session_id=None,trial_id=None):
        stop=next(args for name,args in reversed(self.commands) if name=='end_observation_window')
        t0=stop['observation_ended_at_ms']-int(self.config['observation_window_seconds']*1000)
        values=[]
        if self.node_id==SOURCE:
            values=[ev('SOS_CREATED',SOURCE,t0-1),ev('SOURCE_FIRST_ADVERTISE_STARTED',SOURCE,t0),ev('DATA_BURST_STARTED',SOURCE,t0)]
        else:
            values=[ev('TRIAL_WINDOW_STARTED',self.node_id,t0-1)]
            if self.node_id not in self.config.get('fixture_missing',[]):
                values.append(ev('DATA_RECEIVED',self.node_id,t0+1,tx=adjacency(self.node_id)[0]))
        for seq,e in enumerate(values,1):
            e.update(session_id=session_id,trial_id=trial_id,event_sequence=seq,scope=stable_id(trial_id),message_key='100:200')
        return values

class NeighborExperimentTest(unittest.TestCase):
    def test_adaptive_pilot_has_three_randomized_complete_blocks(self):
        c=json.loads((ROOT/'tools/experiment_controller/config.neighbor.adaptive.pilot.example.json').read_text())
        c['android_build_id']=c['firmware_build_id']='0123456789ab'
        validate_config(c)
        plan=neighbor_plan(c)
        self.assertEqual(36,len(plan))
        canonical=[(m,s) for m in METHODS for s in SCENARIOS]
        for block in range(1,4):
            values=[(p.mode,p.hypothesis) for p in plan if p.block==block]
            self.assertCountEqual(canonical,values)
            self.assertNotEqual(canonical,values)
        self.assertEqual(180,c['observation_window_seconds'])
        self.assertEqual((60000,150000),(neighbor_parameters(c)['status_period_ms'],neighbor_parameters(c)['freshness_ms']))

    def test_adaptive_parameters_fingerprinted_and_legacy_defaults_retained(self):
        old=config(True)
        self.assertEqual('periodic_v1',neighbor_parameters(old)['neighbor_status_policy'])
        self.assertEqual(12000,neighbor_parameters(old)['status_period_ms'])
        new={**old,'neighbor_status_policy':'adaptive_v2','neighbor_parameters':{}}
        self.assertNotEqual(research_fingerprint(old),research_fingerprint(new))
        changed={**new,'neighbor_parameters':{'data_grace_ms':11000}}
        self.assertNotEqual(research_fingerprint(new),research_fingerprint(changed))
        for overrides in ({'freshness_ms':120000},{'empty_retry_max_ms':2000},{'discovery_jitter_ms':True},{'unknown':1}):
            with self.subTest(overrides=overrides),self.assertRaises(ConfigError):
                validate_config({**new,'neighbor_parameters':overrides})
        with self.assertRaises(ConfigError):
            validate_config({**new,'neighbor_status_policy':'unknown'})

    def test_adaptive_config_readiness_and_manifest_match_all_nodes(self):
        c={**config(True),'neighbor_status_policy':'adaptive_v2','neighbor_parameters':{}}
        with tempfile.TemporaryDirectory() as d:
            nodes=[NeighborFake(n,c,None) for n in c['nodes']]
            controller=NeighborExperimentController(c,nodes,Path(d))
            spec=TrialSpec(METHODS[-1],'S0_MAIN',1)
            controller.configure(spec)
            self.assertEqual(neighbor_parameters(c),controller.manifest['neighbor_parameters'])
            self.assertEqual('adaptive_v2',controller.manifest['neighbor_status_policy'])
            for node in nodes:
                args=next(a for name,a in node.commands if name=='configure_session')
                self.assertEqual(10000,args['data_grace_ms'])
                response=node.command('readiness',{})
                self.assertNotIn('neighbor_parameters mismatch',controller._readiness_errors(node,response,spec,False,False))
                response['neighbor_parameters'].pop('neighbor_status_policy')
                self.assertIn('neighbor_parameters mismatch',controller._readiness_errors(node,response,spec,False,False))

    def test_old_readiness_without_policy_still_matches_periodic(self):
        c=config()
        with tempfile.TemporaryDirectory() as d:
            node=NeighborFake(c['nodes'][0],c,None)
            controller=NeighborExperimentController(c,[node],Path(d))
            spec=TrialSpec(METHODS[-1],'S0_MAIN',1)
            controller.configure(spec)
            response=node.command('readiness',{})
            response['neighbor_parameters'].pop('neighbor_status_policy')
            self.assertNotIn('neighbor_parameters mismatch',controller._readiness_errors(node,response,spec,False,False))

    def clock_controller(self, directory):
        c=config()
        node=NeighborFake(c['nodes'][0],c,None)
        controller=NeighborExperimentController(c,[node],Path(directory),sleep=lambda _:None)
        spec=TrialSpec(METHODS[0],'S0_MAIN',1)
        controller.manifest['trials'][spec.trial_id]={}
        return controller,node,spec

    def test_clock_retries_slow_adb_sample_without_raising_tolerance(self):
        with tempfile.TemporaryDirectory() as d:
            controller,node,spec=self.clock_controller(d)
            with patch.object(node,'host_clock_offset_ms',side_effect=[900,940]), \
                 patch('resqmesh_controller.neighbor_experiment.time.time_ns',side_effect=[0,198400000,200000000,240000000,241000000]):
                controller.synchronize_clocks(spec)
            sample=controller.manifest['trials'][spec.trial_id]['clock_samples'][SOURCE]
            self.assertEqual(21,sample['uncertainty_ms'])
            self.assertEqual(940,sample['offset_ms'])
            attempts=controller.manifest['trials'][spec.trial_id]['clock_sync_attempts'][SOURCE]
            self.assertEqual([False,True],[a['accepted'] for a in attempts])
            self.assertAlmostEqual(100.2,attempts[0]['uncertainty_ms'])
            self.assertEqual(100,controller.config['clock_tolerance_ms'])

    def test_clock_all_slow_samples_still_fail_with_attempt_evidence(self):
        with tempfile.TemporaryDirectory() as d:
            controller,node,spec=self.clock_controller(d)
            with patch.object(node,'host_clock_offset_ms',return_value=900) as probe, \
                 patch('resqmesh_controller.neighbor_experiment.time.time_ns',side_effect=[0,220000000,221000000,441000000,442000000,662000000]):
                with self.assertRaisesRegex(DeviceError,'clock uncertainty exceeds tolerance'):
                    controller.synchronize_clocks(spec)
            self.assertEqual(3,probe.call_count)
            self.assertNotIn(SOURCE,controller.clock_offsets)
            attempts=controller.manifest['trials'][spec.trial_id]['clock_sync_attempts'][SOURCE]
            self.assertEqual(3,len(attempts))
            self.assertTrue(all(not a['accepted'] for a in attempts))

    def test_serial_clock_retry_uses_unique_command_and_latest_accepted_offset(self):
        with tempfile.TemporaryDirectory() as d:
            c=config()
            node=SerialNode(TARGETS[0],'RELAY','COM_TEST')
            controller=NeighborExperimentController(c,[node],Path(d))
            spec=TrialSpec(METHODS[0],'S0_MAIN',1)
            controller.manifest['trials'][spec.trial_id]={}
            with patch.object(node,'command',return_value={'ok':True}) as command, \
                 patch('resqmesh_controller.neighbor_experiment.time.time_ns',side_effect=[1000000000,1198400000,1200000000,1220000000,1221000000]):
                controller.synchronize_clocks(spec)
            calls=[call.args[1] for call in command.call_args_list]
            self.assertEqual(2,len({a['command_id'] for a in calls}))
            self.assertEqual([1000,1200],[a['wall_time_ms'] for a in calls])
            self.assertEqual(10,controller.clock_offsets[node.node_id])
            self.assertEqual(11,controller.manifest['trials'][spec.trial_id]['clock_samples'][node.node_id]['uncertainty_ms'])

    def test_serial_clock_rejection_is_not_retried_or_hidden(self):
        with tempfile.TemporaryDirectory() as d:
            c=config()
            node=SerialNode(TARGETS[0],'RELAY','COM_TEST')
            controller=NeighborExperimentController(c,[node],Path(d))
            spec=TrialSpec(METHODS[0],'S0_MAIN',1)
            controller.manifest['trials'][spec.trial_id]={}
            with patch.object(node,'command',return_value={'ok':False}) as command:
                with self.assertRaisesRegex(DeviceError,'clock sync failed'):
                    controller.synchronize_clocks(spec)
            self.assertEqual(1,command.call_count)
            self.assertNotIn(node.node_id,controller.clock_offsets)

    def test_main60_full180_balanced_reproducible_blocks(self):
        for full,length,per_block in [(False,60,4),(True,180,12)]:
            c=config(full);validate_config(c);plan=build_plan(c)
            self.assertEqual(length,len(plan));self.assertEqual(plan,build_plan(c))
            for block in range(1,16):
                rows=[p for p in plan if p.block==block]
                self.assertEqual(per_block,len(rows));self.assertEqual(per_block,len({(r.mode,r.hypothesis) for r in rows}))
            self.assertNotEqual([r.mode for r in plan[:per_block]],list(METHODS)*len(c['scenarios']))

    def test_graph_is_undirected_with_isolated_branch_and_all_esp_relay(self):
        self.assertEqual([stable_id('esp-r2b')],adjacency('esp-destination'))
        self.assertNotIn(stable_id('esp-r1b'),adjacency('esp-r2b'))
        for a,b in EDGES:
            self.assertIn(stable_id(a),adjacency(b));self.assertIn(stable_id(b),adjacency(a))

    def test_config_rejects_duplicate_devices_unsafe_params_and_ack(self):
        for key,value in [('ack_enabled',True),('gateway_enabled',True),('observation_window_seconds',0),('modes',['trickle'])]:
            c=config();c[key]=value
            with self.assertRaises(ConfigError):validate_config(c)
        c=config(True);c['observation_window_seconds']=60
        with self.assertRaises(ConfigError):validate_config(c)
        c=config();c['nodes'][2]['port']=c['nodes'][1]['port']
        with self.assertRaises(ConfigError):validate_config(c)

    def test_metric_dsr80_and_failed_receiver_latency_null(self):
        m=summarize_network(source_events()+[received(n) for n in TARGETS[:4]],record())
        self.assertEqual(80,m['dsr_percent']);self.assertEqual(4,m['U']);self.assertEqual(100,m['e2e_mean_ms'])
        self.assertEqual(4,len(m['per_receiver']))

    def test_zero_receive_ldr_is_undefined_not_zero(self):
        m=summarize_network(source_events(),record())
        self.assertEqual(0,m['dsr_percent']);self.assertIsNone(m['ldr_percent']);self.assertIsNone(m['e2e_mean_ms'])

    def test_clock_uncertainty_is_pair_bound_not_configured_tolerance(self):
        r=record();r['clock_samples']={SOURCE:{'uncertainty_ms':7},TARGETS[0]:{'uncertainty_ms':4}}
        pair=summarize_network(source_events()+[received(TARGETS[0])],r)['per_receiver'][0]
        self.assertEqual(11,pair['clock_uncertainty_ms'])
        self.assertEqual(100,pair['clock_tolerance_ms'])
        self.assertIsNone(summarize_network(source_events()+[received(TARGETS[0])],record())['per_receiver'][0]['clock_uncertainty_ms'])

    def test_disconnected_transport_aborts_wait_instead_of_failed_delivery(self):
        import time
        c=config()
        with tempfile.TemporaryDirectory() as d:
            nodes=[NeighborFake(n,c,None) for n in c['nodes']]
            nodes[3].transport_error='serial disconnected'
            controller=NeighborExperimentController(c,nodes,Path(d))
            with self.assertRaisesRegex(DeviceError,'serial disconnected'):
                controller._wait_until(time.monotonic()+1)

    def test_rf_repetitions_replay_new_burst_and_forwarder(self):
        a=received('esp-r2a');other=received('esp-r2a');other['transmitter_id']=stable_id('esp-r1b')
        m=summarize_network(source_events()+[a,a.copy(),received('esp-r2a',2),other],record())
        self.assertEqual(3,m['R']);self.assertEqual(1,m['U']);self.assertAlmostEqual(200/3,m['ldr_percent'])
        reboot=copy.deepcopy(a);reboot['boot_id']=2
        self.assertEqual(2,summarize_network(source_events()+[a,reboot],record())['R'])

    def test_forbidden_edge_wrong_scope_source_rx_and_status_do_not_count(self):
        forbidden=received('esp-destination');forbidden['transmitter_id']=stable_id(SOURCE)
        wrong=received(TARGETS[0]);wrong['scope']+=1
        status=ev('STATUS_RECEIVED',TARGETS[0],1100,tx=stable_id(SOURCE))
        self.assertEqual(0,summarize_network(source_events()+[forbidden,wrong,status,received(SOURCE)],record())['R'])

    def test_overhead_only_successful_starts_control_setup_separate(self):
        events=source_events()+[ev('STATUS_BURST_STARTED',TARGETS[0],950),ev('STATUS_BURST_STARTED',TARGETS[0],1200,2),ev('STATUS_BURST_STARTED',TARGETS[0],1200,2),ev('DATA_BURST_REQUESTED',TARGETS[0],1500),ev('DATA_BURST_FAILED',TARGETS[0],1500),ev('DATA_BURST_STARTED',TARGETS[0],2000)]
        m=summarize_network(events,record())
        self.assertEqual((1,1,2,1),(m['data_tx'],m['control_tx'],m['network_overhead'],m['setup_control_tx']))

    def test_detail_json_metadata_and_physical_clock_correction(self):
        sample=received(TARGETS[0]);sample['timestamp_ms']=100
        sample['clock_offset_ms']=1000
        detail={k:sample.pop(k) for k in ('scope','transmitter_id','boot_id','transmission_sequence')}
        sample['detail_json']=json.dumps(detail)
        self.assertEqual(100,summarize_network(source_events()+[sample],record())['e2e_mean_ms'])

    def test_workbook_numeric_types_tables_and_raw_preservation(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);raw=p/'raw';raw.mkdir()
            events=source_events()+[received(n) for n in TARGETS[:4]]
            (raw/'events.jsonl').write_text('\n'.join(json.dumps(e) for e in events),encoding='utf-8')
            r=record();manifest={'session_id':'fixture','trials':{'t1':r},'neighbor_scenarios':['S0_MAIN']}
            result=merge_neighbor(raw,p/'merged',manifest)
            wb=load_workbook(result['workbook']);self.addCleanup(wb.close)
            self.assertIn('Method Scenario Summary',wb.sheetnames)
            ws=wb['Trial Metrics'];headers=[c.value for c in ws[1]]
            self.assertEqual(80,ws.cell(2,headers.index('dsr_percent')+1).value)
            self.assertIsInstance(ws.cell(2,headers.index('dsr_percent')+1).value,(int,float))
            receivers=list(wb['Receivers'].values);self.assertEqual(6,len(receivers))
            self.assertIsNone(receivers[-1][receivers[0].index('e2e_latency_ms')])
            self.assertEqual(events,json.loads((p/'merged/all_events.json').read_text()))
            for ws in wb:
                for table in ws.tables.values():self.assertEqual(len(ws[1]),len(table.tableColumns))

    def test_controller_failed_delivery_valid_smoke_and_resume_immutable(self):
        c=config();c.update(valid_trials_per_condition=1,max_attempts_per_condition=1,quiet_period_seconds=0,observation_window_seconds=1,fixture_missing=['esp-destination'])
        with tempfile.TemporaryDirectory() as d:
            nodes=[NeighborFake(n,c,None) for n in c['nodes']]
            controller=NeighborExperimentController(c,nodes,Path(d),sleep=lambda _:None)
            try:
                controller.run()
            except BatchIncompleteError:
                self.fail(json.dumps(controller.manifest['trials'],indent=2))
            self.assertTrue(controller.smoke_report()['passed'])
            self.assertTrue(all(r['result']=='FAILED_DELIVERY' and r['evidence']['dsr_percent']==80 for r in controller.manifest['trials'].values()))
            before={str(p):p.read_bytes() for p in Path(d).rglob('*.jsonl')}
            resumed=NeighborExperimentController(c,nodes,Path(d),sleep=lambda _:None)
            resumed.run()
            self.assertEqual(before,{str(p):p.read_bytes() for p in Path(d).rglob('*.jsonl')})
            changed={**c,'neighbor_parameters':{'freshness_ms':60000}}
            self.assertNotEqual(research_fingerprint(c),research_fingerprint(changed))
            with self.assertRaises(DeviceError):NeighborExperimentController(changed,nodes,Path(d))

    def test_actual_perturbation_commands_do_not_reset_running_trial(self):
        c=config(True);c.update(valid_trials_per_condition=1,quiet_period_seconds=0)
        with tempfile.TemporaryDirectory() as d:
            nodes=[NeighborFake(n,c,None) for n in c['nodes']]
            controller=NeighborExperimentController(c,nodes,Path(d),sleep=lambda _:None)
            for scenario,node_id,command in [('S1_DELAYED_RX','esp-r2b','set_rx_participation'),('S2_LATE_JOIN','esp-destination','set_node_participation')]:
                spec=TrialSpec(METHODS[3],scenario,1)
                controller.manifest['trials'][spec.trial_id]={}
                node=next(n for n in nodes if n.node_id==node_id)
                node.trial_id='running';controller.before_trigger(spec);controller._perturb(spec,True)
                self.assertEqual('running',node.trial_id)
                self.assertEqual([False,True],[a['enabled'] for n,a in node.commands if n==command])

    def test_readiness_preserves_measured_clock_and_rejects_foreign_trial(self):
        c=config()
        with tempfile.TemporaryDirectory() as d:
            nodes=[NeighborFake(n,c,None) for n in c['nodes']]
            controller=NeighborExperimentController(c,nodes,Path(d))
            spec=TrialSpec(METHODS[0],'S0_MAIN',1)
            controller.configure(spec)
            controller.synchronize_clocks(spec)
            count=sum(name=='clock_sync' for n in nodes for name,args in n.commands)
            controller.readiness(spec)
            self.assertEqual(count,sum(name=='clock_sync' for n in nodes for name,args in n.commands))
            nodes[1].trial_id='unarchived-old-session'
            readiness=controller.readiness()
            with self.assertRaisesRegex(DeviceError,'Unarchived/foreign'):
                controller.recover_interrupted_trials(readiness)
            self.assertFalse(any(name=='reset_trial' for n in nodes for name,args in n.commands))

if __name__=='__main__':unittest.main()
