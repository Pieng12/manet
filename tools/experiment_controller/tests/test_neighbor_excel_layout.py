import copy
import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from resqmesh_controller.mpl_config import DEFAULTS, METHODS, SEMANTICS
from resqmesh_controller.neighbor_experiment import TARGETS, merge_neighbor
from test_neighbor_experiment import record, received, source_events


class NeighborExcelLayoutTests(unittest.TestCase):
    def export(self, root, events=None, empty=False):
        events = events if events is not None else source_events() + [received(n) for n in TARGETS]
        raw = root / 'raw'
        raw.mkdir()
        (raw / 'events.jsonl').write_text('\n'.join(json.dumps(e) for e in events), encoding='utf-8')
        manifest = {'session_id': 'fixture', 'synthetic_data': True,
                    'experiment_methods': list(METHODS), 'neighbor_scenarios': [] if empty else ['S0_MAIN', 'S1_DELAYED_RX', 'S2_LATE_JOIN'],
                    'scheduler_semantics': SEMANTICS, 'mpl_parameters': DEFAULTS,
                    'trials': {} if empty else {'t1': {**record(), 'mode': 'trickle_mpl'}}}
        before = copy.deepcopy(manifest)
        result = merge_neighbor(raw, root / 'merged', manifest)
        self.assertEqual(before, manifest)
        self.assertEqual(events, json.loads((root / 'merged/all_events.json').read_text()))
        return result

    def test_reading_order_clear_columns_and_original_metrics_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = self.export(root)
            wb = load_workbook(output['workbook'])
            try:
                self.assertEqual(['Mulai Di Sini', 'Ringkasan Metode', 'Hasil Per Trial', 'Grafik Ringkas'], wb.sheetnames[:4])
                self.assertEqual('All Events', wb.sheetnames[-1])
                self.assertIn('DATA SINTETIS', wb['Mulai Di Sini']['B3'].value)
                self.assertEqual('C2', wb['Ringkasan Metode'].freeze_panes)
                headers = [c.value for c in wb['Hasil Per Trial'][1]]
                values = dict(zip(headers, [c.value for c in wb['Hasil Per Trial'][2]]))
                metrics = json.loads((root / 'merged/network_metrics.json').read_text())[0]
                for label, field in [('DSR (%)', 'dsr_percent'), ('LDR (%)', 'ldr_percent'),
                                     ('Delay pasangan sukses (ms)', 'e2e_mean_ms'),
                                     ('Total overhead window (burst)', 'network_overhead'),
                                     ('M - Pesan sumber', 'M'), ('R - Penerimaan DATA', 'R')]:
                    self.assertEqual(metrics[field], values[label])
                    self.assertIsInstance(values[label], (float, int))
                cell = wb['Hasil Per Trial'].cell(2, headers.index('DSR (%)') + 1)
                self.assertEqual('0.00"%"', cell.number_format)
                self.assertTrue(wb['Hasil Per Trial'].cell(1, 7).comment)
                self.assertEqual(['trial_id', 'method', 'scenario', 'block', 'result', 'valid'],
                                 [c.value for c in wb['Trial Metrics'][1]][:6])
                self.assertIsNotNone(wb['Trial Metrics']['O1'].comment)
                self.assertEqual(12, len(wb['Grafik Ringkas']._charts))
                glossary = list(wb['Kamus Kolom'].values)
                self.assertTrue(any(row[1] == 'U' and 'Pasangan' in row[2] for row in glossary[1:]))
                self.assertEqual(('parameter', 'value', 'unit', 'penjelasan'), tuple(c.value for c in wb['MPL Parameters'][1]))
                for sheet in wb:
                    for table in sheet.tables.values():
                        self.assertEqual(len(sheet[1]), len(table.tableColumns))
                        self.assertEqual([c.value for c in sheet[1]], [c.name for c in table.tableColumns])
                        self.assertEqual(len(table.tableColumns), len(set(c.name for c in table.tableColumns)))
            finally:
                wb.close()
            with ZipFile(output['workbook']) as archive:
                self.assertIsNone(archive.testzip())
                for name in archive.namelist():
                    if name.endswith('.xml'):
                        ET.fromstring(archive.read(name))

    def test_empty_archive_has_no_fabricated_metrics_or_invalid_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            output = self.export(Path(directory), events=[], empty=True)
            wb = load_workbook(output['workbook'])
            try:
                self.assertEqual('Tidak ada data', wb['Hasil Per Trial']['A2'].value)
                self.assertEqual('Tidak ada data', wb['Ringkasan Metode']['A2'].value)
                for sheet in wb:
                    for table in sheet.tables.values():
                        self.assertEqual(len(sheet[1]), len(table.tableColumns))
                self.assertEqual(0, len(wb['Grafik Ringkas']._charts))
            finally:
                wb.close()

    def test_missing_values_stay_blank_and_failed_delivery_is_not_invalid(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = self.export(root, events=source_events())
            wb = load_workbook(output['workbook'])
            try:
                row = dict(zip([c.value for c in wb['Hasil Per Trial'][1]],
                               [c.value for c in wb['Hasil Per Trial'][2]]))
                self.assertEqual('FAILED_DELIVERY', row['Hasil trial'])
                self.assertTrue(row['Trial sah'])
                self.assertIsNone(row['Delay pasangan sukses (ms)'])
                self.assertIsNone(row['LDR (%)'])
                self.assertEqual(0, row['DSR (%)'])
            finally:
                wb.close()


if __name__ == '__main__':
    unittest.main()
