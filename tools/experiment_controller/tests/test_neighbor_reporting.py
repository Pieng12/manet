import copy
import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from resqmesh_controller.neighbor_experiment import METHODS, SOURCE, TARGETS, merge_neighbor, stable_id, summarize_network, neighbor_parameters
from resqmesh_controller.neighbor_reporting import aggregate, descriptive, mechanism_counts, phy_evidence
from resqmesh_controller.neighbor_validation import validate_logs
from test_neighbor_experiment import ev, record, received, source_events


class NeighborReportingTests(unittest.TestCase):
    def test_adaptive_export_policy_numeric_parameters_raw_and_setup_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);raw=root/'raw';raw.mkdir()
            setup={**ev('STATUS_BURST_STARTED',TARGETS[0],950),'reason':'DISCOVERY'}
            status={**ev('STATUS_BURST_STARTED',TARGETS[0],1200,2),'reason':'HEARTBEAT'}
            coalesced={**ev('STATUS_COALESCED_WITH_DATA',SOURCE,1000),'reason':'DATA_STARTED'}
            events=source_events()+[received(TARGETS[0]),setup,status,coalesced,status.copy()]
            (raw/'events.jsonl').write_text('\n'.join(json.dumps(e) for e in events),encoding='utf-8')
            parameters=neighbor_parameters({'neighbor_status_policy':'adaptive_v2'})
            manifest={'synthetic_data':True,'session_id':'fixture','trials':{'t1':record()},
                      'neighbor_status_policy':'adaptive_v2','neighbor_parameters':parameters,
                      'method_parameters':{'trickle_neighbor_status':{'scheduler':'trickle'}},'neighbor_scenarios':['S0_MAIN']}
            manifest['pilot_baseline']={'session_id':'old-fixture','trials_per_scenario':1,
                                       'network_overhead_by_scenario':{'S0_MAIN':93},'s0_reduction_target_percent':50}
            original=copy.deepcopy(manifest)
            output=merge_neighbor(raw,root/'merged',manifest)
            self.assertEqual(original,manifest)
            self.assertEqual(events,json.loads((root/'merged/all_events.json').read_text()))
            metrics=json.loads((root/'merged/network_metrics.json').read_text())[0]
            self.assertEqual((1,1,2,1,3),tuple(metrics[k] for k in ('data_tx','control_tx','network_overhead','setup_control_tx','setup_plus_window_tx')))
            self.assertEqual('adaptive_v2',metrics['neighbor_status_policy'])
            diagnostics=json.loads((root/'merged/status_diagnostics.json').read_text())
            self.assertEqual(3,len(diagnostics))
            self.assertTrue(all(r['count']==1 for r in diagnostics))
            review=json.loads((root/'merged/adaptive_pilot_review.json').read_text())[0]
            self.assertEqual(93,review['baseline_overhead'])
            self.assertEqual(1,review['baseline_trials'])
            self.assertEqual(2,review['neighbor_overhead_mean'])
            self.assertEqual('FAILED_DELIVERY',metrics['result'])
            self.assertTrue(review['exploration_target_met'])
            wb=load_workbook(output['workbook'])
            try:
                rows=list(wb['Method Parameters'].values)
                values=dict(zip(rows[0],rows[1]))
                self.assertEqual('adaptive_v2',values['neighbor_status_policy'])
                for key,number in parameters.items():
                    if type(number) is int:
                        self.assertEqual(number,values[key]);self.assertIsInstance(values[key],int)
                self.assertIn('STATUS Diagnostics',wb.sheetnames)
                for ws in wb:
                    for table in ws.tables.values():
                        self.assertEqual(len(ws[1]),len(table.tableColumns))
            finally:wb.close()

    def test_cost_setup_window_dedup_failed_requested_and_scope_trial_session(self):
        r = {**record(), "trial_id": "t1"}
        events = source_events() + [ev("STATUS_BURST_STARTED", TARGETS[0], 950),
                                   ev("STATUS_BURST_STARTED", TARGETS[0], 1200, 2)]
        events += copy.deepcopy(events)
        for kind in ("STATUS_BURST_REQUESTED", "STATUS_BURST_FAILED", "DATA_BURST_REQUESTED", "DATA_BURST_FAILED"):
            events.append(ev(kind, TARGETS[0], 1300, 3))
        for field, value in (("session_id", "foreign"), ("trial_id", "foreign"), ("scope", r["scope"]+1)):
            for kind in ("SOS_CREATED", "STATUS_BURST_STARTED", "DATA_RECEIVED"):
                e = ev(kind, SOURCE if kind == "SOS_CREATED" else TARGETS[0], 900 if kind == "STATUS_BURST_STARTED" else 1400, 4)
                e[field] = value
                e["message_key"] = "foreign"
                events.append(e)
        m = summarize_network(events, r)
        self.assertEqual((1, 1, 1, 2, 1, 3), tuple(m[k] for k in ("M", "data_tx", "control_tx", "network_overhead", "setup_control_tx", "setup_plus_window_tx")))

    def test_sample_sd_undefined_not_zero(self):
        self.assertIsNone(descriptive([None])["mean"])
        self.assertIsNone(descriptive([3, None])["sample_sd"])
        self.assertAlmostEqual(2**0.5, descriptive([2, 4])["sample_sd"])

    def test_unequal_trials_failure_valid_invalid_excluded_weighted_vs_trial(self):
        rows = []
        receivers = []
        for trial, valid, result, delay, pairs, dsr, r in [("a", True, "SUCCESS", 100, 5, 100, 10),
                ("b", True, "FAILED_DELIVERY", 1000, 1, 20, 1), ("c", True, "FAILED_DELIVERY", None, 0, 0, 0),
                ("d", False, "INVALID", 99999, 5, 100, 500)]:
            rows.append(dict(trial_id=trial, valid=valid, result=result, method=METHODS[0], scenario="S0_MAIN",
                             M=1, U=pairs, R=r, dsr_percent=dsr, ldr_percent=None if r == 0 else 100*(r-pairs)/r,
                             e2e_mean_ms=delay, successful_pairs=pairs, data_tx=2, control_tx=0,
                             network_overhead=2, setup_control_tx=1, setup_plus_window_tx=3))
            receivers.extend(dict(valid=valid, method=METHODS[0], scenario="S0_MAIN", e2e_latency_ms=delay) for _ in range(pairs))
        rows.append({**rows[0], "method": METHODS[1]})
        summaries, stats = aggregate(rows, receivers, ["S0_MAIN"], METHODS)
        s = summaries[0]
        self.assertEqual((3, 1, 2, 6), (s["valid_trials"], s["invalid_trials"], s["failed_delivery_trials"], s["successful_pairs"]))
        self.assertEqual(250, s["e2e_mean_ms"])
        self.assertEqual(550, s["e2e_mean_ms_trial_mean"])
        self.assertEqual(40, s["dsr_percent"])
        self.assertAlmostEqual(500/11, s["ldr_percent"])
        self.assertEqual(25, s["ldr_percent_trial_mean"])
        self.assertEqual(6, s["network_overhead"])
        self.assertEqual(2, s["network_overhead_trial_mean"])
        self.assertEqual(1, summaries[1]["valid_trials"])
        self.assertIsNone(summaries[2]["e2e_mean_ms"])
        self.assertEqual(2, next(v["defined_trials"] for v in stats if v["method"] == METHODS[0] and v["metric"] == "e2e_mean_ms"))

    def test_phy_config_is_not_s8_proof(self):
        p = phy_evidence({"requested_mode": "coded", "configured_mode": "coded", "ready": True,
                          "coding_actual": "S8", "s8_requirement_accepted": True, "on_air_coding_verified": False})
        self.assertTrue(p["api_configuration_accepted"])
        self.assertEqual("UNKNOWN", p["coding_actual"])
        self.assertEqual("UNVERIFIED", p["on_air_coding_status"])
        self.assertIsNone(phy_evidence({"requested_mode": "coded"})["api_configuration_accepted"])

    def test_frequency_reason_scope_and_archive_duplicates(self):
        r = record()
        e = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1100)
        e["reason"] = "FRESH_MISSING"
        wrong = {**e, "trial_id": "other"}
        result = mechanism_counts([e, e.copy(), wrong], {"t1": r})[0]
        self.assertEqual(1, result["allow_missing"])
        self.assertEqual(0, result["initial_failed"])

    def test_workbook_empty_partial_raw_preserved_charts_numeric_tables(self):
        for partial in (False, True):
            with self.subTest(partial=partial), tempfile.TemporaryDirectory() as d:
                p = Path(d)
                raw = p/"raw"
                raw.mkdir()
                events = source_events()+[received(TARGETS[0])] if partial else []
                (raw/"events.jsonl").write_text("\n".join(json.dumps(e) for e in events))
                r = record()
                if not partial:
                    r.update(observation_started_at_ms=None, observation_ended_at_ms=None, result="INVALID")
                manifest = {"synthetic_data": True, "session_id": "fixture", "trials": {"t1": r}, "neighbor_scenarios": ["S0_MAIN"]}
                output = merge_neighbor(raw, p/"merged", manifest)
                wb = load_workbook(output["workbook"])
                try:
                    self.assertIn("Descriptive Statistics", wb.sheetnames)
                    self.assertIn("Log Validation", wb.sheetnames)
                    self.assertEqual(5, len(output["charts"]))
                    for path in output["charts"]:
                        svg = ET.parse(path)
                        self.assertIn("svg", svg.getroot().tag)
                    delay = Path(next(v for v in output["charts"] if "e2e_mean_ms" in v)).read_text()
                    self.assertIn("Pasangan sukses", delay)
                    self.assertIn("DSR agregat", delay)
                    if not partial:
                        self.assertIn("Tidak ada nilai terdefinisi", delay)
                    self.assertIn("DATA SINTETIS", delay)
                    for ws in wb:
                        for table in ws.tables.values():
                            self.assertEqual(len(ws[1]), len(table.tableColumns))
                    self.assertEqual(events, json.loads((p/"merged/all_events.json").read_text()))
                finally:
                    wb.close()


class NeighborLogValidationTests(unittest.TestCase):
    def complete_archive(self):
        events = source_events() + [received(node) for node in TARGETS]
        for node in (SOURCE, *TARGETS):
            events.extend([ev("TRIAL_WINDOW_STARTED", node, 800), ev("TRIAL_WINDOW_ENDED", node, 2000)])
            for sequence, event in enumerate(sorted((e for e in events if e["node_id"] == node),
                                                     key=lambda e: e["timestamp_ms"]), 10):
                event["id" if node == SOURCE else "event_sequence"] = sequence
        return events

    def checks(self, events, r=None):
        result = validate_logs(events, {"session_id": "fixture", "synthetic_data": True, "trials": {"t1": r or record()}})
        self.assertTrue(all(v["evidence_origin"] == "SYNTHETIC_FIXTURE_NOT_HARDWARE" for v in result))
        return {v["check"]: v["result"] for v in result}

    def test_missing_evidence_never_pass(self):
        self.assertEqual({"INCONCLUSIVE"}, set(self.checks([]).values()))

    def test_mixed_named_neighbors_allow_not_suppress(self):
        e = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1100)
        e.update(reason="FRESH_MISSING", neighbor_knowledge={str(stable_id(TARGETS[2])): "HAVE", str(stable_id(TARGETS[3])): "MISSING"})
        self.assertEqual("PASS", self.checks([e])["A_C_HAVE_D_MISSING"])
        e["event_type"] = "NEIGHBOR_TX_SUPPRESSED"
        self.assertEqual("FAIL", self.checks([e])["A_C_HAVE_D_MISSING"])

    def test_all_have_requires_status_still_running(self):
        e = ev("NEIGHBOR_TX_SUPPRESSED", TARGETS[0], 1100)
        e.update(reason="ALL_OBSERVED_HAVE", have_count=2, missing_count=0, unknown_count=0)
        self.assertEqual("INCONCLUSIVE", self.checks([e])["ALL_HAVE_SUPPRESSION_STATUS_CONTINUES"])
        self.assertEqual("PASS", self.checks([e, ev("STATUS_BURST_STARTED", TARGETS[0], 1200)])["ALL_HAVE_SUPPRESSION_STATUS_CONTINUES"])
        self.assertEqual("PASS", self.checks([e])["NEIGHBOR_DECISION_CONSISTENCY"])
        e["missing_count"] = 1
        self.assertEqual("FAIL", self.checks([e])["NEIGHBOR_DECISION_CONSISTENCY"])

    def test_failure_requires_pending_scanner_requires_actual_rx(self):
        e = ev("INITIAL_FORWARD_FAILED", TARGETS[0], 1100)
        self.assertEqual("INCONCLUSIVE", self.checks([e])["INITIAL_FAILURE_REMAINS_PENDING"])
        e["first_forward_pending"] = True
        self.assertEqual("PASS", self.checks([e])["INITIAL_FAILURE_REMAINS_PENDING"])
        e["first_forward_pending"] = False
        self.assertEqual("FAIL", self.checks([e])["INITIAL_FAILURE_REMAINS_PENDING"])
        end = ev("DATA_BURST_ENDED", TARGETS[0], 1100)
        end["transport_burst_id"] = "burst"
        check = {**end, "event_type": "SCANNER_RECOVERY_CHECK", "burst_outcome": "ENDED", "scanner_registered": True, "rx_enabled": True}
        self.assertEqual("INCONCLUSIVE", self.checks([end, check])["SCANNER_RECOVERY_ENDED"])
        self.assertEqual("PASS", self.checks([end, check, ev("STATUS_RECEIVED", TARGETS[0], 1200)])["SCANNER_RECOVERY_ENDED"])
        check["scanner_registered"] = False
        self.assertEqual("FAIL", self.checks([end, check])["SCANNER_RECOVERY_ENDED"])

    def test_s1_s2_need_actual_off_on_not_command_requested(self):
        for scenario, node, kind in [("S1_DELAYED_RX", TARGETS[3], "RX_PARTICIPATION_CHANGED"), ("S2_LATE_JOIN", TARGETS[4], "NODE_PARTICIPATION_CHANGED")]:
            r = {**record(), "hypothesis": scenario}
            off = ev(kind, node, 1100)
            on = ev(kind, node, 1200)
            off.update(confirmed_enabled=False, rx_enabled=False, tx_enabled=False, scanner_registered=False)
            on.update(confirmed_enabled=True, rx_enabled=True, tx_enabled=True)
            self.assertEqual("PASS", self.checks([off, on], r)["ACTUAL_PARTICIPATION_OFF_ON"])
            self.assertEqual("INCONCLUSIVE", self.checks([on], r)["ACTUAL_PARTICIPATION_OFF_ON"])
            off["rx_enabled"] = True
            self.assertEqual("FAIL", self.checks([off, on], r)["ACTUAL_PARTICIPATION_OFF_ON"])

    def test_gaps_fail_metric_partial_inconclusive_and_wrong_metrics_fail(self):
        events = source_events()
        gap1 = {**ev("STATUS_RECEIVED", TARGETS[0], 1100), "event_sequence": 1}
        gap2 = {**gap1, "timestamp_ms": 1200, "event_sequence": 3}
        self.assertEqual("FAIL", self.checks(events+[gap1, gap2])["LOG_COMPLETENESS"])
        r = record()
        r["evidence"] = summarize_network(events, r)
        self.assertEqual("INCONCLUSIVE", self.checks(events, r)["METRICS_RECOMPUTE_MATCH"])
        r["evidence"]["R"] = 100
        self.assertEqual("FAIL", self.checks(events, r)["METRICS_RECOMPUTE_MATCH"])

    def test_complete_marker_archive_and_correct_metrics_pass_only_observed_log(self):
        events = self.complete_archive()
        r = record()
        r["evidence"] = summarize_network(events, r)
        result = self.checks(events, r)
        self.assertEqual("PASS", result["LOG_COMPLETENESS"])
        self.assertEqual("PASS", result["METRICS_RECOMPUTE_MATCH"])
        self.assertEqual("INCONCLUSIVE", result["INITIAL_FAILURE_REMAINS_PENDING"])

    def test_markers_without_sequences_cannot_certify_logs_or_metrics(self):
        events = self.complete_archive()
        for e in events:
            e.pop("event_sequence", None)
        r = record()
        r["evidence"] = summarize_network(events, r)
        result = self.checks(events, r)
        self.assertEqual("INCONCLUSIVE", result["LOG_COMPLETENESS"])
        self.assertEqual("INCONCLUSIVE", result["METRICS_RECOMPUTE_MATCH"])

    def test_each_esp_event_requires_valid_sequence(self):
        for value in (None, True, "10", 0, -1):
            with self.subTest(value=value):
                events = self.complete_archive()
                next(e for e in events if e["node_id"] == TARGETS[2])["event_sequence"] = value
                self.assertEqual("INCONCLUSIVE", self.checks(events)["LOG_COMPLETENESS"])

    def test_foreign_sequences_cannot_repair_missing_node_evidence(self):
        for field, value in (("session_id", "other"), ("trial_id", "other"), ("scope", 999)):
            with self.subTest(field=field):
                events = self.complete_archive()
                original = next(e for e in events if e["node_id"] == TARGETS[2])
                foreign = {**original, field: value}
                original.pop("event_sequence")
                self.assertEqual("INCONCLUSIVE", self.checks(events+[foreign])["LOG_COMPLETENESS"])

    def test_replayed_archive_not_conflicting_sequence_is_allowed(self):
        events = self.complete_archive()
        self.assertEqual("PASS", self.checks(events+copy.deepcopy(events))["LOG_COMPLETENESS"])
        conflict = {**next(e for e in events if e["node_id"] == TARGETS[0]), "event_type": "OTHER"}
        self.assertEqual("FAIL", self.checks(events+[conflict])["LOG_COMPLETENESS"])

    def test_android_global_db_ids_need_not_be_contiguous(self):
        events = self.complete_archive()
        for event in events:
            if event["node_id"] == SOURCE:
                event["id"] *= 3
                self.assertNotIn("event_sequence", event)
        self.assertEqual("PASS", self.checks(events)["LOG_COMPLETENESS"])

    def test_validator_cli_inconclusive_exit_never_opens_devices(self):
        from contextlib import redirect_stdout
        from io import StringIO
        from unittest.mock import patch
        from resqmesh_controller.cli import main
        with tempfile.TemporaryDirectory() as d:
            path = Path(d)
            (path/"raw").mkdir()
            (path/"manifest.json").write_text(json.dumps({"session_id": "fixture", "synthetic_data": True, "trials": {"t1": record()}}))
            args = ["run.py", "validate-neighbor", "--input", str(path/"raw"), "--output", str(path/"validation"), "--manifest", str(path/"manifest.json")]
            with patch("sys.argv", args), patch("resqmesh_controller.cli.build_nodes", side_effect=AssertionError("Must not open devices")), redirect_stdout(StringIO()):
                self.assertEqual(3, main())
            rows = json.loads((path/"validation/log_validation.json").read_text())
            self.assertEqual({"INCONCLUSIVE"}, {r["result"] for r in rows})

    def test_empty_status_repair_requires_same_burst_and_discovery(self):
        empty = ev("STATUS_RECEIVED", TARGETS[0], 1100)
        empty.update(inventory_count=0, snapshot_complete=True, transport_burst_id="empty")
        repair = {**empty, "timestamp_ms": 1101, "event_type": "REPAIR_NEEDED"}
        discovered = {**empty, "event_type": "NEIGHBOR_DISCOVERED"}
        self.assertEqual("INCONCLUSIVE", self.checks([empty, repair])["EMPTY_STATUS_DISCOVERY_REPAIR"])
        self.assertEqual("PASS", self.checks([empty, repair, discovered])["EMPTY_STATUS_DISCOVERY_REPAIR"])
        repair["transport_burst_id"] = "different"
        self.assertEqual("INCONCLUSIVE", self.checks([empty, repair, discovered])["EMPTY_STATUS_DISCOVERY_REPAIR"])

    def test_expiry_unknown_but_refreshed_have_not_failure(self):
        expired = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1200)
        decision.update(reason="UNKNOWN_OR_NO_NEIGHBORS", neighbor_knowledge={str(expired["transmitter_id"]): "UNKNOWN"})
        self.assertEqual("PASS", self.checks([expired, decision])["EXPIRED_STATUS_UNKNOWN"])
        decision.update(reason="ALL_OBSERVED_HAVE", event_type="NEIGHBOR_TX_SUPPRESSED", neighbor_knowledge={str(expired["transmitter_id"]): "HAVE"})
        self.assertEqual("FAIL", self.checks([expired, decision])["EXPIRED_STATUS_UNKNOWN"])
        refresh = {**expired, "event_type": "NEIGHBOR_STATUS_UPDATED", "timestamp_ms": 1150}
        self.assertEqual("INCONCLUSIVE", self.checks([expired, refresh, decision])["EXPIRED_STATUS_UNKNOWN"])

    def test_expiry_have_or_missing_fails_regardless_of_allow_reason(self):
        expired = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        for state in ("HAVE", "MISSING"):
            with self.subTest(state=state):
                decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1200)
                decision.update(reason="UNKNOWN_OR_NO_NEIGHBORS", neighbor_knowledge={str(expired["transmitter_id"]): state})
                self.assertEqual("FAIL", self.checks([expired, decision])["EXPIRED_STATUS_UNKNOWN"])

    def test_refresh_of_one_peer_cannot_mask_stale_other_peer(self):
        a = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        b = {**a, "transmitter_id": stable_id(TARGETS[3])}
        refresh = {**a, "event_type": "NEIGHBOR_STATUS_UPDATED", "timestamp_ms": 1150}
        decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1200)
        decision.update(reason="UNKNOWN_OR_NO_NEIGHBORS", neighbor_knowledge={
            str(a["transmitter_id"]): "UNKNOWN", str(b["transmitter_id"]): "HAVE"})
        self.assertEqual("FAIL", self.checks([a, b, refresh, decision])["EXPIRED_STATUS_UNKNOWN"])

    def test_expiry_does_not_borrow_other_observer_unknown_or_refresh(self):
        a = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        b = {**a, "node_id": TARGETS[1]}
        refresh = {**b, "timestamp_ms": 1150, "event_type": "NEIGHBOR_STATUS_UPDATED"}
        unknown = ev("NEIGHBOR_TX_ALLOWED", TARGETS[1], 1200)
        unknown["neighbor_knowledge"] = {str(a["transmitter_id"]): "UNKNOWN"}
        have = {**unknown, "node_id": TARGETS[0], "neighbor_knowledge": {str(a["transmitter_id"]): "HAVE"}}
        self.assertEqual("FAIL", self.checks([a, b, refresh, unknown, have])["EXPIRED_STATUS_UNKNOWN"])
        self.assertEqual("INCONCLUSIVE", self.checks([a, unknown])["EXPIRED_STATUS_UNKNOWN"])

    def test_each_expiry_episode_requires_its_own_unknown_evidence(self):
        expired = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1200)
        decision["neighbor_knowledge"] = {str(expired["transmitter_id"]): "UNKNOWN"}
        refresh = {**expired, "event_type": "NEIGHBOR_STATUS_UPDATED", "timestamp_ms": 1300}
        second = {**expired, "timestamp_ms": 1400}
        self.assertEqual("INCONCLUSIVE", self.checks([expired, decision, refresh, second])["EXPIRED_STATUS_UNKNOWN"])
        stale = {**decision, "timestamp_ms": 1500, "neighbor_knowledge": {str(expired["transmitter_id"]): "MISSING"}}
        self.assertEqual("FAIL", self.checks([second, stale, decision, refresh, expired])["EXPIRED_STATUS_UNKNOWN"])
        stale["neighbor_knowledge"] = decision["neighbor_knowledge"]
        self.assertEqual("PASS", self.checks([second, stale, decision, refresh, expired])["EXPIRED_STATUS_UNKNOWN"])

    def test_same_millisecond_expiry_uses_node_event_order(self):
        expired = {**ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2])), "event_sequence": 7}
        decision = {**ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1100), "event_sequence": 8,
                    "neighbor_knowledge": {str(expired["transmitter_id"]): "UNKNOWN"}}
        self.assertEqual("PASS", self.checks([decision, expired])["EXPIRED_STATUS_UNKNOWN"])

    def test_missing_peer_identity_or_snapshot_is_inconclusive(self):
        expired = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1100, tx=stable_id(TARGETS[2]))
        decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1200)
        self.assertEqual("INCONCLUSIVE", self.checks([expired, decision])["EXPIRED_STATUS_UNKNOWN"])
        expired.pop("transmitter_id")
        decision["neighbor_knowledge"] = {"None": "UNKNOWN"}
        self.assertEqual("INCONCLUSIVE", self.checks([expired, decision])["EXPIRED_STATUS_UNKNOWN"])

    def test_expiry_android_db_order_and_same_peer_refresh(self):
        expired = {**ev("NEIGHBOR_STATUS_EXPIRED", SOURCE, 1100, tx=stable_id(TARGETS[0])), "id": 20}
        decision = {**ev("NEIGHBOR_TX_ALLOWED", SOURCE, 1100), "id": 21,
                    "neighbor_knowledge": {str(expired["transmitter_id"]): "UNKNOWN"}}
        refresh = {**expired, "id": 22, "event_type": "NEIGHBOR_STATUS_UPDATED"}
        have = {**decision, "id": 23, "neighbor_knowledge": {str(expired["transmitter_id"]): "HAVE"}}
        self.assertEqual("PASS", self.checks([have, refresh, decision, expired])["EXPIRED_STATUS_UNKNOWN"])

    def test_validator_and_workbook_keep_negative_evidence_without_touching_raw(self):
        from contextlib import redirect_stdout
        from io import StringIO
        from unittest.mock import patch
        from resqmesh_controller.cli import main
        for contradiction in (False, True):
            with self.subTest(contradiction=contradiction), tempfile.TemporaryDirectory() as d:
                path = Path(d)
                raw = path/"raw"
                raw.mkdir()
                events = self.complete_archive()
                for event in events:
                    event.pop("event_sequence", None)
                if contradiction:
                    expired = ev("NEIGHBOR_STATUS_EXPIRED", TARGETS[0], 1300, tx=stable_id(TARGETS[2]))
                    decision = ev("NEIGHBOR_TX_ALLOWED", TARGETS[0], 1400)
                    decision.update(reason="UNKNOWN_OR_NO_NEIGHBORS",
                                    neighbor_knowledge={str(expired["transmitter_id"]): "HAVE"})
                    events.extend([expired, decision])
                raw_file = raw/"events.jsonl"
                raw_file.write_text("\n".join(json.dumps(e) for e in events), encoding="utf-8")
                raw_bytes = raw_file.read_bytes()
                r = record()
                r["evidence"] = summarize_network(events, r)
                manifest = {"session_id": "fixture", "synthetic_data": True, "trials": {"t1": r},
                            "neighbor_scenarios": ["S0_MAIN"]}
                manifest_file = path/"manifest.json"
                manifest_file.write_text(json.dumps(manifest), encoding="utf-8")
                args = ["run.py", "validate-neighbor", "--input", str(raw), "--output", str(path/"validation"),
                        "--manifest", str(manifest_file)]
                with patch("sys.argv", args), patch("resqmesh_controller.cli.build_nodes", side_effect=AssertionError("No hardware")), redirect_stdout(StringIO()):
                    self.assertEqual(2 if contradiction else 3, main())
                output = merge_neighbor(raw, path/"merged", manifest)
                wb = load_workbook(output["workbook"])
                try:
                    ws = wb["Log Validation"]
                    rows = [dict(zip([c.value for c in ws[1]], row)) for row in ws.iter_rows(min_row=2, values_only=True)]
                    checks = {row["check"]: row for row in rows}
                    self.assertEqual("INCONCLUSIVE", checks["LOG_COMPLETENESS"]["result"])
                    self.assertEqual("INCONCLUSIVE", checks["METRICS_RECOMPUTE_MATCH"]["result"])
                    if contradiction:
                        self.assertEqual("FAIL", checks["EXPIRED_STATUS_UNKNOWN"]["result"])
                        episode = json.loads(checks["EXPIRED_STATUS_UNKNOWN"]["evidence"])["expiry_episodes"][0]
                        self.assertEqual((TARGETS[0], expired["transmitter_id"], ["HAVE"]),
                                         (episode["node_id"], episode["transmitter_id"], episode["observed_states"]))
                finally:
                    wb.close()
                self.assertEqual(raw_bytes, raw_file.read_bytes())
                self.assertEqual(events, json.loads((path/"merged/all_events.json").read_text()))


if __name__ == "__main__":
    unittest.main()
