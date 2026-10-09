import copy
import json
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from resqmesh_controller.cli import main
from resqmesh_controller.config import ConfigError, validate_config, research_fingerprint, smoke_matches_config
from resqmesh_controller.controller import TrialSpec, build_plan
from resqmesh_controller.devices import DeviceError, write_jsonl
from resqmesh_controller.esp_only_validation import scenario_checks
from resqmesh_controller.mpl_config import DEFAULTS, METHODS
from resqmesh_controller.neighbor_experiment import NeighborExperimentController, merge_neighbor, summarize_network
from resqmesh_controller.neighbor_testbed import (
    RECOVERY_ESP, RECOVERY_ANDROID, RECOVERY_SCENARIOS, RECOVERY_VERSION, ESP_SCENARIOS,
    scenario_parameters, testbed as resolve_testbed,
)
from resqmesh_controller.neighbor_validation import validate_logs
from test_esp_only_experiment import Clock, EspFake, KEY, STATE, ROOT, config as old_config, mechanism_fixture


def config(android=False, stage="smoke"):
    kind = "hp" if android else "esp-only"
    cfg = json.loads((ROOT/f"tools/experiment_controller/config.mpl.recovery-180.{kind}.{stage}.example.json").read_text())
    cfg.update(firmware_build_id="0123456789ab", session_id="recovery-fixture")
    if android:
        cfg["android_build_id"] = cfg["firmware_build_id"]
        cfg["nodes"][0]["serial"] = "phone"
    for i, node in enumerate(cfg["nodes"], 8):
        if node["transport"] == "serial":
            node["port"] = f"COM{i}"
    return cfg


def fixture(android=False):
    manifest, record, events = mechanism_fixture()
    cfg = config(android)
    profile = cfg["testbed_profile"]
    metadata = resolve_testbed(cfg).metadata()
    manifest.update(testbed_profile=profile, **metadata, scenario_parameters=RECOVERY_SCENARIOS,
                    firmware_build_id=cfg["firmware_build_id"], android_build_id=cfg.get("android_build_id"),
                    recovery_measurement_version=RECOVERY_VERSION,
                    mpl_parameters={**DEFAULTS, "mpl_data_expirations": 3})
    record.update(testbed_profile=profile, **metadata, recovery_measurement_version=RECOVERY_VERSION,
                  firmware_build_id=cfg["firmware_build_id"], android_build_id=cfg.get("android_build_id"),
                  observation_ended_at_ms=190000, session_id=cfg["session_id"],
                  scenario_parameters=RECOVERY_SCENARIOS["S2_POST_DATA_STOP"])
    manifest["session_id"] = cfg["session_id"]
    for e in events:
        if android and e["node_id"] == "esp-r1b":
            e["node_id"] = "android-source"
        e["session_id"] = cfg["session_id"]
        relative = e["timestamp_ms"]-10000
        relative = 56000 if e["event_type"] == "MPL_TIMER_STOPPED" else relative-210000 if relative >= 300000 else relative
        e.update(timestamp_ms=10000+relative, monotonic_ms=relative,
                 clock_domain="esp_boot_millis", local_boot_id=2)
        if e["event_type"] == "MPL_TIMER_STOPPED":
            e["expiration_count"] = 3
    record["pre_activation_status"]["captured_at_ms"] = 99990
    record["pre_activation_status"]["response"]["session_id"] = cfg["session_id"]
    record["participation"] = [{"node": "esp-destination", "enabled": True,
                                "response": {"confirmed_enabled": True, "local_boot_id": 2}}]
    return manifest, record, events


class RecoveryFake(EspFake):
    def __init__(self, node, cfg, clock):
        super().__init__(node, cfg, clock)
        self.transport = node["transport"]

    def host_clock_offset_ms(self):
        return 0

    def command(self, name, args):
        result = super().command(name, args)
        if name in {"readiness", "get_status"}:
            result["mpl_parameters"] = self.cfg["mpl_parameters"]
            if self.transport == "adb":
                result.update(bluetooth=True, permissions={"scan": True, "advertise": True})
                result["radio"]["maximum_advertising_data_length"] = 255
        if name == "trigger_sos" and self.transport == "adb":
            result["message_key"] = KEY
        return result


class Recovery180Test(unittest.TestCase):
    def test_examples_plan_blocks_and_old_defaults_unchanged(self):
        for android, stage, count in ((False, "smoke", 9), (True, "smoke", 9), (True, "pilot", 27), (True, "main", 135)):
            cfg = config(android, stage)
            validate_config(cfg)
            plan = build_plan(cfg)
            self.assertEqual(count, len(plan))
            self.assertEqual(plan, build_plan(cfg))
            for block in set(s.block for s in plan):
                self.assertEqual(9, len({(s.mode, s.hypothesis) for s in plan if s.block == block}))
            self.assertEqual(5 if android else 4, len(resolve_testbed(cfg).targets))
            self.assertTrue(all(scenario_parameters(cfg, s)["observation_window_seconds"] == 180 for s in cfg["scenarios"]))
        validate_config(old_config())
        self.assertEqual(5, DEFAULTS["mpl_data_expirations"])
        self.assertEqual(420, ESP_SCENARIOS["S2_POST_DATA_STOP"]["observation_window_seconds"])

    def test_config_rejects_parameter_schedule_role_and_port_changes(self):
        for mutate in (
            lambda c: c["mpl_parameters"].update(mpl_data_expirations=5),
            lambda c: c["mpl_parameters"].update(mpl_control_expirations=3),
            lambda c: c["scenario_parameters"]["S2_POST_DATA_STOP"].update(activate_at_seconds=120),
            lambda c: c["nodes"][-1].update(role="DESTINATION"),
            lambda c: c["nodes"][-1].update(port="COM_X"),
        ):
            cfg = config()
            mutate(cfg)
            with self.assertRaises(ConfigError):
                validate_config(cfg)

    def test_source_transport_activation_duration_and_metrics(self):
        for android in (False, True):
            for scenario in RECOVERY_SCENARIOS:
                with self.subTest(android=android, scenario=scenario), tempfile.TemporaryDirectory() as d:
                    cfg, clock = config(android), Clock()
                    nodes = [RecoveryFake(n, cfg, clock) for n in cfg["nodes"]]
                    ctrl = NeighborExperimentController(cfg, nodes, Path(d), sleep=clock.sleep)
                    with patch.multiple("resqmesh_controller.controller.time", monotonic=clock.monotonic, time_ns=clock.time_ns):
                        record = ctrl.run_trial(TrialSpec("basic_flooding", scenario, 1))
                    self.assertEqual("SUCCESS", record["result"], record.get("invalid_reasons"))
                    self.assertTrue(record["reset_verified"])
                    self.assertEqual(180000, record["observation_ended_at_ms"]-record["observation_started_at_ms"])
                    self.assertEqual(100, record["evidence"]["dsr_percent"])
                    eligible = record["evidence"]["recovery_eligible_targets"]
                    self.assertEqual(len(RECOVERY_SCENARIOS[scenario]["inactive_node_ids"]), eligible)
                    self.assertEqual(eligible, record["evidence"]["recovery_defined_targets"])
                    for row in record["evidence"]["recovery_receivers"]:
                        if row["recovery_eligible"]:
                            self.assertEqual(1000, row["recovery_delay_ms"])
                            due = record["observation_started_at_ms"] + RECOVERY_SCENARIOS[scenario]["activate_at_seconds"]*1000
                            self.assertGreaterEqual(row["recovery_on_confirmed_at_ms"], due)
                    stored = [n for n in nodes if any(cmd == "store_neighbor_metrics" for cmd, _ in n.commands)]
                    self.assertEqual(1 if android else 0, len(stored))
                    self.assertEqual(len(nodes), len(list((Path(d)/"raw").rglob("*.jsonl"))))

    def test_recovery_clock_identity_duplicates_and_missing_rx(self):
        _, record, events = fixture()
        def row(samples):
            return next(r for r in summarize_network(samples, record)["recovery_receivers"] if r["receiver"] == "esp-destination")
        self.assertEqual(7020, row(events)["recovery_delay_ms"])
        self.assertEqual(7020, row(events+[copy.deepcopy(events[-1])])["recovery_delay_ms"])
        for field, value in (("message_key", "wrong"), ("scope", 999), ("trial_id", "wrong"),
                             ("session_id", "wrong"), ("transmitter_id", 999)):
            changed = [dict(e, **{field: value}) if e["event_type"] == "DATA_RECEIVED" else e for e in events]
            result = row(changed)
            self.assertEqual("NOT_RECOVERED_WITHIN_WINDOW", result["recovery_status"])
            self.assertIsNone(result["recovery_delay_ms"])
        for field, value in (("monotonic_ms", None), ("clock_domain", "android_elapsed_realtime"), ("local_boot_id", 3)):
            changed = [dict(e, **{field: value}) if e["event_type"] == "DATA_RECEIVED" else e for e in events]
            self.assertEqual("TIMING_UNVERIFIED", row(changed)["recovery_status"])
        result = row([e for e in events if e["event_type"] != "DATA_RECEIVED"])
        self.assertEqual("NOT_RECOVERED_WITHIN_WINDOW", result["recovery_status"])
        self.assertEqual(90000, result["recovery_followup_ms"])
        self.assertEqual(90000, result["recovery_on_monotonic_ms"])
        early = [dict(e, timestamp_ms=99999, monotonic_ms=89999) if e["event_type"] == "DATA_RECEIVED" else e for e in events]
        self.assertEqual("RX_BEFORE_ON", row(early)["recovery_unavailable_reason"])
        outside = [dict(e, timestamp_ms=record["observation_ended_at_ms"]) if e["event_type"] == "DATA_RECEIVED" else e for e in events]
        self.assertEqual("NOT_RECOVERED_WITHIN_WINDOW", row(outside)["recovery_status"])

    def test_post_stop_validation_for_both_sources_and_pre_on_restart(self):
        for android in (False, True):
            manifest, record, events = fixture(android)
            def verdict(samples):
                return next(c["result"] for c in scenario_checks(samples, record, manifest) if c["check"] == "RECOVERY_POST_STOP_REPAIR")
            self.assertEqual("PASS", verdict(events))
            if android:
                unscoped = [dict(e, scope=None) if e["event_type"] == "SOS_CREATED" else e for e in events]
                self.assertEqual("PASS", verdict(unscoped))
            for kind in ("MPL_TIMER_STOPPED", "MPL_REPAIR_PENDING", "MPL_NATIVE_STARTED", "DATA_RECEIVED"):
                self.assertEqual("INCONCLUSIVE", verdict([e for e in events if e["event_type"] != kind]))
            restart = {**events[4], "event_type": "MPL_TIMER_RESTARTED", "event_sequence": 6, "timestamp_ms": 80000}
            self.assertEqual("INCONCLUSIVE", verdict(events+[restart]))
            wrong = [dict(e, timer_kind="control") if e["event_type"] == "MPL_TIMER_STOPPED" else e for e in events]
            self.assertEqual("INCONCLUSIVE", verdict(wrong))
            wrong = [dict(e, buffer_retained=False) if e["event_type"] == "MPL_TIMER_STOPPED" else e for e in events]
            self.assertEqual("FAIL", verdict(wrong))

    def test_failed_delivery_stays_valid_and_empty_delay_is_preserved(self):
        with tempfile.TemporaryDirectory() as d:
            cfg, clock = config(True), Clock()
            nodes = [RecoveryFake(n, cfg, clock) for n in cfg["nodes"]]
            nodes[-1].missing_delivery = True
            ctrl = NeighborExperimentController(cfg, nodes, Path(d), sleep=clock.sleep)
            with patch.multiple("resqmesh_controller.controller.time", monotonic=clock.monotonic, time_ns=clock.time_ns):
                result = ctrl.run_trial(TrialSpec("trickle", "S2_POST_DATA_STOP", 1))
            self.assertEqual("FAILED_DELIVERY", result["result"])
            self.assertEqual(80, result["evidence"]["dsr_percent"])
            self.assertIsNone(result["evidence"]["recovery_mean_ms"])
            self.assertEqual(1, result["evidence"]["recovery_unreceived_targets"])
            self.assertEqual(1, result["evidence"]["recovery_eligible_targets"])

    def test_new_recovery_does_not_change_four_metric_results(self):
        manifest, record, events = fixture()
        new = summarize_network(events, record)
        old = summarize_network(events, {**record, "testbed_profile": "esp_only_five_v1"})
        for field in ("M", "N", "U", "R", "dsr_percent", "ldr_percent", "e2e_mean_ms",
                      "data_tx", "control_tx", "network_overhead", "setup_control_tx", "setup_plus_window_tx"):
            self.assertEqual(old[field], new[field])

    def test_smoke_gate_allows_valid_delivery_failure_but_requires_proofs(self):
        for android in (False, True):
            with tempfile.TemporaryDirectory() as d:
                cfg = config(android)
                ctrl = NeighborExperimentController(cfg, [], Path(d))
                checks = []
                for spec in build_plan(cfg):
                    record = {"mode": spec.mode, "hypothesis": spec.hypothesis,
                              "result": "FAILED_DELIVERY" if spec.mode == "trickle" else "SUCCESS"}
                    ctrl.manifest["trials"][spec.trial_id] = record
                    names = ["LOG_COMPLETENESS", "METRICS_RECOMPUTE_MATCH", "RECOVERY_SCENARIO_PARTICIPATION", "RECOVERY_METRICS_RECOMPUTE_MATCH"]
                    if spec.mode == "trickle_mpl":
                        names += ["MPL_LISTEN_ONLY_AND_BOUNDS", "MPL_C_K_DECISIONS", "MPL_PARAMETERS_MATCH", "MPL_OPPORTUNITY_ACCOUNTING", "MPL_EXPIRATION_AND_DOUBLING", "ANDROID_CONTROL_RX"]
                        if spec.hypothesis == "S2_POST_DATA_STOP":
                            names.append("RECOVERY_POST_STOP_REPAIR")
                    checks += [{"trial_id": spec.trial_id, "check": name, "result": "PASS"} for name in names]
                with patch("resqmesh_controller.neighbor_validation.validate_logs", return_value=checks):
                    report = ctrl.smoke_report()
                self.assertTrue(report["passed"])
                self.assertFalse(report["delivery_passed"])
                for name in ("RECOVERY_POST_STOP_REPAIR", "LOG_COMPLETENESS", "ANDROID_CONTROL_RX"):
                    if not android and name == "ANDROID_CONTROL_RX":
                        continue
                    lacking = [dict(c, result="INCONCLUSIVE") if c["check"] == name else c for c in checks]
                    with patch("resqmesh_controller.neighbor_validation.validate_logs", return_value=lacking):
                        self.assertFalse(ctrl.smoke_report()["passed"])

    def test_incoming_android_control_requires_actual_classification(self):
        manifest, record, events = fixture(True)
        def result(values):
            return next(c["result"] for c in validate_logs(values, manifest) if c["check"] == "ANDROID_CONTROL_RX")
        self.assertEqual("INCONCLUSIVE", result(events))
        incoming = {**events[6], "node_id": "android-source", "peer_id": resolve_testbed(manifest).adjacency("android-source")[0]}
        self.assertEqual("PASS", result(events+[incoming]))
        self.assertEqual("INCONCLUSIVE", result(events+[dict(incoming, scope=123)]))

    def test_recovery_validation_recomputes_evidence_not_only_labels(self):
        manifest, record, events = fixture()
        record["evidence"] = summarize_network(events, record)
        record["evidence"]["recovery_mean_ms"] = 0
        checks = validate_logs(events, manifest)
        self.assertEqual("FAIL", next(c["result"] for c in checks if c["check"] == "RECOVERY_METRICS_RECOMPUTE_MATCH"))

    def test_workbook_raw_numeric_tables_statistics_and_cumulative(self):
        for android in (False, True):
            manifest, record, events = fixture(android)
            record["evidence"] = summarize_network(events, record)
            with tempfile.TemporaryDirectory() as d:
                raw, out = Path(d)/"raw", Path(d)/"merged"
                write_jsonl(raw/"events.jsonl", events)
                result = merge_neighbor(raw, out, manifest)
                wb = load_workbook(result["workbook"])
                self.addCleanup(wb.close)
                self.assertEqual(["Mulai Di Sini", "Desain Pengujian", "Ringkasan Metode", "Hasil Per Trial", "Pemulihan Node"], wb.sheetnames[:5])
                self.assertEqual((5 if android else 4)+1, wb["Pemulihan Node"].max_row)
                self.assertIn("Grafik Pemulihan", wb.sheetnames)
                for sheet in wb:
                    for table in sheet.tables.values():
                        self.assertEqual(len(sheet[1]), len(table.tableColumns))
                        self.assertEqual(len(table.tableColumns), len({c.name for c in table.tableColumns}))
                        self.assertEqual(table.ref, table.autoFilter.ref)
                with zipfile.ZipFile(result["workbook"]) as archive:
                    self.assertIsNone(archive.testzip())
                    for name in archive.namelist():
                        if name.endswith(".xml"):
                            ET.fromstring(archive.read(name))
                self.assertEqual(events, json.loads((out/"all_events.json").read_text()))
                rows = json.loads((out/"network_metrics.json").read_text())
                self.assertEqual(5 if android else 4, rows[0]["N"])
                self.assertEqual(7020, rows[0]["recovery_mean_ms"])
                curves = json.loads((out/"cumulative_metrics.json").read_text())
                self.assertEqual(180, curves[-1]["elapsed_seconds"])
                self.assertEqual(rows[0]["network_overhead"], curves[-1]["network_overhead"])

    def test_cleanup_failure_and_interrupt_archive_stop_batch(self):
        for android in (False, True):
            for failure in (DeviceError("disconnected"), KeyboardInterrupt()):
                with tempfile.TemporaryDirectory() as d:
                    cfg, clock = config(android), Clock()
                    nodes = [RecoveryFake(n, cfg, clock) for n in cfg["nodes"]]
                    nodes[0].failures["trigger_sos"] = failure
                    ctrl = NeighborExperimentController(cfg, nodes, Path(d), sleep=clock.sleep)
                    spec = TrialSpec("basic_flooding", "S0_STABLE", 1)
                    with patch.multiple("resqmesh_controller.controller.time", monotonic=clock.monotonic, time_ns=clock.time_ns):
                        if isinstance(failure, KeyboardInterrupt):
                            with self.assertRaises(KeyboardInterrupt):
                                ctrl.run_trial(spec)
                        else:
                            ctrl.run_trial(spec)
                    record = ctrl.manifest["trials"][spec.trial_id]
                    self.assertEqual("INVALID", record["result"])
                    self.assertTrue(record["reset_verified"])
                    self.assertEqual(len(nodes), len(list((Path(d)/"raw").rglob("*.jsonl"))))
            with tempfile.TemporaryDirectory() as d:
                cfg, clock = config(android), Clock()
                nodes = [RecoveryFake(n, cfg, clock) for n in cfg["nodes"]]
                nodes[-1].failures["reset_trial"] = DeviceError("stop unconfirmed")
                ctrl = NeighborExperimentController(cfg, nodes, Path(d), sleep=clock.sleep)
                with patch.multiple("resqmesh_controller.controller.time", monotonic=clock.monotonic, time_ns=clock.time_ns):
                    with self.assertRaisesRegex(DeviceError, "CLEANUP_UNCONFIRMED"):
                        ctrl.run()
                self.assertEqual(1, len(ctrl.manifest["trials"]))

    def test_export_rejects_mixed_build_profile_and_fingerprint(self):
        for field, value in (("testbed_profile", "esp_only_five_v1"), ("config_fingerprint", "wrong"),
                             ("firmware_build_id", "ffffffffffff"), ("recovery_measurement_version", "wrong")):
            manifest, record, events = fixture()
            record[field] = value
            with tempfile.TemporaryDirectory() as d:
                raw, out = Path(d)/"raw", Path(d)/"merged"
                write_jsonl(raw/"events.jsonl", events)
                with self.subTest(field=field), self.assertRaises(ConfigError):
                    merge_neighbor(raw, out, manifest)

    def test_separate_smoke_report_checks_profile_build_and_fingerprint(self):
        cfg = config(True)
        report = {"passed": True, "testbed_profile": cfg["testbed_profile"], "firmware_build_id": cfg["firmware_build_id"],
                  "android_build_id": cfg["android_build_id"], "config_fingerprint": research_fingerprint(cfg),
                  "core_checks_passed": True, "mechanism_result": "PASS", "batch_complete": True}
        self.assertTrue(smoke_matches_config(report, cfg))
        for field in ("testbed_profile", "firmware_build_id", "android_build_id", "config_fingerprint", "mechanism_result"):
            self.assertFalse(smoke_matches_config({**report, field: "wrong"}, cfg))
        with tempfile.TemporaryDirectory() as d:
            cfgpath, reportpath = Path(d)/"config.json", Path(d)/"separate.json"
            cfgpath.write_text(json.dumps(cfg))
            reportpath.write_text(json.dumps({**report, "passed": False}))
            with patch("sys.argv", ["run.py", "run", "--config", str(cfgpath), "--smoke-report", str(reportpath)]), patch("builtins.print"), patch("resqmesh_controller.cli.validate_discovered_nodes") as devices:
                self.assertEqual(3, main())
                devices.assert_not_called()


if __name__ == "__main__":
    unittest.main()
