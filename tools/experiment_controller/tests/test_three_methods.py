"""Synthetic software fixtures only; never used as thesis measurements."""
import copy
import json
import tempfile
import unittest
import zipfile
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

from resqmesh_controller.config import MODES, METHOD_PARAMETERS, ConfigError, validate_config
from resqmesh_controller.cli import load_config
from resqmesh_controller.controller import build_plan, ExperimentController
from resqmesh_controller.comparisons import descriptive_comparisons
from resqmesh_controller.excel_report import write_analysis_workbook
from resqmesh_controller.log_merge import aggregate, method_decision_errors, summarize_trial, trickle_timing_errors
from test_controller import physical_config, fake_nodes, standard_provider
from test_log_merge import valid_record, delivery_events, trial_event


class ThreeMethodsTest(unittest.TestCase):
    def test_powershell_bom_config_and_separate_session_guard(self):
        config = physical_config()
        with tempfile.TemporaryDirectory() as path:
            root = Path(path)
            file = root / 'config.json'
            file.write_text(json.dumps(config), encoding='utf-8-sig')
            self.assertEqual(config, load_config(file))
            controller = ExperimentController(config, [], root)
            controller._save_manifest()
            with self.assertRaisesRegex(Exception, 'session differs'):
                ExperimentController(config | {'session_id': 'separate-main-session'}, [], root)
    def test_135_plan_has_15_reproducible_balanced_blocks(self):
        config = physical_config() | {'trial_order': 'balanced_randomized', 'random_seed': 8127}
        plan = build_plan(config)
        self.assertEqual(135, len(plan))
        self.assertEqual(135, len({p.trial_id for p in plan}))
        self.assertEqual({15}, set(Counter((p.mode, p.hypothesis) for p in plan).values()))
        expected = {(mode, hop) for mode in MODES for hop in ('H1', 'H2', 'H3')}
        for block in range(15):
            rows = plan[block * 9:(block + 1) * 9]
            self.assertEqual(expected, {(p.mode, p.hypothesis) for p in rows})
            self.assertTrue(all(p.block == block + 1 for p in rows))
        self.assertEqual(plan, build_plan(config))
        self.assertNotEqual(plan, build_plan(config | {'random_seed': 8128}))

    def test_parameters_differ_only_in_suppression(self):
        full, ablation = copy.deepcopy(METHOD_PARAMETERS['trickle']), copy.deepcopy(METHOD_PARAMETERS['trickle_no_suppression'])
        self.assertTrue(full.pop('suppression_enabled'))
        self.assertFalse(ablation.pop('suppression_enabled'))
        self.assertEqual(full, ablation)
        self.assertEqual(256000, full['imin_ms'] * 2 ** full['imax_doublings'])

    def test_old_two_method_config_is_not_silently_reused(self):
        with self.assertRaises(ConfigError):
            validate_config(physical_config() | {'modes': ['basic_flooding', 'trickle']})

    def test_decision_attribution_and_disabled_suppression(self):
        for mode in ('trickle', 'trickle_no_suppression'):
            for c in (0, 1, 20):
                allowed = mode != 'trickle' or c < 1
                event = trial_event('TRICKLE_TX_OPPORTUNITY', reason='ALLOWED' if allowed else 'SUPPRESSED',
                                    consistency_count=c, k=1, suppression_enabled=mode == 'trickle')
                record = valid_record(mode=mode, method_design_version=3, source_node_id='R1')
                self.assertEqual(set(), method_decision_errors([event], record))
                event['reason'] = 'SUPPRESSED' if allowed else 'ALLOWED'
                self.assertIn('TRICKLE_DECISION_MISMATCH:R1', method_decision_errors([event], record))
        self.assertIn('SUPPRESSION_DISABLED_BUT_SUPPRESSED:R1', method_decision_errors(
            [trial_event('TRICKLE_TX_SUPPRESSED')], valid_record(mode='trickle_no_suppression')))

    def test_same_half_interval_validation_for_both_trickle_variants(self):
        events = [trial_event('TRICKLE_INTERVAL_STARTED', timestamp_ms=10000, monotonic_ms=10000,
                             interval_ms=8000, interval_started_at_monotonic_ms=10000, transmit_at_monotonic_ms=15000),
                  trial_event('ADVERTISE_BURST_REQUESTED', timestamp_ms=10017, monotonic_ms=10017)]
        for mode in ('trickle', 'trickle_no_suppression'):
            self.assertIn('TRICKLE_TRANSMIT_BEFORE_HALF_INTERVAL:R1', trickle_timing_errors(events, valid_record(mode=mode)))

    def test_closed_interval_requires_decision_or_explicit_missed_event(self):
        first = trial_event('TRICKLE_INTERVAL_STARTED', message_id='synthetic', timestamp_ms=10000,
                            monotonic_ms=10000, interval_ms=8000,
                            interval_started_at_monotonic_ms=10000, transmit_at_monotonic_ms=17906)
        second = trial_event('TRICKLE_INTERVAL_STARTED', message_id='synthetic', timestamp_ms=18000,
                             monotonic_ms=18000, interval_ms=16000,
                             interval_started_at_monotonic_ms=18000, transmit_at_monotonic_ms=27000)
        missed = trial_event('TRICKLE_TX_MISSED', message_id='synthetic', timestamp_ms=18000,
                             monotonic_ms=18000, interval_ms=8000,
                             interval_started_at_monotonic_ms=10000, transmit_at_monotonic_ms=17906)
        for mode in ('trickle', 'trickle_no_suppression'):
            record = valid_record(mode=mode, method_design_version=3)
            self.assertIn('TRICKLE_OPPORTUNITY_UNACCOUNTED:R1', trickle_timing_errors([first, second], record))
            self.assertEqual(set(), trickle_timing_errors([first, missed, second], record))
            self.assertIn('TRICKLE_MISSED_BEFORE_INTERVAL_END:R1',
                          trickle_timing_errors([first, missed | {'monotonic_ms': 17999}, second], record))

    def test_missed_is_not_suppression_or_successful_burst(self):
        events = delivery_events() + [trial_event('TRICKLE_TX_MISSED', node_id='source', timestamp_ms=900,
                                                 monotonic_ms=18000, interval_ms=8000,
                                                 interval_started_at_monotonic_ms=10000)]
        row = summarize_trial('trial-1', events, valid_record(observation_started_at_ms=1000, observation_ended_at_ms=2000))
        self.assertEqual('SUCCESS', row['result'])
        self.assertEqual(1, row['missed_opportunities_total'])
        self.assertEqual(0, row['transmission_missed'])
        self.assertEqual(0, row['transmission_suppressed'])
        self.assertEqual(0, row['transmission_bursts'])
        self.assertEqual(300, row['e2e_latency_ms'])

    def test_failed_delivery_is_valid_and_zero_suppression_smoke_is_valid(self):
        config = physical_config() | {'valid_trials_per_condition': 1, 'max_attempts_per_condition': 1}
        def no_suppression(node, session, trial):
            return [e for e in standard_provider()(node, session, trial)
                    if e['event_type'] not in {'TRICKLE_CONSISTENT_HEARD', 'TRICKLE_TX_SUPPRESSED'}]
        with tempfile.TemporaryDirectory() as path:
            ctrl = ExperimentController(config, fake_nodes(config, no_suppression), Path(path), sleep=lambda _: None)
            ctrl.run()
            self.assertTrue(ctrl.smoke_report()['passed'])
            self.assertEqual(9, len(ctrl.smoke_report()['conditions']))
        events = [delivery_events()[0]]
        row = summarize_trial('trial-1', events, valid_record('FAILED_DELIVERY'))
        self.assertTrue(row['valid'])
        self.assertIsNone(row['e2e_latency_ms'])

    def test_invalid_replacements_preserve_balanced_original_plan(self):
        config = physical_config() | {'valid_trials_per_condition': 2, 'max_attempts_per_condition': 3,
                                     'trial_order': 'balanced_randomized'}
        with tempfile.TemporaryDirectory() as path:
            ctrl = ExperimentController(config, fake_nodes(config, standard_provider(
                lambda n: 'INVALID' if n == 1 else 'FAILED_DELIVERY')), Path(path), sleep=lambda _: None)
            original = copy.deepcopy(ctrl.manifest['trial_order'])
            ctrl.run()
            self.assertEqual(original, ctrl.manifest['trial_order'][:18])
            self.assertTrue(ctrl.batch_summary()['complete'])
            for condition in ctrl.batch_summary()['conditions']:
                self.assertEqual(2, condition['valid'])
                self.assertEqual(2, condition['failed_delivery'])
                self.assertEqual(1, condition['invalid'])

    def test_ratios_counts_charts_and_raw_phy_preserved(self):
        trials = []
        for mode in MODES:
            for n, (accepted, duplicates) in enumerate(((1, 9), (9, 0))):
                trials.append(dict(trial_id=f'synthetic-{mode}-{n}', mode=mode, hypothesis='H1', valid=True,
                                   result='SUCCESS' if n == 0 else 'FAILED_DELIVERY', e2e_latency_ms=450 if n == 0 else None,
                                   accepted=accepted, duplicates=duplicates, ldr=duplicates / (accepted + duplicates),
                                   transmission_bursts=4, transmission_opportunities=2 if mode != 'basic_flooding' else 0,
                                   transmission_allowed=1 if mode == 'trickle' else 2,
                                   transmission_suppressed=1 if mode == 'trickle' else 0, other_cancellations=1))
        rows = aggregate(trials)
        self.assertEqual(10, len(descriptive_comparisons(rows)))
        for row in rows:
            self.assertEqual(2, row['valid_trials'])
            self.assertEqual(1, row['success_trials'])
            self.assertEqual(.5, row['dsr'])
            self.assertEqual(450, row['e2e_median'])
            self.assertAlmostEqual(9 / 19, row['ldr'])
            self.assertAlmostEqual(.45, row['ldr_mean_per_trial'])
        phy = dict(event_type='BLE_RX_PHY_OBSERVED', observation_id='synthetic-obs', primary_phy=3,
                   secondary_phy=3, legacy=False, coding='unknown')
        with tempfile.TemporaryDirectory() as path:
            output = Path(path) / 'synthetic.xlsx'
            write_analysis_workbook(output, manifest={'method_parameters': METHOD_PARAMETERS}, events=[phy],
                                    trial_summaries=trials, aggregates=rows, attempts=[], invalid_trials=[])
            workbook = load_workbook(output)
            try:
                self.assertEqual(4, len(workbook['Charts']._charts))
                self.assertEqual(3, len(workbook['Charts']._charts[0].series))
                headers = [c.value for c in workbook['RX PHY'][1]]
                self.assertEqual(3, workbook['RX PHY'].cell(2, headers.index('primary_phy') + 1).value)
                self.assertIs(workbook['RX PHY'].cell(2, headers.index('legacy') + 1).value, False)
                self.assertEqual(4, workbook['Method Parameters'].max_row)
            finally:
                workbook.close()
            with zipfile.ZipFile(output) as archive:
                self.assertIsNone(archive.testzip())


if __name__ == '__main__':
    unittest.main()
