import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from resqmesh_controller.cli import build_nodes, main, validate_discovered_nodes
from resqmesh_controller.config import ConfigError, research_fingerprint, validate_config
from resqmesh_controller.controller import TrialSpec, build_plan
from resqmesh_controller.devices import DeviceError, SerialNode, write_jsonl
from resqmesh_controller.esp_only_validation import scenario_checks
from resqmesh_controller.mpl_config import DEFAULTS, METHODS, SEMANTICS
from resqmesh_controller.neighbor_experiment import (
    NeighborExperimentController, merge_neighbor, neighbor_parameters, stable_id, summarize_network,
)
from resqmesh_controller.neighbor_testbed import ESP_ONLY, ESP_SCENARIOS, testbed as resolve_testbed
from resqmesh_controller.neighbor_validation import validate_logs

ROOT = Path(__file__).resolve().parents[3]
KEY = "100:200"
STATE = "100:200:1:0:0"


def config():
    value = json.loads((ROOT / "tools/experiment_controller/config.mpl.esp-only.smoke.example.json").read_text())
    value.update(firmware_build_id="0123456789ab", session_id="esp-fixture")
    for i, node in enumerate(value["nodes"], 8):
        node["port"] = f"COM{i}"
    return value


class Clock:
    def __init__(self):
        self.now = 1000.0

    def monotonic(self):
        return self.now

    def time_ns(self):
        return int((1791500000 + self.now) * 1_000_000_000)

    def sleep(self, seconds):
        self.now += max(seconds, .001)


class EspFake:
    def __init__(self, node, cfg, clock):
        self.node_id, self.role = node["node_id"], node["role"]
        self.transport = "serial"
        self.cfg, self.clock = cfg, clock
        self.design = resolve_testbed(cfg)
        self.commands, self.events = [], []
        self.session_id = self.trial_id = self.mode = ""
        self.started = None
        self.rx = self.tx = True
        self.boot = 1
        self.duration = 180
        self.failures = {}
        self.missing_delivery = False
        self.transport_error = None

    def event(self, kind, at=None, **fields):
        at = self.clock.time_ns()//1_000_000 if at is None else at
        value = {"kind": "event", "node_id": self.node_id, "session_id": self.session_id,
                 "trial_id": self.trial_id, "event_type": kind, "event_sequence": len(self.events)+1,
                 "timestamp_ms": at, "monotonic_ms": at-1791500000000, "clock_domain": "esp_boot_millis",
                 "clock_sync_valid": True, "clock_offset_ms": 0, "local_boot_id": self.boot,
                 "message_key": KEY, "state_identity": STATE, "scope": stable_id(self.trial_id),
                 "transmitter_id": stable_id(self.node_id), "boot_id": 1, "transmission_sequence": 1, **fields}
        self.events.append(value)
        return value

    def command(self, name, args):
        self.commands.append((name, copy.deepcopy(args)))
        self.clock.now += .002
        if name in self.failures:
            raise self.failures[name]
        if name == "configure_session":
            self.session_id, self.mode, self.role = args["session_id"], args["mode"], args["role"]
            self.duration = args["observation_window_ms"] / 1000
        elif name == "start_trial":
            self.trial_id = args["trial_id"]
            self.events = []
            self.event("TRIAL_WINDOW_STARTED")
        elif name == "trigger_sos":
            self.event("SOS_CREATED")
            self.started = self.clock.time_ns()//1_000_000
            self.event("SOURCE_FIRST_ADVERTISE_STARTED")
            self.event("DATA_BURST_STARTED")
            return {"ok": True}  # Real ESP response has no message_key.
        elif name == "set_node_participation":
            self.rx = self.tx = args["enabled"]
            if self.rx and self.mode == "trickle_mpl":
                self.boot += 1
            self.event("NODE_PARTICIPATION_CHANGED", confirmed_enabled=self.rx,
                       rx_enabled=self.rx, tx_enabled=self.tx, scanner_registered=self.rx)
            return {"ok": True, "confirmed_enabled": self.rx, "local_boot_id": self.boot}
        elif name == "end_observation_window":
            end = args["observation_ended_at_ms"]
            start = end-int(self.duration*1000)
            if self.role != "SOURCE" and not self.missing_delivery:
                activated = [e["timestamp_ms"] for e in self.events if e.get("confirmed_enabled") is True]
                received_at = max(start+1000, max(activated, default=start)+1000)
                parent = self.design.adjacency(self.node_id)[0]
                self.event("DATA_RECEIVED", received_at, transmitter_id=parent)
            self.event("TRIAL_WINDOW_ENDED", end)
            self.events.sort(key=lambda e: e["timestamp_ms"])
            for i, e in enumerate(self.events, 1):
                e["event_sequence"] = i
        elif name == "reset_trial":
            self.trial_id = ""
            self.started = None
            self.rx = self.tx = True
        if name in {"readiness", "get_status"}:
            return {"ok": True, "node_id": self.node_id, "session_id": self.session_id,
                    "trial_id": self.trial_id or None, "scope": stable_id(self.trial_id),
                    "mode": self.mode, "role": self.role, "protocol_active": True,
                    "epoch_valid": True, "clock_valid": True, "scanner": self.rx,
                    "packet_pending": bool(self.trial_id and self.node_id == "esp-r2b"),
                    "advertising": False, "queue_size": 0, "quiet_period_complete": True,
                    "build_id": self.cfg["firmware_build_id"], "neighbor_design_version": 1,
                    "transport_profile": "neighbor_graph_v1", "transport_version": "resqmesh-neighbor-v1",
                    "transmitter_id": stable_id(self.node_id), "allowed_transmitters": self.design.adjacency(self.node_id),
                    "protocol_version": "resqmesh-ble17-v1", "payload_length": 17, "manufacturer_id": 65535,
                    "data_frame_length": 39, "trickle_imin_ms": 8000, "trickle_imax_ms": 256000,
                    "trickle_k": 1, "burst_duration_ms": 2000, "suppression_enabled": self.mode in {"trickle", "trickle_mpl"},
                    "neighbor_parameters": neighbor_parameters(self.cfg), "rx_burst_gap_ms": 1000,
                    "scheduler_semantics": SEMANTICS, "supported_scheduler_semantics": [SEMANTICS],
                    "buffer_retention": "persistent_until_supersession_ack_admin", "mpl_parameters": DEFAULTS,
                    "source_first_advertise_started_at_ms": self.started,
                    "source_first_advertise_message_key": KEY if self.started else None,
                    "radio": {"requested_mode": "coded", "configured_mode": "coded", "ready": True,
                              "primary_phy": "coded", "secondary_phy": "coded", "scan_phy": "coded",
                              "advertising_interval_units": 400, "coding_selection_support": "unsupported",
                              "s8_requirement_accepted": False, "on_air_coding_verified": False}}
        return {"ok": True}

    def diagnostic_events(self, session_id, trial_id):
        return [dict(e) for e in self.events if e["session_id"] == session_id and e["trial_id"] == trial_id]

    def collect_events(self, session_id, trial_id):
        events = self.diagnostic_events(session_id, trial_id)
        self.events = []
        return events


def mechanism_fixture():
    cfg = config()
    spec = TrialSpec("trickle_mpl", "S2_POST_DATA_STOP", 1)
    device_id = "esp-fixture-t1"
    scope = stable_id(device_id)
    start = 10000
    record = {"testbed_profile": ESP_ONLY, **resolve_testbed(cfg).metadata(), "trial_id": spec.trial_id,
              "device_trial_id": device_id, "session_id": cfg["session_id"], "mode": spec.mode,
              "hypothesis": spec.hypothesis, "result": "SUCCESS", "message_key": KEY,
              "config_fingerprint": "fixture-fingerprint", "observation_started_at_ms": start,
              "observation_ended_at_ms": start+420000, "scope": scope,
              "pre_activation_status": {"node_id": "esp-r2b", "captured_at_ms": start+299990,
                                        "response": {"ok": True, "scope": scope, "packet_pending": True,
                                                     "session_id": cfg["session_id"], "trial_id": device_id}}}
    manifest = {"testbed_profile": ESP_ONLY, **resolve_testbed(cfg).metadata(), "session_id": cfg["session_id"],
                "activation_tolerance_ms": 10000, "config_fingerprint": "fixture-fingerprint",
                "mpl_parameters": DEFAULTS, "scheduler_semantics": SEMANTICS,
                "synthetic_data": True, "experiment_methods": list(METHODS),
                "neighbor_scenarios": list(ESP_SCENARIOS), "trials": {spec.trial_id: record}}
    values = []

    def event(node, kind, at, **fields):
        value = {"session_id": cfg["session_id"], "trial_id": spec.trial_id, "node_id": node,
                 "event_type": kind, "event_sequence": len(values)+1, "scope": scope,
                 "timestamp_ms": start+at, "monotonic_ms": at, "clock_sync_valid": True,
                 "clock_offset_ms": 0, "message_key": KEY, "state_identity": STATE,
                 "timer_kind": "data", "timer_key": STATE, "generation": 1,
                 "transmitter_id": stable_id(node), "boot_id": 1, "transmission_sequence": 1, **fields}
        values.append(value)
        return value

    event("esp-r1b", "SOS_CREATED", -20)
    event("esp-r1b", "SOURCE_FIRST_ADVERTISE_STARTED", 0)
    event("esp-r1b", "DATA_BURST_STARTED", 0)
    event("esp-destination", "NODE_PARTICIPATION_CHANGED", -10, confirmed_enabled=False,
          rx_enabled=False, tx_enabled=False, scanner_registered=False)
    event("esp-r2b", "MPL_TIMER_STOPPED", 250000, active=False, buffer_retained=True, expiration_count=5)
    event("esp-destination", "NODE_PARTICIPATION_CHANGED", 300000, confirmed_enabled=True,
          rx_enabled=True, tx_enabled=True, scanner_registered=True, local_boot_id=2)
    event("esp-r2b", "MPL_RX_CLASSIFIED", 303000, peer_id=stable_id("esp-destination"), peer_boot=2,
          transmission_sequence=2, frame_type="status", inventory=[], snapshot_complete=True)
    for kind, at in (("MPL_REPAIR_PENDING", 303010), ("MPL_REPAIR_RESET", 303020)):
        event("esp-r2b", kind, at, peer_id=stable_id("esp-destination"), peer_boot=2, transmission_sequence=2)
    event("esp-r2b", "MPL_TX_ALLOWED", 307000)
    event("esp-r2b", "MPL_NATIVE_STARTED", 307001)
    event("esp-r2b", "DATA_BURST_STARTED", 307002, transmission_sequence=3)
    event("esp-destination", "DATA_RECEIVED", 307020, transmission_sequence=3, transmitter_id=stable_id("esp-r2b"))
    return manifest, record, values


class EspOnlyExperimentTest(unittest.TestCase):
    def controller(self, directory):
        cfg, clock = config(), Clock()
        nodes = [EspFake(n, cfg, clock) for n in cfg["nodes"]]
        return NeighborExperimentController(cfg, nodes, Path(directory), sleep=clock.sleep), nodes, clock

    def clock_patch(self, clock):
        return patch.multiple("resqmesh_controller.controller.time", monotonic=clock.monotonic, time_ns=clock.time_ns)

    def test_cli_config_read_errors_do_not_touch_devices(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            invalid = root / "invalid.json"
            invalid.write_text("{", encoding="utf-8")
            encoding = root / "encoding.json"
            encoding.write_bytes(b"\xff")
            for config_path in (root / "missing.json", invalid, encoding):
                with self.subTest(config_path=config_path), \
                     patch("sys.argv", ["run.py", "plan", "--config", str(config_path), "--output", str(root / "output")]), \
                     patch("sys.stdout", new_callable=io.StringIO) as capture, \
                     patch("resqmesh_controller.cli.validate_discovered_nodes") as discovery, \
                     patch("resqmesh_controller.cli.build_nodes") as devices:
                    self.assertEqual(2, main())
                    result = json.loads(capture.getvalue())
                    self.assertFalse(result["ok"])
                    self.assertEqual(str(config_path.resolve()), result["config"])
                    self.assertIn("COM", result["instruction"])
                    discovery.assert_not_called()
                    devices.assert_not_called()
            self.assertFalse((root / "output").exists())

    def test_plan_nine_randomized_conditions_and_39_minutes(self):
        cfg = config()
        validate_config(cfg)
        plan = build_plan(cfg)
        self.assertEqual(9, len(plan))
        self.assertEqual(plan, build_plan(cfg))
        self.assertEqual(9, len({(p.mode, p.hypothesis) for p in plan}))
        self.assertEqual(2340, sum(ESP_SCENARIOS[p.hypothesis]["observation_window_seconds"] for p in plan))
        self.assertTrue(all(p.block == 1 for p in plan))
        self.assertNotEqual([(m, s) for m in METHODS for s in ESP_SCENARIOS], [(p.mode, p.hypothesis) for p in plan])

    def test_cli_preserves_cleanup_warning_when_partial_export_fails(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config()), encoding="utf-8")
            for error in (DeviceError("ESP_ONLY_CLEANUP_UNCONFIRMED"), KeyboardInterrupt()):
                with self.subTest(error=type(error).__name__):
                    ctrl, nodes, _ = self.controller(root / "smoke_run")
                    (ctrl.output_dir / "raw").mkdir(parents=True, exist_ok=True)
                    capture = io.StringIO()
                    with patch("sys.argv", ["run.py", "smoke", "--config", str(config_path), "--output", d]), \
                         patch("sys.stdout", capture), \
                         patch("resqmesh_controller.cli.validate_discovered_nodes"), \
                         patch("resqmesh_controller.cli.build_nodes", return_value=nodes), \
                         patch("resqmesh_controller.neighbor_experiment.NeighborExperimentController", return_value=ctrl), \
                         patch.object(ctrl, "run", side_effect=error), \
                         patch("resqmesh_controller.cli.merge_directory", side_effect=ValueError("broken partial archive")):
                        for node in nodes:
                            node.close = unittest.mock.Mock()
                        result = main()
                    self.assertEqual(130 if isinstance(error, KeyboardInterrupt) else 2, result)
                    report = json.loads(capture.getvalue())
                    self.assertEqual("broken partial archive", report["partial_export_error"])
                    self.assertIn("matikan ESP", report["instruction"])
                    self.assertTrue(ctrl.manifest_path.exists())
                    for node in nodes:
                        node.close.assert_called_once()

    def test_profile_validation_and_fingerprint_isolation(self):
        cfg = config()
        for change in ({"nodes": cfg["nodes"][:-1]}, {"max_attempts_per_condition": 2},
                       {"android_build_id": "unknown", "firmware_build_id": "unknown"},
                       {"scenarios": ["S0_MAIN"]}, {"activation_tolerance_ms": 0},
                       {"scenario_parameters": {}}, {"quiet_period_seconds": 0},
                       {"source_node_id": "android-source"}, {"graph": []}):
            with self.subTest(change=change), self.assertRaises(ConfigError):
                validate_config({**cfg, **change})
        for field, value in (("port", "COM_A"), ("transport", "adb"), ("role", "SOURCE")):
            bad = copy.deepcopy(cfg)
            bad["nodes"][1][field] = value
            with self.subTest(field=field), self.assertRaises(ConfigError):
                validate_config(bad)
        self.assertNotEqual(research_fingerprint(cfg), research_fingerprint({**cfg, "random_seed": 99}))

    def test_discovery_and_build_nodes_never_use_adb(self):
        cfg = config()
        records = [{"port": n["port"]} for n in cfg["nodes"]]
        with patch("resqmesh_controller.cli.discover_adb", side_effect=AssertionError("ADB invoked")), patch("resqmesh_controller.cli.discover_serial", return_value=records):
            validate_discovered_nodes(cfg)
        self.assertTrue(all(isinstance(n, SerialNode) for n in build_nodes(cfg)))

    def test_configure_readiness_source_and_durations(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            for scenario, parameters in ESP_SCENARIOS.items():
                spec = TrialSpec("trickle_mpl", scenario, 1)
                ctrl.configure(spec)
                self.assertEqual(parameters["observation_window_seconds"], ctrl.observation_window_seconds(spec))
                for node in nodes:
                    args = next(a for n, a in reversed(node.commands) if n == "configure_session")
                    self.assertEqual(parameters["observation_window_seconds"]*1000, args["observation_window_ms"])
                    self.assertEqual(ctrl.testbed.adjacency(node.node_id), args["allowed_transmitters"])
                    self.assertEqual([], ctrl._readiness_errors(node, node.command("readiness", {}), spec, False, False))
            self.assertEqual(4, len(ctrl.manifest["target_node_ids"]))
            self.assertEqual("SOURCE", nodes[0].role)

    def test_readiness_archives_actual_builds_and_rejects_mismatch(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, _ = self.controller(d)
            for node in nodes:
                original = node.command
                def mismatch(name, args, original=original):
                    result = original(name, args)
                    if name == "readiness":
                        result["build_id"] = "fedcba987654"
                    return result
                node.command = mismatch
            with self.assertRaisesRegex(DeviceError, "expected=0123456789ab, actual=fedcba987654"):
                ctrl.readiness()
            saved = json.loads(ctrl.manifest_path.read_text())
            self.assertEqual(5, len(saved["readiness_results"]))
            self.assertTrue(all(r["build_id"] == "fedcba987654" for r in saved["readiness_results"].values()))
            self.assertEqual({}, saved["trials"])
            self.assertTrue(all([name for name, _ in node.commands] == ["readiness"] for node in nodes))

    def test_integration_all_nine_conditions_with_actual_timing_and_no_android_calls(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            with self.clock_patch(clock):
                results = ctrl.run()
            self.assertEqual(9, len(results))
            self.assertTrue(all(r["result"] == "SUCCESS" for r in results))
            for r in results:
                params = ESP_SCENARIOS[r["hypothesis"]]
                self.assertEqual(params["observation_window_seconds"]*1000, r["observation_ended_at_ms"]-r["observation_started_at_ms"])
                self.assertEqual(4, r["evidence"]["N"])
                self.assertEqual(100, r["evidence"]["dsr_percent"])
                self.assertTrue(r["reset_verified"])
                for change in r.get("participation", []):
                    if change["enabled"]:
                        elapsed = change["requested_at_ms"]-r["observation_started_at_ms"]
                        self.assertGreaterEqual(elapsed, params["activate_at_seconds"]*1000)
                        self.assertLess(elapsed, params["activate_at_seconds"]*1000+1000)
            forbidden = {"store_neighbor_metrics", "export_trial", "finalize_trial"}
            self.assertFalse(any(name in forbidden for n in nodes for name, _ in n.commands))
            # Successful delivery alone cannot claim the synthetic nodes demonstrated MPL repair.
            report = ctrl.smoke_report()
            self.assertTrue(report["batch_complete"])
            self.assertTrue(report["delivery_passed"])
            self.assertFalse(report["passed"])
            self.assertEqual("INCONCLUSIVE", report["mechanism_result"])
            exported = merge_neighbor(Path(d)/"raw", Path(d)/"merged", ctrl.manifest)
            metrics = json.loads((Path(d)/"merged/network_metrics.json").read_text())
            self.assertEqual(9, len(metrics))
            self.assertTrue(all(r["N"] == 4 for r in metrics))
            self.assertTrue(all(r["analysis_group"] == "utama" for r in metrics if r["scenario"] == "S0_STABLE"))
            wb = load_workbook(exported["workbook"])
            try:
                self.assertEqual(6, len(wb["Grafik Kumulatif"]._charts))
            finally:
                wb.close()

    def test_source_identity_snapshot_keeps_events_and_rejects_wrong_scope(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            node = nodes[0]
            spec = TrialSpec("basic_flooding", "S0_STABLE", 1)
            ctrl.configure(spec)
            node.command("start_trial", {"trial_id": ctrl._device_trial_id(spec)})
            node.command("trigger_sos", {})
            count = len(node.events)
            self.assertEqual(KEY, ctrl.source_message_key(node, spec, {"ok": True}))
            self.assertEqual(count, len(node.events))
            node.events.append(dict(node.events[1]))
            with self.assertRaisesRegex(DeviceError, "AMBIGUOUS"):
                ctrl.source_message_key(node, spec, {})
            node.events = [{**node.events[1], "scope": 999}]
            with self.clock_patch(clock), self.assertRaisesRegex(DeviceError, "MISSING"):
                ctrl.source_message_key(node, spec, {})

    def test_clock_uncertainty_cannot_be_relaxed_for_esp_only(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            spec = TrialSpec("basic_flooding", "S0_STABLE", 1)
            ctrl.manifest["trials"][spec.trial_id] = {}
            original = nodes[0].command
            def delayed(name, args):
                if name == "clock_sync":
                    clock.now += .3
                return original(name, args)
            nodes[0].command = delayed
            with self.clock_patch(clock), self.assertRaisesRegex(DeviceError, "clock uncertainty"):
                ctrl.synchronize_clocks(spec)
            self.assertEqual(3, len(ctrl.manifest["trials"][spec.trial_id]["clock_sync_attempts"][nodes[0].node_id]))
            self.assertEqual(100, ctrl.config["clock_tolerance_ms"])

    def test_metric_duplicates_scope_filter_and_source_rx_exclusion(self):
        manifest, record, events = mechanism_fixture()
        rx = events[-1]
        duplicate = {**rx, "timestamp_ms": rx["timestamp_ms"]+1}
        foreign = {**rx, "scope": 999, "transmission_sequence": 9}
        source_rx = {**rx, "node_id": "esp-r1b", "transmitter_id": stable_id("esp-r1a")}
        metric = summarize_network(events+[duplicate, foreign, source_rx], record)
        self.assertEqual((4, 1, 1, 0), (metric["N"], metric["R"], metric["U"], metric["ldr_percent"]))
        repeated = {**rx, "timestamp_ms": rx["timestamp_ms"]+2, "transmission_sequence": 4}
        metric = summarize_network(events+[repeated], record)
        self.assertEqual((2, 1, 50), (metric["R"], metric["U"], metric["ldr_percent"]))

    def test_smoke_ready_requires_delivery_core_and_mechanism(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            checks = []
            for spec in build_plan(ctrl.config):
                ctrl.manifest["trials"][spec.trial_id] = {"mode": spec.mode, "hypothesis": spec.hypothesis, "result": "SUCCESS"}
                for name in ("LOG_COMPLETENESS", "METRICS_RECOMPUTE_MATCH", "ESP_SCENARIO_PARTICIPATION"):
                    checks.append({"trial_id": spec.trial_id, "method": spec.mode, "check": name, "result": "PASS"})
                if spec.mode == "trickle_mpl":
                    for name in ("MPL_LISTEN_ONLY_AND_BOUNDS", "MPL_C_K_DECISIONS", "MPL_PARAMETERS_MATCH", "MPL_OPPORTUNITY_ACCOUNTING", "MPL_EXPIRATION_AND_DOUBLING"):
                        checks.append({"trial_id": spec.trial_id, "method": spec.mode, "check": name, "result": "PASS"})
            checks.append({"trial_id": TrialSpec("trickle_mpl", "S2_POST_DATA_STOP", 1).trial_id,
                           "method": "trickle_mpl", "check": "ESP_POST_STOP_REPAIR", "result": "PASS"})
            with patch("resqmesh_controller.neighbor_validation.validate_logs", return_value=checks):
                self.assertTrue(ctrl.smoke_report()["passed"])
                checks[-1]["result"] = "INCONCLUSIVE"
                self.assertFalse(ctrl.smoke_report()["passed"])
                checks[-1]["result"] = "PASS"
                checks[0]["result"] = "INCONCLUSIVE"
                self.assertFalse(ctrl.smoke_report()["passed"])
                checks[0]["result"] = "PASS"
                required = checks.pop(0)
                self.assertFalse(ctrl.smoke_report()["passed"])
                checks.insert(0, required)
                next(iter(ctrl.manifest["trials"].values()))["result"] = "FAILED_DELIVERY"
                report = ctrl.smoke_report()
                self.assertTrue(report["batch_complete"])
                self.assertFalse(report["delivery_passed"])
                self.assertFalse(report["passed"])

    def test_serial_port_locked_is_reported_without_adb_or_device_mutation(self):
        import serial
        node = build_nodes(config())[0]
        with patch("serial.Serial", side_effect=serial.SerialException("Access to COM8 denied")):
            with self.assertRaisesRegex(OSError, "denied"):
                node.command("readiness", {"command_id": "locked"})
        node.close()

    def test_cli_serial_only_discovery_does_not_require_adb(self):
        from resqmesh_controller.cli import main
        with patch("sys.argv", ["run.py", "discover", "--serial-only"]), patch("resqmesh_controller.cli.discover_adb", side_effect=AssertionError("ADB invoked")), patch("resqmesh_controller.cli.discover_serial", return_value=[]), patch("builtins.print"):
            self.assertEqual(0, main())

    def test_activation_evidence_rejects_wrong_window_early_on_and_rx_while_off(self):
        manifest, record, events = mechanism_fixture()
        def verdict(samples, rec=record):
            return scenario_checks(samples, rec, manifest)[0]["result"]
        self.assertEqual("PASS", verdict(events))
        early = [dict(e, timestamp_ms=record["observation_started_at_ms"]+290000)
                 if e.get("event_type") == "NODE_PARTICIPATION_CHANGED" and e.get("confirmed_enabled") is True else e for e in events]
        self.assertEqual("FAIL", verdict(early))
        rx_while_off = {**events[-1], "timestamp_ms": record["observation_started_at_ms"]+1000}
        self.assertEqual("FAIL", verdict(events+[rx_while_off]))
        self.assertEqual("INCONCLUSIVE", verdict(events, {**record, "observation_ended_at_ms": record["observation_started_at_ms"]+180000}))

    def test_failed_delivery_is_valid_and_delay_is_not_zero(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            nodes[-1].missing_delivery = True
            with self.clock_patch(clock):
                r = ctrl.run_trial(TrialSpec("basic_flooding", "S0_STABLE", 1))
            self.assertEqual("FAILED_DELIVERY", r["result"])
            self.assertEqual(75, r["evidence"]["dsr_percent"])
            self.assertEqual(1, r["evidence"]["failed_pairs"])
            self.assertEqual(3, len(r["evidence"]["per_receiver"]))

    def test_interrupt_archives_and_resets_without_replaying(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            nodes[0].failures["trigger_sos"] = KeyboardInterrupt()
            spec = TrialSpec("basic_flooding", "S0_STABLE", 1)
            with self.clock_patch(clock), self.assertRaises(KeyboardInterrupt):
                ctrl.run_trial(spec)
            r = ctrl.manifest["trials"][spec.trial_id]
            self.assertEqual(["USER_INTERRUPTED"], r["invalid_reasons"])
            self.assertTrue(r["terminal"])
            self.assertTrue(r["reset_verified"])
            self.assertEqual(5, len(list((Path(d)/"raw"/spec.trial_id).glob("*.jsonl"))))
            self.assertTrue(all(any(n == "reset_trial" for n, _ in node.commands) for node in nodes))

    def test_serial_timeout_aborts_attempt_and_cleanup_failure_stops_batch(self):
        for command in ("trigger_sos", "reset_trial"):
            with self.subTest(command=command), tempfile.TemporaryDirectory() as d:
                ctrl, nodes, clock = self.controller(d)
                nodes[0].failures[command] = DeviceError("serial disconnected / response timeout")
                with self.clock_patch(clock):
                    if command == "reset_trial":
                        with self.assertRaisesRegex(DeviceError, "CLEANUP_UNCONFIRMED"):
                            ctrl.run()
                        self.assertEqual(1, len(ctrl.manifest["trials"]))
                    else:
                        result = ctrl.run_trial(TrialSpec("basic_flooding", "S0_STABLE", 1))
                        self.assertEqual("INVALID", result["result"])
                        self.assertTrue(result["reset_verified"])

    def test_foreign_active_trial_is_not_reset_or_modified(self):
        with tempfile.TemporaryDirectory() as d:
            ctrl, nodes, clock = self.controller(d)
            with self.assertRaisesRegex(DeviceError, "foreign active trial"):
                ctrl.recover_interrupted_trials({nodes[0].node_id: {"trial_id": "foreign"}})
            self.assertFalse(nodes[0].commands)

    def test_post_stop_repair_chain_is_exact_and_conservative(self):
        manifest, record, events = mechanism_fixture()
        def verdict(values, rec=record):
            return next(c["result"] for c in scenario_checks(values, rec, manifest) if c["check"] == "ESP_POST_STOP_REPAIR")
        self.assertEqual("PASS", verdict(events))
        for kind in ("MPL_TIMER_STOPPED", "MPL_RX_CLASSIFIED", "MPL_REPAIR_PENDING", "MPL_REPAIR_RESET", "MPL_NATIVE_STARTED", "DATA_RECEIVED"):
            with self.subTest(kind=kind):
                self.assertEqual("INCONCLUSIVE", verdict([e for e in events if e["event_type"] != kind]))
        for kind, field, value in (("MPL_TIMER_STOPPED", "timer_kind", "control"),
                                   ("MPL_RX_CLASSIFIED", "peer_id", 999),
                                   ("MPL_RX_CLASSIFIED", "inventory", [STATE]),
                                   ("MPL_REPAIR_PENDING", "peer_boot", 999),
                                   ("DATA_RECEIVED", "transmission_sequence", 999),
                                   ("DATA_RECEIVED", "scope", 999)):
            changed = [dict(e, **{field: value}) if e["event_type"] == kind else e for e in events]
            with self.subTest(kind=kind, field=field):
                self.assertEqual("INCONCLUSIVE", verdict(changed))
        for field, value in (("active", True), ("buffer_retained", False), ("expiration_count", 99)):
            changed = [dict(e, **{field: value}) if e["event_type"] == "MPL_TIMER_STOPPED" else e for e in events]
            self.assertEqual("FAIL", verdict(changed))
        restart = {**events[4], "event_type": "MPL_TIMER_RESTARTED", "event_sequence": 6,
                   "timestamp_ms": events[4]["timestamp_ms"]+1000}
        self.assertEqual("FAIL", verdict(events+[restart]))
        missing = copy.deepcopy(record)
        missing.pop("pre_activation_status")
        self.assertEqual("INCONCLUSIVE", verdict(events, missing))

    def test_export_numeric_tables_curves_denominator_and_raw_preservation(self):
        manifest, record, events = mechanism_fixture()
        record["evidence"] = summarize_network(events, record)
        with tempfile.TemporaryDirectory() as d:
            raw = Path(d)/"raw"
            write_jsonl(raw/"events.jsonl", events)
            result = merge_neighbor(raw, Path(d)/"merged", manifest)
            self.assertEqual("resqmesh_esp_only_analysis.xlsx", Path(result["workbook"]).name)
            wb = load_workbook(result["workbook"])
            self.addCleanup(wb.close)
            for name in ("Topologi dan Skenario", "Bukti Skenario", "Kurva Kumulatif", "Grafik Kumulatif"):
                self.assertIn(name, wb.sheetnames)
            self.assertIn("BUKAN DATA UTAMA ANDROID", wb["Mulai Di Sini"]["A1"].value)
            for sheet in wb:
                for table in sheet.tables.values():
                    self.assertEqual(len(sheet[1]), len(table.tableColumns))
            exported = json.loads((Path(d)/"merged/network_metrics.json").read_text())
            self.assertEqual((4, 25), (exported[0]["N"], exported[0]["dsr_percent"]))
            summary = json.loads((Path(d)/"merged/method_scenario_summary.json").read_text())
            self.assertTrue(all(s["N"] == 4 for s in summary))
            self.assertEqual(events, json.loads((Path(d)/"merged/all_events.json").read_text()))
            curves = json.loads((Path(d)/"merged/cumulative_metrics.json").read_text())
            self.assertEqual(420, curves[-1]["elapsed_seconds"])
            self.assertEqual(exported[0]["network_overhead"], curves[-1]["network_overhead"])
            self.assertEqual(exported[0]["dsr_percent"], curves[-1]["dsr_percent"])
            self.assertIsInstance(wb["Trial Metrics"].cell(2, 7).value, (int, float))

    def test_export_rejects_mixed_profile_fingerprint_and_foreign_nodes(self):
        for field, value in (("testbed_profile", "android_plus_five_esp32"), ("config_fingerprint", "foreign")):
            manifest, record, events = mechanism_fixture()
            record[field] = value
            with tempfile.TemporaryDirectory() as d:
                write_jsonl(Path(d)/"raw/events.jsonl", events)
                with self.subTest(field=field), self.assertRaises(ConfigError):
                    merge_neighbor(Path(d)/"raw", Path(d)/"merged", manifest)
        manifest, record, events = mechanism_fixture()
        events[0]["node_id"] = "android-source"
        with tempfile.TemporaryDirectory() as d:
            write_jsonl(Path(d)/"raw/events.jsonl", events)
            with self.assertRaises(ConfigError):
                merge_neighbor(Path(d)/"raw", Path(d)/"merged", manifest)

    def test_source_event_sequence_is_required_too(self):
        manifest, record, events = mechanism_fixture()
        for node in resolve_testbed(manifest).node_ids:
            node_events = [e for e in events if e["node_id"] == node]
            for kind, at in (("TRIAL_WINDOW_STARTED", 9000), ("TRIAL_WINDOW_ENDED", 430000)):
                marker = {**events[0], "node_id": node, "event_type": kind, "timestamp_ms": at}
                events.append(marker)
                node_events.append(marker)
            for i, e in enumerate(sorted(node_events, key=lambda e: e["timestamp_ms"]), 1):
                e["event_sequence"] = i
        self.assertEqual("PASS", next(c["result"] for c in validate_logs(events, manifest) if c["check"] == "LOG_COMPLETENESS"))
        for e in events:
            if e["node_id"] == "esp-r1b":
                e.pop("event_sequence")
        result = validate_logs(events, manifest)
        self.assertEqual("INCONCLUSIVE", next(c["result"] for c in result if c["check"] == "LOG_COMPLETENESS"))


if __name__ == "__main__":
    unittest.main()
