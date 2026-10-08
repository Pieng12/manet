import copy
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from resqmesh_controller.config import validate_config, research_fingerprint, ConfigError
from resqmesh_controller.controller import build_plan, TrialSpec
from resqmesh_controller.mpl_config import SEMANTICS, METHODS, DEFAULTS
from resqmesh_controller.mpl_validation import checks
from resqmesh_controller.neighbor_experiment import NeighborExperimentController, merge_neighbor, summarize_network, TARGETS
from test_neighbor_experiment import ROOT, NeighborFake, record, ev, received, source_events, config as historical


def config(main=False):
    value=json.loads((ROOT/f'tools/experiment_controller/config.mpl.{"main" if main else "pilot"}.example.json').read_text())
    value.update(android_build_id='0123456789ab',firmware_build_id='0123456789ab',session_id='fixture')
    return value


class MplFake(NeighborFake):
    def command(self,name,args):
        value=super().command(name,args)
        if name in ('readiness','get_status'):
            value.update(supported_scheduler_semantics=[SEMANTICS],scheduler_semantics=SEMANTICS,
                         buffer_retention='persistent_until_supersession_ack_admin',
                         mpl_parameters=DEFAULTS,suppression_enabled=self.mode in ('trickle','trickle_mpl'))
        return value


def timer(kind='MPL_INTERVAL_STARTED',**overrides):
    return {'event_type':kind,'node_id':'esp-r1a','timer_kind':'data','timer_key':'123:456:1:0:0',
            'scheduler_semantics':SEMANTICS,'generation':1,'interval_ms':8000,
            'interval_started_at_monotonic_ms':1000,'transmit_at_monotonic_ms':5000,
            'interval_end_at_monotonic_ms':9000,'consistency_count':0,'k':1,
            'expiration_count':0,'monotonic_ms':1000,**overrides}


class MplExperimentTest(unittest.TestCase):
    def test_balanced_main135_pilot27_smoke9_and_historical_profiles(self):
        for main,n in ((True,135),(False,27)):
            c=config(main);validate_config(c);plan=build_plan(c)
            self.assertEqual(n,len(plan));self.assertEqual(plan,build_plan(c))
            for b in range(1,c['valid_trials_per_condition']+1):
                self.assertEqual(9,len({(p.mode,p.hypothesis) for p in plan if p.block==b}))
            self.assertEqual(set(METHODS),{p.mode for p in plan})
        c=config();c.update(valid_trials_per_condition=1,max_attempts_per_condition=1)
        self.assertEqual(9,len(build_plan(c)))
        self.assertEqual(180,len(build_plan(historical(True))))

    def test_parameters_retention_and_window_validated_fingerprinted(self):
        c=config();old=research_fingerprint(c)
        changed=copy.deepcopy(c);changed['mpl_parameters']['mpl_control_k']=2
        self.assertNotEqual(old,research_fingerprint(changed))
        for params in ({'mpl_control_imin_ms':1},{'mpl_data_expirations':0},
                       {'mpl_control_imax_ms':31000},{'mpl_repair_budget':0},{'unknown':1}):
            bad={**c,'mpl_parameters':params}
            with self.subTest(params=params),self.assertRaises(ConfigError):validate_config(bad)
        for override in ({'buffer_retention':'ttl'},{'observation_window_seconds':60},
                         {'scheduler_semantics':'old'},{'modes':list(METHODS)+['trickle_neighbor_status']}):
            with self.subTest(override=override),self.assertRaises(ConfigError):validate_config({**c,**override})

    def test_configure_readiness_and_manifest_reject_old_artifacts(self):
        c=config()
        with tempfile.TemporaryDirectory() as directory:
            nodes=[MplFake(n,c,None) for n in c['nodes']]
            ctrl=NeighborExperimentController(c,nodes,Path(directory))
            self.assertEqual(27,ctrl.manifest['target_valid_trials'])
            self.assertEqual(list(METHODS),ctrl.manifest['experiment_methods'])
            spec=TrialSpec('trickle_mpl','S0_MAIN',1)
            ctrl.configure(spec)
            for node in nodes:
                args=next(a for n,a in node.commands if n=='configure_session')
                self.assertEqual(SEMANTICS,args['scheduler_semantics'])
                self.assertEqual(4,args['mpl_control_expirations'])
                response=node.command('readiness',{})
                self.assertEqual([],ctrl._readiness_errors(node,response,spec,False,False))
                response.pop('supported_scheduler_semantics')
                self.assertIn('MPL scheduler unavailable',ctrl._readiness_errors(node,response,spec,False,False))
                response['mpl_parameters']={}
                self.assertIn('mpl_parameters mismatch',ctrl._readiness_errors(node,response,spec,False,False))

    def test_smoke_report_contains_only_nine_conditions(self):
        c=config()
        with tempfile.TemporaryDirectory() as directory:
            ctrl=NeighborExperimentController(c,[],Path(directory))
            for method in METHODS:
                for scenario in c['scenarios']:
                    ctrl.manifest['trials'][method+scenario]={'mode':method,'hypothesis':scenario,'result':'FAILED_DELIVERY'}
            report=ctrl.smoke_report()
            self.assertTrue(report['passed']);self.assertEqual(9,len(report['conditions']))

    def test_recovery_uses_confirmed_on_native_time_with_clock_correction(self):
        r=record();r['participation']=[{'node':TARGETS[0],'enabled':True,
            'response':{'confirmed_enabled':True,'timestamp_ms':1070}}]
        r['clock_samples']={TARGETS[0]:{'offset_ms':10}}
        metric=summarize_network(source_events()+[received(TARGETS[0])],r)
        self.assertEqual(20,metric['per_receiver'][0]['recovery_delay_ms'])
        self.assertEqual(4,metric['failed_pairs'])
        r['participation'][0]['response']['confirmed_enabled']=False
        self.assertIsNone(summarize_network(source_events()+[received(TARGETS[0])],r)['per_receiver'][0]['recovery_delay_ms'])

    def test_export_raw_control_setup_numeric_tables_three_methods(self):
        events=source_events()+[received(n) for n in TARGETS]+[
            ev('STATUS_BURST_STARTED',TARGETS[0],950),ev('STATUS_BURST_STARTED',TARGETS[0],1300,2),
            ev('MPL_TX_SUPPRESSED',TARGETS[0],1400,3)]
        events[-1].update(timer_kind='control',scheduler_semantics=SEMANTICS)
        manifest={'session_id':'fixture','scheduler_semantics':SEMANTICS,'mpl_parameters':DEFAULTS,
                  'experiment_methods':list(METHODS),'neighbor_scenarios':['S0_MAIN'],
                  'synthetic_data':True,'trials':{'t1':{**record(),'mode':'trickle_mpl'}}}
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);raw=root/'raw';raw.mkdir()
            (raw/'events.jsonl').write_text('\n'.join(json.dumps(e) for e in events))
            result=merge_neighbor(raw,root/'merged',manifest)
            wb=load_workbook(result['workbook']);self.addCleanup(wb.close)
            self.assertIn('MPL Parameters',wb.sheetnames);self.assertIn('MPL Diagnostics',wb.sheetnames)
            self.assertEqual(5,len(result['charts']))
            rows=json.loads((root/'merged/network_metrics.json').read_text())
            self.assertEqual((1,1,2,1,3),(rows[0]['data_tx'],rows[0]['control_tx'],rows[0]['network_overhead'],rows[0]['setup_control_tx'],rows[0]['setup_plus_window_tx']))
            self.assertEqual(events,json.loads((root/'merged/all_events.json').read_text()))
            self.assertTrue((root/'merged/mpl_diagnostics.csv').exists())
            self.assertEqual(4,wb['Method Scenario Summary'].max_row)
            headers=[c.value for c in wb['Trial Metrics'][1]]
            self.assertIsInstance(wb['Trial Metrics'].cell(2,headers.index('dsr_percent')+1).value,(int,float))
            for ws in wb:
                for table in ws.tables.values():self.assertEqual(len(ws[1]),len(table.tableColumns))

    def test_validator_injected_bounds_decisions_retention_and_missing(self):
        cases=[('MPL_LISTEN_ONLY_AND_BOUNDS',timer(transmit_at_monotonic_ms=1001)),
               ('MPL_C_K_DECISIONS',timer('MPL_TX_ALLOWED',consistency_count=1)),
               ('MPL_EXPIRATION_RETAINS_BUFFER',timer('MPL_TIMER_STOPPED',buffer_retained=False,expiration_count=5)),
               ('MPL_NATIVE_INSIDE_INTERVAL',timer('MPL_NATIVE_STARTED',monotonic_ms=9000)),
               ('MPL_PHYSICAL_FRESHNESS',timer('MPL_RX_CLASSIFIED',peer_id=3,peer_boot=1,physical_received_monotonic_ms=1,monotonic_ms=200000)),
               ('MPL_SEMANTICS_PROVENANCE',timer(scheduler_semantics='renamed_old_mode'))]
        for name,event in cases:
            with self.subTest(name=name):self.assertTrue(next(r[2] for r in checks([event],DEFAULTS) if r[0]==name))
        self.assertTrue(next(r[1] for r in checks([timer()],DEFAULTS) if r[0]=='MPL_LISTEN_ONLY_AND_BOUNDS'))
        partial=timer();partial.pop('generation')
        self.assertFalse(next(r[1] for r in checks([partial],DEFAULTS) if r[0]=='MPL_LISTEN_ONLY_AND_BOUNDS'))

    def test_validator_expiration_doubling_reset_storm_and_peer_evidence(self):
        values=[timer(),timer('MPL_INTERVAL_ENDED',expiration_count=3)]
        self.assertTrue(next(r[2] for r in checks(values,DEFAULTS) if r[0]=='MPL_EXPIRATION_AND_DOUBLING'))
        values=[timer('MPL_REPAIR_RESET',peer_id=3,peer_boot=1,episode_until=60000,monotonic_ms=1000+i) for i in range(3)]
        self.assertTrue(next(r[2] for r in checks(values,DEFAULTS) if r[0]=='MPL_REPAIR_RESET_STORM'))
        values=[timer('MPL_REPAIR_PENDING',peer_id=3,peer_boot=1),
                timer('MPL_RX_CLASSIFIED',peer_id=4,peer_boot=1,physical_received_monotonic_ms=1000)]
        self.assertFalse(next(r[1] for r in checks(values,DEFAULTS) if r[0]=='MPL_REPAIR_PEER_EVIDENCE'))

    def test_validator_repair_success_budget_is_per_episode(self):
        def result(values):
            return next(r for r in checks(values,DEFAULTS) if r[0]=='MPL_REPAIR_BOUNDS')
        one=timer('MPL_REPAIR_COMPLETED',peer_id=3,peer_boot=1,episode_until=60000,
                  budget_used=1,budget_limit=2)
        two={**one,'budget_used':2}
        self.assertTrue(result([one,two])[1])
        self.assertTrue(result([one,one,one])[2])
        missing=dict(one);missing.pop('episode_until')
        self.assertFalse(result([missing])[1])
        different_peer={**one,'peer_id':4}
        self.assertTrue(result([one,two,different_peer])[1])


if __name__=='__main__': unittest.main()
