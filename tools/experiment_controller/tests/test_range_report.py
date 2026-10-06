import importlib.util
import json
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location("range_report", Path(__file__).parents[1] / "range_report.py")
report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


class RangeReportTests(unittest.TestCase):
    def dataset(self, folder):
        data = {"run": {"run_id": "pilot", "session_id": "pilot", "trial_id": "pilot-T1",
                        "algorithm": "basic_flooding", "source_position_method": "manual",
                        "source_accuracy_m": None, "android_build_id": "apk"},
                "points": [{"point_id": "p1", "receive_count": 1, "receive_result": "received"}],
                "receives": [{"event_id": 1, "observation_id": "obs", "timestamp_ms": 1790000000123,
                              "distance_m": 123.4, "rssi": -101, "phase": "session",
                              "coded_verified": True, "primary_phy": 3, "secondary_phy": 3,
                              "location": {"latitude": 3.5, "longitude": 98.5,
                                           "timestamp_ms": 1790000000100, "accuracy_m": 5}}],
                "positions": [{"latitude": 3.5, "accuracy_m": 5}],
                "diagnostics": [{"code": "MAP_TILE_UNAVAILABLE", "detail": "=never_execute()"}]}
        for name, value in {"range_trial.json": data,
                            "manifest.json": {"session_id": "pilot", "esp_build_id": "fw"},
                            "source_status.json": {"stop_confirmed": True, "completed_duration": True,
                                                   "final_readiness": {"advertising": False}}}.items():
            (folder / name).write_text(json.dumps(value), encoding="utf-8-sig")
        return data

    def test_six_sheets_numeric_types_valid_tables_and_raw_preservation(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            original = (folder / "range_trial.json").read_bytes()
            path = report.export_range(folder)
            workbook = load_workbook(path)
            self.assertEqual(list(report.SHEETS), workbook.sheetnames)
            rx = dict(zip([c.value for c in workbook["RX Samples"][1]],
                          [c.value for c in workbook["RX Samples"][2]]))
            self.assertEqual(123.4, rx["distance_m"])
            self.assertEqual(-101, rx["rssi"])
            self.assertEqual(1790000000123, rx["timestamp_ms"])
            self.assertEqual("fw", rx["esp_build_id"])
            self.assertEqual("unknown_s2_or_s8", rx["coding"])
            self.assertEqual(original, (folder / "range_trial.json").read_bytes())
            self.assertEqual(data["receives"][0]["location"], json.loads(rx["location"]))
            for sheet in workbook:
                self.assertEqual("A2", sheet.freeze_panes)
                self.assertTrue(all(c.data_type != "f" for row in sheet for c in row))
                self.assertTrue((folder / (sheet.title.lower().replace(" ", "_") + ".csv")).exists())
            with ZipFile(path) as archive:
                tables = [name for name in archive.namelist() if name.startswith("xl/tables/")]
                self.assertEqual(6, len(tables))
                for name in tables:
                    table = ET.fromstring(archive.read(name))
                    columns = table.find("{*}tableColumns")
                    self.assertEqual(int(columns.attrib["count"]), len(columns))
                    self.assertEqual(len(columns), len({c.attrib["name"] for c in columns}))
                    self.assertIsNotNone(table.find("{*}autoFilter"))
            workbook.close()

    def test_empty_sheets_do_not_create_invalid_header_only_tables(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            for key in ("points", "receives", "positions", "diagnostics"):
                data[key] = []
            (folder / "range_trial.json").write_text(json.dumps(data))
            workbook = load_workbook(report.export_range(folder))
            for name in report.SHEETS[2:]:
                self.assertFalse(workbook[name].tables)
            workbook.close()

    def test_after_session_receive_not_used_for_farthest_summary(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            data["receives"].append({"distance_m": 9999, "phase": "after_session"})
            (folder / "range_trial.json").write_text(json.dumps(data))
            workbook = load_workbook(report.export_range(folder))
            overview = dict(workbook["Overview"].values)
            self.assertEqual(123.4, overview["farthest_observed_horizontal_m"])
            workbook.close()

    def test_wrong_run_rejected_without_changing_raw_archive(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            self.dataset(folder)
            (folder / "manifest.json").write_text('{"session_id":"other"}')
            with self.assertRaises(ValueError):
                report.export_range(folder)
            self.assertFalse((folder / "resqmesh_range_analysis.xlsx").exists())

    def test_nearby_gps_uncertainty_is_flagged_without_changing_raw_distance(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            data["run"]["source_accuracy_m"] = 4.3
            data["receives"][0]["distance_m"] = 11.6
            data["receives"][0]["location"]["accuracy_m"] = 10.7
            (folder / "range_trial.json").write_text(json.dumps(data))
            raw = (folder / "range_trial.json").read_bytes()
            workbook = load_workbook(report.export_range(folder))
            rx = dict(zip([c.value for c in workbook["RX Samples"][1]],
                          [c.value for c in workbook["RX Samples"][2]]))
            overview = dict(workbook["Overview"].values)
            self.assertEqual(11.6, rx["distance_m"])
            self.assertEqual(15, rx["gps_accuracy_radii_sum_m"])
            self.assertEqual("not_distinguishable_from_location_uncertainty", rx["gps_distance_quality"])
            self.assertIn("belum dapat dibedakan", rx["gps_distance_explanation"])
            self.assertIsNone(overview["farthest_resolved_gps_estimate_m"])
            self.assertEqual(11.6, overview["farthest_observed_horizontal_m"])
            self.assertEqual(raw, (folder / "range_trial.json").read_bytes())
            workbook.close()

    def test_manual_measurements_numeric_and_separate_from_gps(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            manual = {"measured_horizontal_m": 3, "height_difference_m": -4,
                      "measured_3d_m": 5, "measurement_method": "tape_measure",
                      "point_note": "Lantai 1, dekat tangga"}
            data["points"][0].update(status="completed", **manual)
            data["receives"][0].update(point_id="p1", **manual)
            (folder / "range_trial.json").write_text(json.dumps(data))
            workbook = load_workbook(report.export_range(folder))
            overview = dict(workbook["Overview"].values)
            self.assertEqual(3, overview["farthest_received_manual_horizontal_m"])
            self.assertEqual(5, overview["farthest_received_manual_3d_m"])
            self.assertEqual(123.4, overview["farthest_observed_horizontal_m"])
            for sheet in ("Test Points", "RX Samples"):
                row = dict(zip([c.value for c in workbook[sheet][1]],
                               [c.value for c in workbook[sheet][2]]))
                for key, value in manual.items():
                    self.assertEqual(value, row[key])
            workbook.close()

    def test_cancelled_and_nonreceiving_points_excluded_from_manual_summary(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            data = self.dataset(folder)
            data["points"] = [
                {"status": "invalid", "receive_count": 10, "measured_horizontal_m": 1000},
                {"status": "completed", "receive_count": 0, "measured_horizontal_m": 999},
            ]
            (folder / "range_trial.json").write_text(json.dumps(data))
            workbook = load_workbook(report.export_range(folder))
            overview = dict(workbook["Overview"].values)
            self.assertIsNone(overview["farthest_received_manual_horizontal_m"])
            self.assertIsNone(overview["farthest_received_manual_3d_m"])
            workbook.close()

    def test_gps_quality_rejects_bad_accuracy_and_preserves_unknown_source(self):
        self.assertEqual("source_uncertainty_unknown", report.gps_distance_quality(
            {"distance_m": 12, "location": {"accuracy_m": 5}}, {})["gps_distance_quality"])
        for value in (21, float("nan"), -1, True):
            self.assertEqual("unavailable_or_inaccurate", report.gps_distance_quality(
                {"distance_m": 12, "location": {"accuracy_m": value}},
                {"source_accuracy_m": 4})["gps_distance_quality"])

    def test_old_archive_does_not_invent_manual_distance_or_height(self):
        with TemporaryDirectory() as temporary:
            folder = Path(temporary)
            self.dataset(folder)
            workbook = load_workbook(report.export_range(folder))
            overview = dict(workbook["Overview"].values)
            self.assertIsNone(overview["farthest_received_manual_horizontal_m"])
            self.assertIsNone(overview["farthest_received_manual_3d_m"])
            self.assertIn("diisi manual", overview["height_note"])
            workbook.close()


if __name__ == "__main__":
    unittest.main()
