"""Reporting-only fixtures; never used as physical experiment measurements."""

import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

from resqmesh_controller.config import METHOD_PARAMETERS, research_fingerprint
from resqmesh_controller.excel_report import write_analysis_workbook
from resqmesh_controller.log_merge import summarize_trial
from test_log_merge import delivery_events, trial_event, valid_record


class ExcelReportTest(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.output = Path(temporary.name) / "synthetic.xlsx"

    def workbook(self, *, manifest=None, events=None, trials=None):
        write_analysis_workbook(
            self.output,
            manifest=manifest if manifest is not None else {
                "method_parameters": copy.deepcopy(METHOD_PARAMETERS),
            },
            events=events or [],
            trial_summaries=trials or [],
            aggregates=[],
            attempts=[],
            invalid_trials=[],
        )
        workbook = load_workbook(self.output)
        self.addCleanup(workbook.close)
        return workbook

    @staticmethod
    def rows(workbook, name):
        values = iter(workbook[name].values)
        headers = next(values)
        return [dict(zip(headers, row)) for row in values]

    def test_only_applicable_parameters_are_numeric(self):
        workbook = self.workbook()
        rows = {row["mode"]: row for row in self.rows(workbook, "Method Parameters")}
        basic = rows["basic_flooding"]
        for field in ("imin_ms", "imax_ms", "imax_doublings", "k"):
            self.assertEqual("N/A", basic[field])
        self.assertEqual(2000, basic["basic_wait_ms"])
        self.assertEqual(300, basic["jitter_min_ms"])
        self.assertEqual(1500, basic["jitter_max_ms"])
        for mode in ("trickle", "trickle_no_suppression"):
            row = rows[mode]
            for field in ("basic_wait_ms", "jitter_min_ms", "jitter_max_ms"):
                self.assertEqual("N/A", row[field])
            self.assertEqual(8000, row["imin_ms"])
            self.assertEqual(256000, row["imax_ms"])
            self.assertEqual(5, row["imax_doublings"])
            self.assertIn("[I/2, I)", row["transmission_timing"])
        self.assertEqual(1, rows["trickle"]["k"])
        self.assertEqual("N/A", rows["trickle_no_suppression"]["k"])
        self.assertIs(rows["trickle"]["suppression_enabled"], True)
        self.assertIs(rows["trickle_no_suppression"]["suppression_enabled"], False)
        for row in rows.values():
            self.assertEqual(2000, row["burst_ms"])

    def test_projection_preserves_manifest_parameters_and_fingerprint(self):
        original = copy.deepcopy(METHOD_PARAMETERS)
        manifest = {"method_parameters": copy.deepcopy(original), "trials": {}}
        before = copy.deepcopy(manifest)
        config = {"radio_mode": "coded"}
        fingerprint = research_fingerprint(config)
        workbook = self.workbook(manifest=manifest)
        self.assertEqual(before, manifest)
        self.assertEqual(original, METHOD_PARAMETERS)
        self.assertEqual(fingerprint, research_fingerprint(config))
        metadata = {row["field"]: row["value"]
                    for row in self.rows(workbook, "Manifest Metadata")}
        self.assertEqual(original, json.loads(metadata["method_parameters"]))

    def test_missing_values_are_not_fabricated_and_unknown_modes_are_preserved(self):
        manifest = {"method_parameters": {
            "basic_flooding": {"scheduler": "basic", "burst_ms": 2000},
            "future_mode": {"scheduler": "future", "imin_ms": 123},
        }}
        workbook = self.workbook(manifest=manifest)
        rows = {row["mode"]: row for row in self.rows(workbook, "Method Parameters")}
        self.assertEqual("N/A", rows["basic_flooding"]["imin_ms"])
        self.assertIsNone(rows["basic_flooding"]["basic_wait_ms"])
        self.assertEqual(123, rows["future_mode"]["imin_ms"])

    def test_invalid_and_incomplete_attempts_are_not_delivery_failures(self):
        outcomes = [
            ("SUCCESS", True, 1), ("FAILED_DELIVERY", True, 0),
            ("INVALID", False, None), (None, False, None),
            ("SUCCESS", False, None), ("FAILED_DELIVERY", False, None),
        ]
        trials = [{"trial_id": f"synthetic-{index}", "mode": "basic_flooding",
                   "result": result, "valid": valid, "e2e_latency_ms": None}
                  for index, (result, valid, _) in enumerate(outcomes)]
        before = copy.deepcopy(trials)
        workbook = self.workbook(trials=trials)
        self.assertEqual([value for _, _, value in outcomes],
                         [row["delivery_success"] for row in self.rows(workbook, "Trial Metrics")])
        self.assertEqual(before, trials)

    def test_raw_phy_and_events_preserved_and_workbook_tables_are_valid(self):
        event = {"event_type": "BLE_RX_PHY_OBSERVED", "observation_id": "synthetic-obs",
                 "primary_phy": 3, "secondary_phy": 3, "legacy": False,
                 "coding": "unknown", "timestamp_ms": 12345,
                 "radio": {"primary_phy": "coded", "coding": "unknown"}}
        before = copy.deepcopy(event)
        workbook = self.workbook(events=[event])
        for name in ("RX PHY", "All Events"):
            row = self.rows(workbook, name)[0]
            for key, value in event.items():
                self.assertEqual(value, json.loads(row[key]) if isinstance(value, dict) else row[key])
        self.assertEqual(before, event)
        table_names = []
        for sheet in workbook:
            for table in sheet.tables.values():
                table_names.append(table.displayName)
                headers = [column.name for column in table.tableColumns]
                self.assertEqual(sheet.dimensions, table.ref)
                self.assertEqual(sheet.max_column, len(headers))
                self.assertEqual(len(headers), len(set(headers)))
                self.assertTrue(all(isinstance(header, str) for header in headers))
        self.assertEqual(len(table_names), len(set(table_names)))
        with zipfile.ZipFile(self.output) as archive:
            self.assertIsNone(archive.testzip())
            for name in archive.namelist():
                if name.endswith(".xml"):
                    ElementTree.fromstring(archive.read(name))

    def test_indonesian_explanations_and_claim_limits(self):
        workbook = self.workbook(manifest={"method_parameters": METHOD_PARAMETERS,
                                          "trial_order_strategy": "blocked"})
        definitions = self.rows(workbook, "Metric Definitions")
        self.assertEqual("Semakin tinggi semakin baik", definitions[0]["interpretation"])
        self.assertEqual("Semakin rendah semakin baik", definitions[1]["interpretation"])
        self.assertEqual("Tidak ada data", workbook["Invalid Trials"]["A2"].value)
        overview = {r["field"]: r["value"] for r in self.rows(workbook, "Overview")}
        self.assertIn("smoke", overview["trial_order_note"])
        self.assertIn("belum terverifikasi", overview["phy_claim_note"])
        self.assertIn("bukan bukti jangkauan fisik", overview["topology_claim_note"])

    def test_creation_latency_includes_pre_window_wait_without_changing_e2e(self):
        record = valid_record(observation_started_at_ms=1000, observation_ended_at_ms=2000)
        created = trial_event("SOS_CREATED", node_id="source", message_key="1:2",
                              timestamp_ms=-4100, clock_offset_ms=100)
        events = [created, *delivery_events()]
        baseline = summarize_trial("trial-1", delivery_events(), record)
        result = summarize_trial("trial-1", events, record)
        self.assertEqual(-4000, result["sos_created_at_ms"])
        self.assertEqual(5000, result["source_wait_before_first_advertise_ms"])
        self.assertEqual(5300, result["sos_creation_to_destination_ms"])
        for key in ("result", "e2e_latency_ms", "accepted", "duplicates", "ldr", "transmission_bursts"):
            self.assertEqual(baseline[key], result[key])
        workbook = self.workbook(trials=[result])
        self.assertEqual(5300, self.rows(workbook, "Latency Diagnostics")[0]["sos_creation_to_destination_ms"])
        for change in ({"message_key": "other:2"}, {"node_id": "other"},
                       {"clock_sync_valid": False}):
            missing = summarize_trial("trial-1", [created | change, *delivery_events()], record)
            self.assertIsNone(missing["sos_creation_to_destination_ms"])

    def test_burst_diagnostics_do_not_rewrite_unknown_historical_stop_reason(self):
        events = [
            {"event_type": "ADVERTISE_BURST_STARTED", "trial_id": "synthetic-1", "mode": "basic_flooding",
             "node_id": "source", "burst_id": "burst-1", "elapsed_realtime_ms": 1000,
             "detail_json": '{"target_duration_ms":2000}'},
            {"event_type": "ADVERTISE_BURST_ENDED", "trial_id": "synthetic-1", "mode": "basic_flooding",
             "node_id": "source", "burst_id": "burst-1", "elapsed_realtime_ms": 1510},
        ]
        before = copy.deepcopy(events)
        workbook = self.workbook(events=events)
        row = self.rows(workbook, "Burst Diagnostics")[0]
        self.assertIs(row["shortened"], True)
        self.assertEqual(510, row["actual_duration_ms"])
        self.assertIsNone(row["burst_stop_reason"])
        self.assertIn("alasan tidak tercatat", row["explanation"])
        self.assertEqual(before, events)
        events[1]["detail_json"] = '{"stop_reason":"OBSERVATION_WINDOW_ENDED"}'
        workbook = self.workbook(events=events)
        self.assertEqual("OBSERVATION_WINDOW_ENDED", self.rows(workbook, "Burst Diagnostics")[0]["burst_stop_reason"])


if __name__ == "__main__":
    unittest.main()
